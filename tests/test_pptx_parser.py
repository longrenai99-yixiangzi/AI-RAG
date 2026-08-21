from pptx import Presentation
from pptx.util import Inches

from app.parsers import parse_file


def test_pptx_parser_keeps_slide_location(tmp_path) -> None:
    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "设计管理策划"
    presentation.save(path)

    parsed = parse_file(path, 10)

    assert parsed.parse_status == "parsed"
    assert parsed.blocks[0].location["slide"] == 1
