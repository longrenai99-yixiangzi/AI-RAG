import pytest
from openpyxl import Workbook

from app.document_intelligence.v2 import DocumentIntelligenceV2Builder, sha256
from app.ingestion.xlsx_regions import extract_xlsx_regions, apply_native_display
from app.knowledge_engineering.structured_v2 import build_structured_layer
from app.knowledge_engineering.structured_v2 import _xlsx_answer_cells


@pytest.fixture
def native_book(tmp_path):
    path = tmp_path / 'native.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '组合表'
    sheet['A1'] = '独立合成区域表'
    for ref, value in {'A2':'序号','B2':'名称','D2':'数量','A3':1,'B3':'父项','C3':'子项','D3':0,'A4':'小节标题','D4':2,'G2':'金额区间','G3':'费率','G4':'结果（万元）','H2':'100以下','H3':0.01,'H4':0,'I2':'100以上','I3':0.02,'I4':'=SUM(H4:H4)','M10':999}.items():
        sheet[ref] = value
    sheet.merge_cells('B2:C2')
    sheet.merge_cells('A4:C4')
    workbook.save(path)
    workbook.close()
    return path


def definitions():
    return [{'sheet_name':'组合表','range':'A2:D4','header_rows':[2],'caption_cell':'A1','context_ranges':['A1']}, {'sheet_name':'组合表','range':'G2:I4','orientation':'columns','context_ranges':['A1']}]


def test_region_orientation_groups_and_native_provenance(native_book):
    digest = sha256(native_book)
    result = extract_xlsx_regions(native_book,digest,definitions())
    assert sha256(native_book)==digest
    assert result['tables'][0]['header']==['序号','名称','数量']
    assert result['tables'][1]['header']==['金额区间','费率','结果（万元）']
    first=result['table_rows'][0]
    assert first['cells'][1]['value']=='父项 / 子项'
    assert [cell['cell_reference'] for cell in first['cells'][1]['native_cells']]==["'组合表'!B3","'组合表'!C3"]
    assert first['cells'][2]['value']==0
    assert result['table_rows'][1]['row_context'][0]['value']=='小节标题'
    assert result['table_rows'][-1]['cells'][-1]['value_status']=='FORMULA_CACHE_MISSING'
    assert 'M10' in result['coverage'][0]['unassigned_cells']
    changed=definitions()
    changed[0]['context_ranges'].append('M10')
    assert extract_xlsx_regions(native_book,digest,changed)['tables'][0]['context_id']!=result['tables'][0]['context_id']


def test_scoped_regions_reach_chunks_with_parents_without_auxiliary_leakage(native_book):
    result=extract_xlsx_regions(native_book,sha256(native_book),definitions())
    base=DocumentIntelligenceV2Builder(native_book.parent).build([native_book])
    base.update(tables=result['tables'],table_rows=result['table_rows'],paragraphs=[],atomic_evidence=[])
    layer=build_structured_layer(**{key:base[key] for key in ('documents','sections','paragraphs','tables','table_rows','atomic_evidence')},source_registry={})
    assert len({chunk['chunk_id'] for chunk in layer['semantic_chunks']})==4
    matrix=[chunk for chunk in layer['semantic_chunks'] if chunk['table_id']==result['tables'][1]['table_id']]
    assert {chunk['source_location']['record_column'] for chunk in matrix}=={'H','I'}
    assert all(chunk['parent_chunk_id'] for chunk in matrix)
    assert any(cell['cell_reference']=="'组合表'!H4" and cell['value']==0 for chunk in matrix for cell in chunk['source_cells'])
    assert all('999' not in chunk['raw_text'] and '=SUM' not in chunk['raw_text'] for chunk in layer['semantic_chunks'])


def test_percent_display_requires_native_format_evidence():
    cell={'column':'C','header':'预留金比例','value':0.05,'formula':None,'cached_value':0.05,'display_required':True}
    selected,omitted=_xlsx_answer_cells([cell],['预留金比例'])
    assert selected==[] and omitted[0]['reason']=='DISPLAY_FORMAT_REVIEW_REQUIRED'
    cell['display_text']='5%'
    selected,omitted=_xlsx_answer_cells([cell],['预留金比例'])
    assert selected[0]['value']==0.05 and selected[0]['display_text']=='5%'


def test_empty_merged_formatting_does_not_define_a_header(native_book):
    from openpyxl import load_workbook
    workbook=load_workbook(native_book)
    workbook.active.merge_cells('E2:F2')
    workbook.save(native_book)
    workbook.close()
    specs=definitions()[:1]
    specs[0]['range']='A2:E4'
    result=extract_xlsx_regions(native_book,sha256(native_book),specs)
    assert result['tables'][0]['header']==['序号','名称','数量','列5']


def test_first_workbook_closes_if_cached_reader_fails(native_book,monkeypatch):
    from types import SimpleNamespace
    from app.ingestion import xlsx_regions
    closed=[]
    def load(path,*,data_only):
        if data_only:
            raise OSError('synthetic cached reader failure')
        return SimpleNamespace(close=lambda:closed.append(True))
    monkeypatch.setattr(xlsx_regions,'load_workbook',load)
    with pytest.raises(OSError):
        extract_xlsx_regions(native_book,sha256(native_book),definitions())
    assert closed==[True]


@pytest.mark.parametrize('computed,allowed',[(None,True),('0.0500000000000001',True),('0.25',False)])
def test_native_display_does_not_adopt_a_changed_formula_result(computed,allowed):
    cell={'sheet_name':'S','origin_coordinate':'A1','value':0.05,'display_required':True,'value_status':'NATIVE_VALUE'}
    lookup={'/S/A1':{'text':'5%','format':{'cachedValue':'0.05','computedValue':computed}}}
    apply_native_display(cell,lookup)
    assert ('display_text' in cell) is allowed
    assert cell['value']==0.05


@pytest.mark.parametrize('fault',['version','overlap','header_cut','orientation','oversize'])
def test_invalid_region_profile_fails_closed(native_book,fault):
    digest=sha256(native_book)
    specs=definitions()
    if fault=='version':
        digest='0'*64
    elif fault=='overlap':
        specs.append(specs[0])
    elif fault=='header_cut':
        specs[0]['range']='A2:B4'
    elif fault=='orientation':
        specs[0]['orientation']='other'
    else:
        specs[0]['range']='A1:XFD1048576'
    with pytest.raises(ValueError):
        extract_xlsx_regions(native_book,digest,specs)


@pytest.mark.parametrize('fault',[None,'missing','unknown','duplicate','orientation'])
def test_field_groups_keep_original_columns_and_same_row_link(tmp_path,fault):
    path=tmp_path/'groups.xlsx'
    workbook=Workbook()
    sheet=workbook.active
    sheet.append(['项目名称','金额','单方造价','预留金','合计','单方造价'])
    sheet.append(['独立合成项目',100,10,20,120,12])
    workbook.save(path)
    workbook.close()
    spec={'sheet_name':sheet.title,'range':'A1:F2','header_rows':[1],'field_groups':[['A','B','C'],['A','D','E','F']]}
    if fault=='missing':spec['field_groups'][1].remove('F')
    if fault=='unknown':spec['field_groups'][0].append('G')
    if fault=='duplicate':spec['field_groups'][0].append('A')
    if fault=='orientation':spec['orientation']='columns'
    if fault:
        with pytest.raises(ValueError):extract_xlsx_regions(path,sha256(path),[spec])
        return
    parsed=extract_xlsx_regions(path,sha256(path),[spec])
    assert len(parsed['tables'])==2
    assert all(len(table['header'])==len(set(table['header'])) for table in parsed['tables'])
    assert all(table['requires_column_scope'] for table in parsed['tables'])
    assert parsed['table_rows'][0]['source_table_row_id']==parsed['table_rows'][1]['source_table_row_id']
    assert [[cell['column'] for cell in row['cells']] for row in parsed['table_rows']]==[['A','B','C'],['A','D','E','F']]
    assert parsed['table_rows'][0]['cells'][-1]['value']==10
    assert parsed['table_rows'][1]['cells'][-1]['value']==12
