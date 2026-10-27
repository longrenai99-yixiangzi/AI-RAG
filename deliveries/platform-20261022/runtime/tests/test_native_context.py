import hashlib
from docx import Document
from app.retrieval.native_context import retain_timed_heading


def native(tmp_path):
    p=tmp_path/'会议安排.docx';doc=Document();doc.add_heading('会后落实（24小时内完成）',1);doc.add_paragraph('管理部编制纪要，经确认后分发。');doc.save(p)
    return {'source_path':str(p),'source_version':hashlib.sha256(p.read_bytes()).hexdigest(),'heading_path':'会议流程 > 会后落实（24小时内完成）','text':'管理部编制纪要，经确认后分发。','location':{'native_paragraph_paths':['/body/p[2]']}}


def test_actual_adjacent_timed_heading_is_preserved_without_changing_input(tmp_path):
    r=native(tmp_path);out=retain_timed_heading(r)
    assert out['text'].startswith('会后落实（24小时内完成）\n') and out['location']['native_context_paragraph']==1
    assert '24小时' not in r['text']


def test_unverified_nonadjacent_or_changed_heading_never_enters_evidence(tmp_path):
    r=native(tmp_path)
    for wrong in [{**r,'source_version':'0'*64},{**r,'heading_path':'会后落实（48小时内完成）'},{**r,'location':{}},{**r,'location':{'native_paragraph_paths':['/body/p[1]']}},{**r,'text':'其他段落。'}]:
        assert retain_timed_heading(wrong)==wrong
