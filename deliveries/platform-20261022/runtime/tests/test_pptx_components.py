from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import pytest
from app.ingestion.pptx_components import slide_components,resolve_inherited_text,native_chart_cache


def presentation():
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);return p,s


def test_grouped_text_is_kept_and_notes_are_a_separate_component():
    p,s=presentation();g=s.shapes.add_group_shape();shape=g.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1));shape.text='阶段要求：审核后执行。'
    s.notes_slide.notes_text_frame.text='仅为讲稿说明，不是正文。'
    result=slide_components(s,p.slide_width,p.slide_height)
    assert result['nodes'][0]['text']=='阶段要求：审核后执行。' and result['nodes'][0]['groups']==[str(g.shape_id)]
    assert '讲稿说明' in result['speaker_notes']['text'] and all('讲稿说明' not in n['text'] for n in result['nodes'])


def test_reading_a_slide_without_notes_does_not_create_a_notes_part():
    p,s=presentation();assert not s.has_notes_slide
    assert slide_components(s,p.slide_width,p.slide_height)['speaker_notes']['text']==''
    assert not s.has_notes_slide


def test_missing_geometry_keeps_native_text_in_a_review_component():
    p,s=presentation();shape=s.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1));shape.text='母版继承位置未解析'
    shape._element.spPr.remove(shape._element.spPr.xfrm)
    assert shape.left is None
    result=slide_components(s,p.slide_width,p.slide_height)
    assert not result['nodes'] and result['unlocated_text'][0]['text']=='母版继承位置未解析'
    assert result['unlocated_text'][0]['status']=='TEXT_GEOMETRY_UNRESOLVED'


def test_hidden_slide_and_unsupported_group_rotation_are_marked():
    p,s=presentation();s._element.set('show','0');g=s.shapes.add_group_shape();shape=g.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1));shape.text='隐藏旧版内容';g._element.grpSpPr.xfrm.set('rot','5400000')
    result=slide_components(s,p.slide_width,p.slide_height);reasons=result['nodes'][0]['review_reasons']
    assert result['hidden'] and 'HIDDEN_SLIDE_REVIEW_REQUIRED' in reasons and 'GROUP_TRANSFORM_REVIEW_REQUIRED' in reasons


def test_native_table_blanks_and_merges_do_not_become_zero_or_no():
    p,s=presentation();table=s.shapes.add_table(3,2,Inches(1),Inches(1),Inches(5),Inches(2)).table
    table.cell(0,0).text='阶段';table.cell(0,1).text='人数';table.cell(1,0).text='专项设计';table.cell(1,0).merge(table.cell(2,0))
    result=slide_components(s,p.slide_width,p.slide_height);node=result['nodes'][0]
    assert node['matrix'][1][1]==node['matrix'][2][1]==''
    assert node['native_cells'][2]['is_merge_origin'] and node['native_cells'][4]['is_spanned']
    assert 'MERGED_TABLE_REVIEW_REQUIRED' in node['review_reasons']


def test_inherited_geometry_uses_version_bound_officecli_provenance():
    item={'shape_id':'7','group_ids':[],'text':'设计资源管理'}
    fmt={'id':7,'effective.x':'48pt','effective.y':'42pt','effective.width':'30.48cm','effective.height':'48pt','effective.size':'36pt'}
    for key in ['effective.x','effective.y','effective.width','effective.height']:fmt[key+'.src']='/slide[8]/layout/ph[@type=ctrTitle]'
    native={'success':True,'data':{'matches':1,'results':[{'text':item['text'],'format':fmt}]}}
    node=resolve_inherited_text(item,native,'V','V',12192000,6858000)
    assert node['bbox'][0]==.05 and node['font_size']==36 and node['geometry_origin']=='OFFICECLI_EFFECTIVE_INHERITED'
    with pytest.raises(ValueError):resolve_inherited_text(item,native,'OLD','V',12192000,6858000)
    native['data']['results'][0]['text']='其他文字'
    with pytest.raises(ValueError):resolve_inherited_text(item,native,'V','V',12192000,6858000)


def test_native_chart_cache_preserves_categories_zero_and_missing_values_without_calculation():
    p,s=presentation();data=CategoryChartData();data.categories=['甲','乙','丙'];data.add_series('原始数量',[0,None,3])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(1),Inches(5),Inches(3),data)
    result=slide_components(s,p.slide_width,p.slide_height);cache=result['assets'][0]['native_chart_cache'];plot=cache['plots'][0]
    assert plot['categories']==['甲','乙','丙'] and plot['series'][0]['values']==[0,None,3]
    assert plot['status']=='NATIVE_CATEGORY_CACHE_NOT_RECALCULATED' and not cache['derived_ratios_or_totals']


def test_unsupported_chart_cache_does_not_abort_reading_other_slide_components():
    class Unsupported:
        @property
        def plots(self):raise ValueError('unsupported plot type ofPieChart')
    assert native_chart_cache(Unsupported())['status']=='UNSUPPORTED_NATIVE_CHART_CACHE_REVIEW_REQUIRED'
