from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.document_intelligence.v2 import DocumentIntelligenceV2Builder, document_id


def test_markdown_heading_tree_and_table(tmp_path: Path):
    path = tmp_path / "知识.md"
    path.write_text("# 根\n## 子节\n正文\n\n|字段|值|\n|---|---|\n|A|B|\n", encoding="utf-8")
    result = DocumentIntelligenceV2Builder(tmp_path)._parse_md(path, document_id(path))
    assert [item["heading_level"] for item in result["headings"]] == [1, 2]
    assert result["sections"][1]["parent_section_id"] == result["sections"][0]["section_id"]
    assert result["tables"][0]["header"] == ["字段", "值"]


def test_docx_heading_and_table_parent_section(tmp_path: Path):
    path = tmp_path / "策划书.docx"
    document = Document()
    document.add_heading("项目策划", level=1)
    document.add_paragraph("说明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "专业"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "建筑"
    table.cell(1, 1).text = "1"
    document.save(path)
    result = DocumentIntelligenceV2Builder(tmp_path)._parse_docx(path, document_id(path))
    assert result["headings"][0]["heading_text"] == "项目策划"
    assert result["tables"][0]["section_id"] == result["sections"][0]["section_id"]
    assert result["table_rows"][1]["values"] == ["建筑", "1"]


def test_pdf_page_and_heading_candidate(tmp_path: Path):
    path = tmp_path / "制度.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "1. General\nBody text")
    pdf.save(path)
    pdf.close()
    result = DocumentIntelligenceV2Builder(tmp_path)._parse_pdf(path, document_id(path))
    assert result["sections"][0]["location_start"] == {"page": 1}
    assert result["paragraphs"][0]["location"]["page"] == 1
    assert result["heading_candidates"][0]["text"] == "1. General"


def test_xlsx_sheet_header_merge_and_summary(tmp_path: Path):
    path = tmp_path / "台账.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "价值创造"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "项目台账"
    sheet.append(["专业", "数量"])
    sheet.append(["建筑", 2])
    sheet.append([None, 3])
    sheet.append([None, "=SUM(B3:B4)"])
    workbook.save(path)
    workbook.close()
    result = DocumentIntelligenceV2Builder(tmp_path)._parse_xlsx(path, document_id(path))
    assert result["headings"][0]["heading_text"] == "价值创造"
    assert result["tables"][0]["header"] == ["专业", "数量"]
    assert result["table_rows"][1]["cells"][0]["value"] == "建筑"
    assert result["table_rows"][-1]["is_summary"] is True


def test_pptx_slide_title_and_content_block(tmp_path: Path):
    path = tmp_path / "汇报.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "设计管理"
    slide.placeholders[1].text = "关键要求"
    presentation.save(path)
    result = DocumentIntelligenceV2Builder(tmp_path)._parse_pptx(path, document_id(path))
    assert result["headings"][0]["heading_text"] == "设计管理"
    assert result["sections"][0]["location_start"] == {"slide": 1}
    assert any("关键要求" in item["text"] for item in result["paragraphs"])


def test_registration_query_role_and_metadata_conflict(tmp_path: Path):
    builder = DocumentIntelligenceV2Builder(tmp_path)
    assert builder._build_profile
    from app.document_intelligence.v2 import _document_type

    assert _document_type(Path(r"D:\设计管理\wiki\sources\登记.md"), "", {}) == "REGISTER_PAGE"
    assert _document_type(Path(r"D:\设计管理\wiki\queries\问答.md"), "", {}) == "QUERY_PAGE"
    document = {"document_id": "d", "document_profile": {}}
    conflicts = builder._detect_conflicts(document, Path(r"D:\2025\资料.md"), "正文记载2026年要求")
    assert conflicts[0]["status"] == "METADATA_CONFLICT"


def test_external_lineage_is_partial_without_external_root_resolution(tmp_path: Path):
    path = tmp_path / "登记.md"
    path.write_text("# 登记\n[外部正文](D:/工作/未批准/正文.docx)\n", encoding="utf-8")
    builder = DocumentIntelligenceV2Builder(tmp_path)
    _, relations = builder._build_lineage(path, path.read_text(encoding="utf-8"), document_id(path))
    assert relations[0]["lineage_status"] == "LINEAGE_PARTIAL"
    assert relations[0]["resolved_path"] is None


def test_authority_does_not_override_project_scope():
    from app.ingestion.metadata.governance import GovernanceClassifier

    classifier = GovernanceClassifier()
    policy = classifier.classify(file_name="设计管理制度.md", source_path="D:/设计管理", text="项目案例仅适用于星谷项目")
    case = classifier.classify(file_name="星谷项目复盘.md", source_path="D:/设计管理", text="公司年度复盘")
    assert policy.authority_level == "L1"
    assert case.document_role == "项目案例"
    assert case.authority_level == "L4"


def test_atomic_evidence_legacy_id_gets_parent_links(tmp_path: Path):
    path = tmp_path / "知识.md"
    path.write_text("# 根\n正文\n", encoding="utf-8")
    did = document_id(path)
    builder = DocumentIntelligenceV2Builder(tmp_path)
    atomic = [{
        "evidence_id": "legacy-1",
        "document_id": did,
        "source_path": str(path),
        "file_name": path.name,
        "file_type": ".md",
        "heading_path": "根",
        "location": {"line_start": 2, "line_end": 2},
        "text": "正文",
    }]
    bundle = builder.build([path], atomic)
    evidence = bundle["atomic_evidence"][0]
    assert evidence["section_id"]
    assert evidence["parent_heading_path"] == "根"
    assert bundle["legacy_evidence_id_map"][0]["v2_evidence_id"] == "legacy-1"
