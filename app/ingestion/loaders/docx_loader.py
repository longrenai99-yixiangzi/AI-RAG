from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as WordDocument

from app.domain import SourceBlock


@dataclass(slots=True)
class DOCXLoadResult:
    """Parser result; document content is emitted only as standard SourceBlock objects."""

    status: str
    blocks: list[SourceBlock] = field(default_factory=list)
    error: str | None = None


class DOCXLoader:
    """Read-only DOCX loader kept separate from the active parser and indexer."""

    def load(self, path: Path, document_id: str) -> DOCXLoadResult:
        if path.suffix.lower() != ".docx":
            return DOCXLoadResult(
                status="read_error",
                error=f"不支持的文件扩展名：{path.suffix or '无扩展名'}；旧 .doc 不自动转换",
            )

        try:
            document = WordDocument(path)
            blocks = self._build_blocks(document, path, document_id)
        except OSError as error:
            return DOCXLoadResult(
                status="read_error",
                error=f"{type(error).__name__}: {error}",
            )
        except Exception as error:  # Isolate malformed DOCX packages and parser failures.
            return DOCXLoadResult(
                status="read_error",
                error=f"{type(error).__name__}: {error}",
            )

        return DOCXLoadResult(status="parsed" if blocks else "empty", blocks=blocks)

    @classmethod
    def _build_blocks(
        cls, document: WordDocument, path: Path, document_id: str
    ) -> list[SourceBlock]:
        blocks: list[SourceBlock] = []
        headings: list[str] = []
        section_heading_path = ""
        section: list[tuple[int, str]] = []

        def flush_section() -> None:
            nonlocal section
            non_empty = [(number, text) for number, text in section if text.strip()]
            if not non_empty:
                section = []
                return
            blocks.append(
                SourceBlock(
                    document_id=document_id,
                    source_path=str(path),
                    file_name=path.name,
                    text="\n".join(text for _, text in non_empty).strip(),
                    heading_path=section_heading_path,
                    location={
                        "paragraph_start": non_empty[0][0],
                        "paragraph_end": non_empty[-1][0],
                    },
                )
            )
            section = []

        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            heading_level = cls._heading_level(paragraph.style.name if paragraph.style else "")
            if heading_level is not None:
                flush_section()
                headings = headings[: heading_level - 1] + [text]
                section_heading_path = " > ".join(headings)
            section.append((paragraph_number, text))
        flush_section()

        for table_number, table in enumerate(document.tables, start=1):
            rows = [
                " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells).strip()
                for row in table.rows
            ]
            rows = [row for row in rows if row.strip(" |")]
            if not rows:
                continue
            columns = max((len(row.cells) for row in table.rows), default=0)
            blocks.append(
                SourceBlock(
                    document_id=document_id,
                    source_path=str(path),
                    file_name=path.name,
                    text="表格：\n" + "\n".join(rows),
                    heading_path=section_heading_path,
                    location={"table": table_number, "rows": len(rows), "columns": columns},
                )
            )
        return blocks

    @staticmethod
    def _heading_level(style_name: str) -> int | None:
        match = re.search(r"heading\s*(\d+)|标题\s*(\d+)", style_name or "", re.IGNORECASE)
        if not match:
            return None
        return int(next(value for value in match.groups() if value))


def load_docx(path: Path, document_id: str) -> DOCXLoadResult:
    """Convenience function for callers that do not need a loader instance."""

    return DOCXLoader().load(path, document_id)
