from __future__ import annotations

import hashlib
import os
import re
import stat
import uuid
from pathlib import Path

import pymupdf
from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation

from .domain import ParsedDocument, SourceBlock


SUPPORTED_EXTENSIONS = {".md", ".pdf", ".docx", ".xlsx", ".pptx"}
DEFAULT_EXCLUDED_DIRECTORIES = {
    ".agents",
    ".ai-growth",
    ".claude",
    ".claudian",
    ".copilot",
    ".obsidian",
    ".opencode",
    ".rag",
    "copilot",
    ".git",
    ".trash",
}
TEMPORARY_SUFFIXES = {".tmp", ".bak", ".part"}


def _is_reparse_point(path: Path) -> bool:
    try:
        attributes = getattr(path.stat(follow_symlinks=False), "st_file_attributes", 0)
        return bool(attributes & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0))
    except OSError:
        return True


def iter_source_files(
    root: Path,
    excluded_directories: set[str] = DEFAULT_EXCLUDED_DIRECTORIES,
    scan_errors: list[str] | None = None,
) -> list[Path]:
    files: list[Path] = []
    excluded = {name.lower() for name in excluded_directories}

    def onerror(error: OSError) -> None:
        if scan_errors is not None:
            scan_errors.append(f"{type(error).__name__}: {error.filename or '未知目录'}")

    for current, directory_names, file_names in os.walk(
        root, topdown=True, followlinks=False, onerror=onerror
    ):
        current_path = Path(current)
        directory_names[:] = [
            name
            for name in directory_names
            if name.lower() not in excluded and not _is_reparse_point(current_path / name)
        ]
        for name in file_names:
            path = current_path / name
            if (
                name.startswith("~$")
                or path.suffix.lower() in TEMPORARY_SUFFIXES
                or path.is_symlink()
                or _is_reparse_point(path)
                or path.suffix.lower() not in SUPPORTED_EXTENSIONS
            ):
                continue
            files.append(path)
    return sorted(files, key=lambda item: str(item).lower())


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1_048_576), b""):
            digest.update(block)
    return digest.hexdigest()


def _document_identity(path: Path) -> str:
    try:
        identity = str(path.resolve()).lower()
    except OSError:
        identity = str(path.absolute()).lower()
    return str(uuid.uuid5(uuid.NAMESPACE_URL, identity))


def _block(
    document_id: str,
    path: Path,
    text: str,
    heading_path: str = "",
    **location: int | str,
) -> SourceBlock:
    return SourceBlock(
        document_id=document_id,
        source_path=str(path),
        file_name=path.name,
        text=text.strip(),
        heading_path=heading_path,
        location=location,
    )


def _parse_markdown(path: Path, document_id: str) -> list[SourceBlock]:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    line_offset = 0
    if lines[:1] == ["---"]:
        for index in range(1, len(lines)):
            if lines[index] == "---":
                lines = lines[index + 1 :]
                line_offset = index + 1
                break

    headings: list[str] = []
    section_lines: list[str] = []
    section_start = line_offset + 1
    blocks: list[SourceBlock] = []

    def flush() -> None:
        content = "\n".join(section_lines).strip()
        if content:
            blocks.append(
                _block(
                    document_id,
                    path,
                    content,
                    " > ".join(headings),
                    line_start=section_start,
                )
            )

    for line_number, line in enumerate(lines, start=line_offset + 1):
        match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if not match:
            section_lines.append(line)
            continue
        flush()
        level, title = len(match.group(1)), match.group(2).strip()
        headings[level - 1 :] = [title]
        section_lines = [line]
        section_start = line_number
    flush()
    return blocks


def _parse_pdf(path: Path, document_id: str) -> tuple[list[SourceBlock], str]:
    blocks: list[SourceBlock] = []
    with pymupdf.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if text:
                blocks.append(_block(document_id, path, text, page=page_number))
    return blocks, "parsed" if blocks else "needs_ocr"


def _parse_docx(path: Path, document_id: str) -> list[SourceBlock]:
    document = WordDocument(path)
    blocks: list[SourceBlock] = []
    headings: list[str] = []
    paragraph_buffer: list[str] = []
    paragraph_start = 1

    def flush() -> None:
        content = "\n".join(paragraph_buffer).strip()
        if content:
            blocks.append(
                _block(
                    document_id,
                    path,
                    content,
                    " > ".join(headings),
                    paragraph=paragraph_start,
                )
            )

    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower()
        heading_match = re.search(r"heading\s*(\d+)|标题\s*(\d+)", style_name)
        if heading_match:
            flush()
            level = int(next(value for value in heading_match.groups() if value))
            headings[level - 1 :] = [text]
            paragraph_buffer = [text]
            paragraph_start = paragraph_number
        else:
            if not paragraph_buffer:
                paragraph_start = paragraph_number
            paragraph_buffer.append(text)
    flush()

    for table_number, table in enumerate(document.tables, start=1):
        rows = [
            " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells).strip()
            for row in table.rows
        ]
        rows = [row for row in rows if row.strip(" |")]
        if rows:
            blocks.append(
                _block(
                    document_id,
                    path,
                    "\n".join(rows),
                    " > ".join(headings),
                    table=table_number,
                )
            )
    return blocks


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("\n", " ")


def _parse_xlsx(path: Path, document_id: str) -> list[SourceBlock]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[SourceBlock] = []
    try:
        for sheet in workbook.worksheets:
            headers: list[str] | None = None
            header_row = 0
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [_cell_text(value) for value in row]
                if not any(values):
                    continue
                if headers is None:
                    header_row = row_number
                    headers = [
                        header or f"列{index + 1}" for index, header in enumerate(values)
                    ]
                    continue
                fields = [
                    f"{headers[index]}：{value}"
                    for index, value in enumerate(values)
                    if value and index < len(headers)
                ]
                if fields:
                    blocks.append(
                        _block(
                            document_id,
                            path,
                            f"工作表：{sheet.title}\n" + "\n".join(fields),
                            sheet.title,
                            sheet=sheet.title,
                            row=row_number,
                            header_row=header_row,
                        )
                    )
    finally:
        workbook.close()
    return blocks


def _parse_pptx(path: Path, document_id: str) -> list[SourceBlock]:
    presentation = Presentation(path)
    blocks: list[SourceBlock] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_parts.append(shape.text.strip())
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    text_parts.append(
                        " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
                    )
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except (AttributeError, ValueError):
            notes = ""
        if notes and "Click to add notes" not in notes:
            text_parts.append(f"讲者备注：{notes}")
        content = "\n".join(part for part in text_parts if part.strip())
        if content:
            title = text_parts[0].splitlines()[0][:120] if text_parts else ""
            blocks.append(
                _block(document_id, path, content, title, slide=slide_number, title=title)
            )
    return blocks


def parse_file(path: Path, max_file_size_mb: int) -> ParsedDocument:
    document_id = _document_identity(path)
    try:
        stat_result = path.stat()
    except OSError as error:
        return ParsedDocument(
            document_id=document_id,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            file_size=0,
            mtime_ns=0,
            sha256="",
            blocks=[],
            parse_status="parse_error",
            error=f"{type(error).__name__}: 文件无法读取",
        )
    if stat_result.st_size > max_file_size_mb * 1_024 * 1_024:
        return ParsedDocument(
            document_id=document_id,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            file_size=stat_result.st_size,
            mtime_ns=stat_result.st_mtime_ns,
            sha256="",
            blocks=[],
            parse_status="skipped_large_file",
            error=f"文件大于 {max_file_size_mb} MiB 解析上限",
        )
    try:
        sha256 = _sha256(path)
    except OSError as error:
        return ParsedDocument(
            document_id=document_id,
            source_path=str(path),
            file_name=path.name,
            file_type=path.suffix.lower(),
            file_size=stat_result.st_size,
            mtime_ns=stat_result.st_mtime_ns,
            sha256="",
            blocks=[],
            parse_status="parse_error",
            error=f"{type(error).__name__}: 文件无法读取",
        )
    common = {
        "document_id": document_id,
        "source_path": str(path),
        "file_name": path.name,
        "file_type": path.suffix.lower(),
        "file_size": stat_result.st_size,
        "mtime_ns": stat_result.st_mtime_ns,
        "sha256": sha256,
    }
    try:
        match path.suffix.lower():
            case ".md":
                blocks, status = _parse_markdown(path, document_id), "parsed"
            case ".pdf":
                blocks, status = _parse_pdf(path, document_id)
            case ".docx":
                blocks, status = _parse_docx(path, document_id), "parsed"
            case ".xlsx":
                blocks, status = _parse_xlsx(path, document_id), "parsed"
            case ".pptx":
                blocks, status = _parse_pptx(path, document_id), "parsed"
            case _:
                blocks, status = [], "unsupported"
        return ParsedDocument(**common, blocks=blocks, parse_status=status)
    except Exception as error:  # The index report retains failures without stopping the whole vault.
        return ParsedDocument(
            **common,
            blocks=[],
            parse_status="parse_error",
            error=f"{type(error).__name__}: {error}",
        )
