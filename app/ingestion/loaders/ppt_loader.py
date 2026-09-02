from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from app.domain import SourceBlock


@dataclass(slots=True)
class PPTLoadResult:
    """Parser result; slide content is emitted only as standard SourceBlock objects."""

    status: str
    blocks: list[SourceBlock] = field(default_factory=list)
    error: str | None = None


class PPTLoader:
    """Read-only PPTX loader kept separate from the active parser and indexer."""

    def __init__(self, max_file_size_mb: int | None = None) -> None:
        self.max_file_size_mb = max_file_size_mb

    def load(self, path: Path, document_id: str) -> PPTLoadResult:
        extension = path.suffix.lower()
        if extension == ".ppt":
            return PPTLoadResult(status="unsupported_legacy_format")
        if extension != ".pptx":
            return PPTLoadResult(
                status="read_error",
                error=f"不支持的文件扩展名：{path.suffix or '无扩展名'}",
            )

        try:
            size = path.stat().st_size
            if self.max_file_size_mb is not None and size > self.max_file_size_mb * 1_024 * 1_024:
                return PPTLoadResult(
                    status="read_error",
                    error=f"文件大于 {self.max_file_size_mb} MiB 解析上限",
                )
            presentation = Presentation(path)
            blocks = self._build_blocks(presentation, path, document_id)
        except OSError as error:
            return PPTLoadResult(
                status="read_error",
                error=f"{type(error).__name__}: {error}",
            )
        except Exception as error:  # Isolate malformed or oversized presentations.
            return PPTLoadResult(
                status="read_error",
                error=f"{type(error).__name__}: {error}",
            )

        return PPTLoadResult(status="parsed" if blocks else "empty", blocks=blocks)

    @classmethod
    def _build_blocks(
        cls, presentation: Presentation, path: Path, document_id: str
    ) -> list[SourceBlock]:
        blocks: list[SourceBlock] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text_parts: list[str] = []
            title = ""
            fallback_title = ""
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    rows = [
                        " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
                        for row in shape.table.rows
                    ]
                    rows = [row for row in rows if row.strip(" |")]
                    if rows:
                        text_parts.append("表格：\n" + "\n".join(rows))
                    continue

                if not getattr(shape, "has_text_frame", False):
                    continue
                text = shape.text.strip()
                if not text:
                    continue
                text_parts.append(text)
                if not fallback_title:
                    fallback_title = text.splitlines()[0].strip()
                if not title and cls._is_title_placeholder(shape):
                    title = text.splitlines()[0].strip()

            notes = cls._notes_text(slide)
            if notes:
                text_parts.append(f"讲者备注：{notes}")
            content = "\n".join(part for part in text_parts if part.strip()).strip()
            title = (title or fallback_title)[:120]
            blocks.append(
                SourceBlock(
                    document_id=document_id,
                    source_path=str(path),
                    file_name=path.name,
                    text=content,
                    heading_path=title,
                    location={"slide": slide_number, "title": title},
                )
            )
        return blocks

    @staticmethod
    def _is_title_placeholder(shape: Any) -> bool:
        try:
            if not shape.is_placeholder:
                return False
            return shape.placeholder_format.type in {
                PP_PLACEHOLDER.TITLE,
                PP_PLACEHOLDER.CENTER_TITLE,
            }
        except (AttributeError, ValueError):
            return False

    @staticmethod
    def _notes_text(slide: Any) -> str:
        try:
            notes = slide.notes_slide.notes_text_frame.text
        except (AttributeError, ValueError):
            return ""
        lines = [
            line.strip()
            for line in notes.splitlines()
            if line.strip() and line.strip().lower() != "click to add notes"
        ]
        return "\n".join(lines).strip()


def load_pptx(path: Path, document_id: str) -> PPTLoadResult:
    """Convenience function for callers that do not need a loader instance."""

    return PPTLoader().load(path, document_id)
