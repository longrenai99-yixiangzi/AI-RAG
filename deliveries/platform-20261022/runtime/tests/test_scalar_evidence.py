from copy import deepcopy
import hashlib
import pymupdf

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.ingestion.scalar_evidence import extract_scalars, requested_metrics
from app.scalar_answers import scalar_proposal, render_scalars, packages
from app.layout_answers import layout_count_mismatch, render_layout
from app.retrieval.query_planner_v1 import plan_query
from app.verified_answer_engine_v2 import validate


def config(path, locations):
    return {'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'scope':{'project':'甲产业园项目一期工程','aliases':['甲产业园项目','甲产业园项目一期工程','甲产业园'],'phases':['一期']},
            'admitted_locations':locations}


def ppt_facts(tmp_path):
    path=tmp_path/'native.pptx';prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2)).text='初设概算总投资为54321.67万元，其中工程费43210.56万元。初设概算费用共计43210万元。'
    prs.save(path)
    return extract_scalars(config(path,[{'slide':1}]))[1]


def test_native_amounts_bind_to_distinct_labels_and_unseen_values(tmp_path):
    facts=ppt_facts(tmp_path)
    for name,value in [('初设概算总投资','54321.67'),('工程费','43210.56'),('初设概算费用','43210')]:
        p=scalar_proposal('甲产业园项目一期工程，'+name+'是多少万元？',facts)
        assert p and len(p['facts'])==1 and p['facts'][0]['value']==value
    assert requested_metrics('初设概算总投资和初设概算费用分别是多少')==['初设概算总投资','初设概算费用']
    p=scalar_proposal('甲产业园一期的初设概算费用是多少亿元？',facts)
    assert p['facts'][0]['value']=='4.321'


def test_ppt_key_value_amounts_preserve_conditions_and_distinct_cost_metrics(tmp_path):
    path=tmp_path/'table.pptx';prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6])
    table=s.shapes.add_table(4,2,Inches(1),Inches(1),Inches(8),Inches(2)).table
    values=[['合同总价','123456789.12元（暂定，按施工图预算核对）'],['工程造价','4.2亿'],['付款方式','每月支付75%，另预留100万元'],['价格组成','工程建设其他费用为321万元，预备费为42万元']]
    for row,items in zip(table.rows,values):
        for cell,value in zip(row.cells,items):cell.text=value
    prs.save(path);_,facts=extract_scalars(config(path,[{}]))
    assert len(facts)==4
    proposal=scalar_proposal('甲产业园一期合同金额多少万元？',facts)
    assert proposal['facts'][0]['value']=='12345.678912' and '暂定' in proposal['facts'][0]['text']
    assert proposal['facts'][0]['record']['location']['row']==1
    assert scalar_proposal('甲产业园一期工程造价多少亿元？',facts)['facts'][0]['value']=='4.2'


def test_wrong_phase_unknown_project_and_missing_metric_do_not_borrow_amounts(tmp_path):
    facts=ppt_facts(tmp_path)
    assert scalar_proposal('甲产业园二期初设概算费用是多少？',facts)['scope_missing']
    assert scalar_proposal('武汉甲产业园项目一期工程，初设概算费用是多少？',facts)['scope_missing']
    assert scalar_proposal('武汉甲产业园一期初设概算费用是多少？',facts)['scope_missing']
    assert scalar_proposal('请问甲产业园一期初设概算费用是多少？',facts)['facts']
    assert scalar_proposal('乙产业园一期的初设概算费用是多少？',facts) is None
    missing=scalar_proposal('甲产业园一期的合同金额是多少？',facts)
    assert missing['missing'] and not missing['facts']
    ambiguous=deepcopy(facts[0]);ambiguous['scalar']['value']='99'
    p=scalar_proposal('甲产业园一期初设概算总投资是多少？',[facts[0],ambiguous])
    assert p['missing'] and not p['facts']


def test_unit_confirmation_is_bound_to_source_hash_and_exact_cell(tmp_path):
    path=tmp_path/'native.docx';doc=Document();t=doc.add_table(rows=3,cols=3)
    for row,values in zip(t.rows,[['序号','创效方法','创效金额'],['1','甲支护优化','731W'],['2','乙支护优化','9W']]):
        for c,v in zip(row.cells,values):c.text=v
    doc.save(path);cfg=config(path,[{'table':1}])
    cfg['unit_confirmations']=[{'source_sha256':cfg['sha256'],'location':{'table':1,'row':2,'column':3},'raw_unit':'W','unit':'万元','authority':'PRIOR_OWNER_CONFIRMATION'}]
    _,facts=extract_scalars(cfg)
    assert scalar_proposal('甲产业园一期，甲支护优化创效多少万元？',facts)['facts'][0]['value']=='731'
    unresolved=scalar_proposal('甲产业园一期，乙支护优化创效多少万元？',facts)
    assert unresolved['missing'] and unresolved['reference']['location']['row']==3
    assert not scalar_proposal('甲产业园一期，丙支护优化创效多少万元？',facts)['facts']
    partial=scalar_proposal('甲产业园一期，甲支护优化创效多少，并说明负责人？',facts)
    assert partial['facts'] and partial['missing'] and '负责人' in partial['missing'][0]


def test_scalar_render_needs_retrieved_facts_and_does_not_hide_missing_metrics(tmp_path):
    facts=ppt_facts(tmp_path);question='甲产业园一期的初设概算总投资和合同金额分别是多少万元？'
    bundle={'query_id':'q','question':question,'bundle_status':'VERIFIED_PARTIAL','candidate_evidence':[{**r,'role':'DIRECT'} for r in facts],
            'scalar_source_records':{r['evidence_id']:r for r in facts}}
    answer=render_scalars(bundle)
    assert answer['answer_status']=='PARTIAL_ANSWER' and answer['validation_status']=='VALID'
    assert validate(answer,bundle)['valid']
    assert '54321.67万元' in answer['answer_text'] and '合同金额' in answer['answer_text']
    bundle['candidate_evidence']=[]
    assert render_scalars(bundle) is None


def test_commas_in_single_fact_context_are_not_parallel_metrics():
    for question in ['甲产业园项目一期工程，初设概算费用多少','泸州市垃圾焚烧发电厂扩能工程项目，填埋场边坡支护优化，创效多少万元','西安医学院第一附属医院沣东院区（一期）项目，通过一次成优，助力主体结构提前多少天封顶']:
        assert plan_query(question).subquestions==[]
    assert len(plan_query('甲项目合同金额、工期分别是多少？').subquestions)==2


def test_native_pdf_milestones_bind_to_the_adjacent_date_and_advance(tmp_path):
    path=tmp_path/'native.pdf';pdf=pymupdf.open();page=pdf.new_page()
    page.insert_text((45,60),'02. 地下室封顶 2026/01/15\n提前17天\n03. 主体结构封顶 2026/06/30\n提前83天',fontname='china-s',fontsize=11)
    pdf.save(path);pdf.close();_,facts=extract_scalars(config(path,[{'page':1}]))
    assert len(facts)==2
    p=scalar_proposal('甲产业园一期主体结构提前多少天封顶？',facts)
    assert p['facts'][0]['value']=='83' and p['facts'][0]['record']['location']['date']=='2026/06/30'


def test_package_scope_and_count_premise_guard(tmp_path):
    assert packages('标包二十八')==packages('标包28')=={28}
    facts=ppt_facts(tmp_path)
    for r in facts:r['source_scope']['packages']=[27]
    assert scalar_proposal('甲产业园一期标包28初设概算费用多少？',facts)['scope_missing']
    from tests.test_layout_evidence import sample
    source=sample('控制要点',['甲','乙','丙','丁','戊'],'暖通专业 初步设计阶段管控要点')
    source['layout']['region_count']=1
    q='暖通初步设计阶段管控有哪六个方面？'
    mismatch=layout_count_mismatch(q,plan_query(q).to_dict(),[source])
    assert mismatch and '5项' in mismatch['count_mismatch'] and '6项' in mismatch['count_mismatch']
    bundle={'query_id':'q','question':q,'query_plan':plan_query(q).to_dict(),'bundle_status':'VERIFIED',
            'candidate_evidence':[{**source,'role':'DIRECT'}],'layout_source_records':{source['evidence_id']:source}}
    answer=render_layout(bundle)
    assert answer['answer_status']=='PARTIAL_ANSWER' and answer['uncovered_subquestions']==['SQ1']
    assert not answer['covered_subquestions'] and '不能补造' in answer['answer_text']
