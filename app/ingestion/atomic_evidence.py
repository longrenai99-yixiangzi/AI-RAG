from __future__ import annotations

import hashlib
import re
import uuid
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.ingestion.loaders.markdown_loader import MarkdownLoader
from app.ingestion.loaders.docx_loader import DOCXLoader
from app.ingestion.loaders.pdf_loader import PDFLoader
from app.ingestion.loaders.ppt_loader import PPTLoader
from app.ingestion.normalization.markdown_normalizer import normalize_markdown_text
from app.ingestion.pipeline import build_metadata


ATOMIC_EXTENSIONS = {".md", ".pdf", ".docx", ".xlsx", ".pptx"}
_HEADING_RE = re.compile(r"^(#{1,6})[ \t]+(.+?)\s*$")


def build_atomic_evidence(path: Path, root: Path) -> dict[str, Any]:
    if path.suffix.lower() == ".md":
        return build_markdown_records(path, root)
    if path.suffix.lower() == ".xlsx":
        return build_xlsx_records(path, root)
    if path.suffix.lower() in {".pdf", ".docx", ".pptx"}:
        return build_block_records(path, root)
    return {"path": str(path), "file_type": path.suffix.lower(), "status": "unsupported", "records": [], "error": None}


def build_markdown_records(path: Path, root: Path) -> dict[str, Any]:
    loaded = MarkdownLoader().load(path, _document_id(path))
    base = _base_result(path, loaded.status, loaded.error)
    if loaded.status != "parsed":
        base["metadata"] = build_metadata(path, root, parse_status=loaded.status)
        return base

    try:
        lines = normalize_markdown_text(path.read_bytes().decode("utf-8-sig")).splitlines()
    except (OSError, UnicodeDecodeError) as error:
        base["status"] = "read_error"
        base["error"] = f"{type(error).__name__}: {error}"
        return base

    body_start, front_matter, front_matter_error = MarkdownLoader._parse_front_matter(lines)
    if front_matter_error:
        base["status"] = "front_matter_error"
        base["error"] = front_matter_error
        base["metadata"] = build_metadata(path, root, parse_status="front_matter_error", front_matter=front_matter)
        return base

    headings: list[str] = []
    records: list[dict[str, Any]] = []
    in_fence = False
    first_heading = ""
    document_id = _document_id(path)
    sha256 = _sha256(path)
    for line_number, line in enumerate(lines[body_start:], start=body_start + 1):
        stripped = line.strip()
        if stripped.startswith(("```", "~~~")):
            in_fence = not in_fence
        match = _HEADING_RE.match(line) if not in_fence else None
        if match:
            title = re.sub(r"[ \t]+#+[ \t]*$", "", match.group(2)).strip()
            if title:
                headings = headings[: len(match.group(1)) - 1] + [title]
                first_heading = first_heading or title
        if not stripped:
            continue
        records.append(
            _record(
                path=path,
                document_id=document_id,
                sha256=sha256,
                evidence_key=f"line:{line_number}:{stripped}",
                text=stripped,
                granularity="line",
                heading_path=" > ".join(headings),
                location={"line_start": line_number, "line_end": line_number},
                metadata=None,
            )
        )

    metadata = build_metadata(
        path,
        root,
        parse_status="parsed",
        title=first_heading,
        headers=headings,
        text="\n".join(record["text"] for record in records[:200]),
        front_matter=front_matter,
    )
    for record in records:
        record["metadata"] = metadata
    base["metadata"] = metadata
    base["records"] = records
    return base


def build_xlsx_records(path: Path, root: Path) -> dict[str, Any]:
    base = _base_result(path, "parsed", None)
    document_id = _document_id(path)
    sha256 = _sha256(path)
    records: list[dict[str, Any]] = []
    empty_sheets: list[str] = []
    valid_sheets: list[str] = []
    try:
        formula_workbook = load_workbook(path, read_only=True, data_only=False)
        value_workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        return _base_result(path, "read_error", f"{type(error).__name__}: {error}")

    try:
        for formula_sheet, value_sheet in zip(formula_workbook.worksheets, value_workbook.worksheets, strict=True):
            formula_rows = list(formula_sheet.iter_rows(values_only=True))
            value_rows = list(value_sheet.iter_rows(values_only=True))
            rows = []
            max_columns = 0
            for row_number, formula_values in enumerate(formula_rows, start=1):
                cached_values = value_rows[row_number - 1] if row_number <= len(value_rows) else ()
                if not _row_has_value(formula_values, cached_values):
                    continue
                rows.append((row_number, list(formula_values), list(cached_values)))
                max_columns = max(max_columns, _last_non_empty_column(formula_values, cached_values))
            if not rows:
                empty_sheets.append(formula_sheet.title)
                continue
            valid_sheets.append(formula_sheet.title)
            header_row, header_values = _detect_header(rows, max_columns)
            headers = [
                _cell_text(header_values[index]) or f"列{index + 1}"
                for index in range(max_columns)
            ]
            sheet_text = []
            for row_number, formula_values, cached_values in rows:
                cells, pairs = _cells_and_pairs(headers, formula_values, cached_values, max_columns)
                if not cells:
                    continue
                text = f"第{row_number}行：" + " | ".join(pairs)
                sheet_text.append(text)
                records.append(
                    _record(
                        path=path,
                        document_id=document_id,
                        sha256=sha256,
                        evidence_key=f"sheet:{formula_sheet.title}:row:{row_number}",
                        text=f"工作表：{formula_sheet.title}\n{text}",
                        granularity="row",
                        heading_path=formula_sheet.title,
                        location={
                            "sheet_name": formula_sheet.title,
                            "row_start": row_number,
                            "row_end": row_number,
                            "column_count": max_columns,
                            "header_row": header_row,
                        },
                        metadata=None,
                        extra={"sheet_name": formula_sheet.title, "row_number": row_number, "header": headers, "cells": cells},
                    )
                )
            metadata = build_metadata(
                path,
                root,
                parse_status="parsed",
                title=formula_sheet.title,
                headers=headers,
                text="\n".join(sheet_text[:200]),
            )
            for record in records[-len(sheet_text) :]:
                record["metadata"] = metadata
    except Exception as error:
        return _base_result(path, "read_error", f"{type(error).__name__}: {error}")
    finally:
        formula_workbook.close()
        value_workbook.close()

    base["status"] = "parsed" if records else "empty"
    base["records"] = records
    base["valid_sheets"] = valid_sheets
    base["empty_sheets"] = empty_sheets
    return base


def build_block_records(path: Path, root: Path) -> dict[str, Any]:
    document_id = _document_id(path)
    loaders = {".pdf": PDFLoader, ".docx": DOCXLoader, ".pptx": PPTLoader}
    try:
        loaded = loaders[path.suffix.lower()]().load(path, document_id)
    except Exception as error:
        return _base_result(path, "read_error", f"{type(error).__name__}: {error}")
    base = _base_result(path, loaded.status, loaded.error)
    if not loaded.blocks:
        base["metadata"] = build_metadata(path, root, parse_status=loaded.status)
        return base

    sha256 = base["sha256"]
    records: list[dict[str, Any]] = []
    document_text = "\n".join(block.text for block in loaded.blocks)
    metadata = build_metadata(
        path,
        root,
        parse_status=loaded.status,
        title=next((block.heading_path for block in loaded.blocks if block.heading_path), ""),
        text=document_text[:4_000],
    )
    for block in loaded.blocks:
        lines = [line.strip() for line in block.text.splitlines() if line.strip()]
        if path.suffix.lower() == ".docx" and "table" in block.location and lines[:1] == ["表格："]:
            lines = lines[1:]
        for offset, text in enumerate(lines, start=1):
            location = dict(block.location or {})
            granularity = "text_line"
            if path.suffix.lower() == ".pdf":
                location.update({"text_line_start": offset, "text_line_end": offset})
                granularity = "page_line"
            elif path.suffix.lower() == ".docx":
                if "paragraph_start" in location:
                    paragraph = int(location["paragraph_start"]) + offset - 1
                    location.update({"paragraph_start": paragraph, "paragraph_end": paragraph})
                    granularity = "paragraph"
                elif "table" in location:
                    location.update({"row_start": offset, "row_end": offset})
                    granularity = "table_row"
            elif path.suffix.lower() == ".pptx":
                location["text_index"] = offset
                granularity = "slide_line"
            records.append(
                _record(
                    path=path,
                    document_id=document_id,
                    sha256=sha256,
                    evidence_key=f"{granularity}:{json_location_key(location)}:{offset}:{text}",
                    text=text,
                    granularity=granularity,
                    heading_path=block.heading_path,
                    location=location,
                    metadata=metadata,
                )
            )
    base["metadata"] = metadata
    base["records"] = records
    base["quality"] = getattr(loaded, "quality", None)
    return base


def _base_result(path: Path, status: str, error: str | None) -> dict[str, Any]:
    return {
        "path": str(path),
        "file_name": path.name,
        "file_type": path.suffix.lower(),
        "document_id": _document_id(path),
        "sha256": _sha256(path),
        "status": status,
        "error": error,
        "records": [],
    }


def _record(
    *,
    path: Path,
    document_id: str,
    evidence_key: str,
    sha256: str,
    text: str,
    granularity: str,
    heading_path: str,
    location: dict[str, Any],
    metadata: dict[str, Any] | None,
    extra: dict[str, Any] | None = None,
) -> dict[str, Any]:
    record = {
        "evidence_id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"{document_id}|{evidence_key}")),
        "document_id": document_id,
        "source_path": str(path),
        "file_name": path.name,
        "file_type": path.suffix.lower(),
        "sha256": sha256,
        "granularity": granularity,
        "heading_path": heading_path,
        "location": location,
        "text": text,
        "metadata": metadata,
    }
    if extra:
        record.update(extra)
    return record


def json_location_key(location: dict[str, Any]) -> str:
    return "|".join(f"{key}={location[key]}" for key in sorted(location))


def _cells_and_pairs(
    headers: list[str], formula_values: list[Any], cached_values: list[Any], max_columns: int
) -> tuple[list[dict[str, Any]], list[str]]:
    cells: list[dict[str, Any]] = []
    pairs: list[str] = []
    for index in range(max_columns):
        formula = formula_values[index] if index < len(formula_values) else None
        cached = cached_values[index] if index < len(cached_values) else None
        value = cached if isinstance(formula, str) and formula.startswith("=") and cached is not None else formula
        if value in (None, "") and formula in (None, ""):
            continue
        cell = {"column": get_column_letter(index + 1), "header": headers[index], "value": _json_value(value)}
        if isinstance(formula, str) and formula.startswith("="):
            cell["formula"] = formula
            cell["cached_value"] = _json_value(cached)
        cells.append(cell)
        pairs.append(f"{headers[index]}：{_cell_text(value)}")
    return cells, pairs


def _detect_header(rows: list[tuple[int, list[Any], list[Any]]], max_columns: int) -> tuple[int, list[Any]]:
    for row_number, formula_values, cached_values in rows[:20]:
        values = [
            cached if isinstance(formula, str) and formula.startswith("=") and cached is not None else formula
            for formula, cached in zip(formula_values, cached_values, strict=False)
        ]
        text_values = {_cell_text(value) for value in values}
        if len([value for value in text_values if value]) >= 2 and (
            any("项目名称" in value or "专业类别" in value for value in text_values)
            or any("序号" == value or "设计阶段" in value for value in text_values)
        ):
            return row_number, values[:max_columns]
    row_number, formula_values, cached_values = rows[0]
    values = [
        cached if isinstance(formula, str) and formula.startswith("=") and cached is not None else formula
        for formula, cached in zip(formula_values, cached_values, strict=False)
    ]
    return row_number, values[:max_columns]


def _row_has_value(formula_values: tuple[Any, ...], cached_values: tuple[Any, ...]) -> bool:
    return any(value not in (None, "") for value in (*formula_values, *cached_values))


def _last_non_empty_column(*values_list: tuple[Any, ...]) -> int:
    last = 0
    for values in values_list:
        for index, value in enumerate(values, start=1):
            if value not in (None, ""):
                last = max(last, index)
    return last


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(_json_value(value)).strip().replace("\n", " ")


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _document_id(path: Path) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, str(path.resolve()).lower()))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1_048_576), b""):
                digest.update(block)
    except OSError:
        return ""
    return digest.hexdigest()
