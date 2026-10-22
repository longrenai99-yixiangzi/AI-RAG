from pathlib import Path
from docx import Document
from app.document_intelligence.v2 import DocumentIntelligenceV2Builder,document_id


def test_table_caption_and_section_follow_body_order_not_first_and_last(tmp_path):
    doc=Document();doc.add_heading('人员组织',level=1);doc.add_paragraph('表 人员配置')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='岗位';table.cell(0,1).text='人数';table.cell(1,0).text='设计负责人';table.cell(1,1).text='1'
    doc.add_heading('实施效果',level=1);doc.add_paragraph('具体措施如下：')
    other=doc.add_table(rows=3,cols=3);other.cell(0,0).merge(other.cell(0,2)).text='降概措施清单'
    other.cell(1,0).text='序号';other.cell(1,1).text='内容';other.cell(1,2).text='金额'
    other.cell(2,0).text='1';other.cell(2,1).text='措施甲';other.cell(2,2).text='100'
    path=tmp_path/'source.docx';doc.save(path);data=DocumentIntelligenceV2Builder(tmp_path)._parse_docx(path,document_id(path))
    sections={s['section_id']:s for s in data['sections']}
    assert data['tables'][0]['table_title']=='表 人员配置'
    assert data['tables'][1]['table_title']=='降概措施清单'
    assert '人员组织' in sections[data['tables'][0]['section_id']]['heading_path']
    assert '实施效果' in sections[data['tables'][1]['section_id']]['heading_path']
    assert data['tables'][1]['source_location']['body_order']>data['tables'][0]['source_location']['body_order']


def test_caption_is_consumed_and_numeric_data_is_not_promoted_to_header(tmp_path):
    doc=Document();doc.add_paragraph('表 首表')
    first=doc.add_table(rows=2,cols=2);first.cell(0,0).text='字段';first.cell(0,1).text='内容'
    second=doc.add_table(rows=2,cols=2);second.cell(0,0).text='字段';second.cell(0,1).text='内容'
    third=doc.add_table(rows=2,cols=2);third.cell(0,0).merge(third.cell(0,1)).text='措施清单';third.cell(1,0).text='1';third.cell(1,1).text='业务数据'
    path=tmp_path/'source.docx';doc.save(path);data=DocumentIntelligenceV2Builder(tmp_path)._parse_docx(path,document_id(path))
    assert data['tables'][0]['table_title']=='表 首表'
    assert '未标明' in data['tables'][1]['table_title']
    assert data['tables'][2]['header']==[] and data['tables'][2]['quality_flags']==['HEADER_NOT_DECLARED']
    rows=[r for r in data['table_rows'] if r['table_id']==data['tables'][2]['table_id']]
    assert rows[0]['is_table_preamble'] and not rows[1]['is_header'] and rows[1]['values']==['1','业务数据']
