from docx import Document

from app.parsers import parse_file


def test_docx_parser_keeps_paragraph_location(tmp_path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("设计策划", level=1)
    document.add_paragraph("项目启动后应编制设计策划。")
    document.save(path)

    parsed = parse_file(path, 10)

    assert parsed.parse_status == "parsed"
    assert parsed.blocks
    assert parsed.blocks[0].location["paragraph"] == 1
