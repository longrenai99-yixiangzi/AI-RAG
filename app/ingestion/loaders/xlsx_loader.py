from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.domain import SourceBlock


@dataclass(slots=True)
class XLSXLoadResult:
    """Parser result; worksheet content is emitted only as standard SourceBlock objects."""

    status: str
    blocks: list[SourceBlock] = field(default_factory=list)
    workbook_name: str = ""
    sheet_names: list[str] = field(default_factory=list)
    error: str | None = None


class XLSXLoader:
    """Read-only XLSX loader kept separate from the active parser and indexer."""

    def load(self, path: Path, document_id: str) -> XLSXLoadResult:
        workbook_name = path.name
        if path.suffix.lower() != ".xlsx":
            return XLSXLoadResult(
                status="read_error",
                workbook_name=workbook_name,
                error=f"不支持的文件扩展名：{path.suffix or '无扩展名'}；不自动转换旧格式",
            )

        workbook: Any = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_names = list(workbook.sheetnames)
            blocks = [
                block
                for sheet in workbook.worksheets
                if (block := self._build_sheet_block(sheet, path, document_id)) is not None
            ]
        except OSError as error:
            return XLSXLoadResult(
                status="read_error",
                workbook_name=workbook_name,
                error=f"{type(error).__name__}: {error}",
            )
        except Exception as error:  # Isolate malformed workbooks and worksheet failures.
            return XLSXLoadResult(
                status="read_error",
                workbook_name=workbook_name,
                error=f"{type(error).__name__}: {error}",
            )
        finally:
            if workbook is not None:
                workbook.close()

        return XLSXLoadResult(
            status="parsed" if blocks else "empty",
            blocks=blocks,
            workbook_name=workbook_name,
            sheet_names=sheet_names,
        )

    @classmethod
    def _build_sheet_block(cls, sheet: Any, path: Path, document_id: str) -> SourceBlock | None:
        non_empty_rows: list[tuple[int, list[str]]] = []
        max_column = 0
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [cls._cell_text(value) for value in row]
            last_column = cls._last_non_empty_column(values)
            if last_column == 0:
                continue
            non_empty_rows.append((row_number, values[:last_column]))
            max_column = max(max_column, last_column)

        if not non_empty_rows:
            return None

        header_row, header_values = non_empty_rows[0]
        header_values = header_values + [""] * (max_column - len(header_values))
        headers = [
            header_values[index].strip() or f"列{index + 1}"
            for index in range(max_column)
        ]
        lines = [
            f"工作表：{sheet.title}",
            f"表头（第{header_row}行）：" + " | ".join(headers),
        ]
        for row_number, values in non_empty_rows[1:]:
            normalized = values + [""] * (max_column - len(values))
            pairs = [f"{headers[index]}：{normalized[index]}" for index in range(max_column)]
            lines.append(f"第{row_number}行：" + " | ".join(pairs))

        return SourceBlock(
            document_id=document_id,
            source_path=str(path),
            file_name=path.name,
            text="\n".join(lines),
            heading_path=sheet.title,
            location={
                "sheet_name": sheet.title,
                "row_start": non_empty_rows[0][0],
                "row_end": non_empty_rows[-1][0],
                "column_count": max_column,
                "header_row": header_row,
            },
        )

    @staticmethod
    def _cell_text(value: object) -> str:
        if value is None:
            return ""
        return str(value).strip().replace("\n", " ")

    @staticmethod
    def _last_non_empty_column(values: list[str]) -> int:
        for index in range(len(values) - 1, -1, -1):
            if values[index]:
                return index + 1
        return 0


def load_xlsx(path: Path, document_id: str) -> XLSXLoadResult:
    """Convenience function for callers that do not need a loader instance."""

    return XLSXLoader().load(path, document_id)
