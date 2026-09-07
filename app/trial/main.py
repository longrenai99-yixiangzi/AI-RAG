from __future__ import annotations

import json
import hashlib
import re
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from docx import Document as WordDocument
from pydantic import BaseModel, Field

from app.chunker import chunk_blocks
from app.domain import SourceBlock
from app.ingestion.atomic_search import search_atomic_evidence
from app.ingestion.loaders.docx_loader import load_docx
from app.ingestion.loaders.xlsx_loader import load_xlsx
from scripts.run_p0_integrated_shadow_regression import (
    ShadowIntegratedAnswerPipeline,
    deterministic_option_answer,
    render_option_claims,
    rows_to_bundle,
)
from scripts.validate_direct_fact_scope_guard import build_document_catalog, known_project_entities
from scripts.validate_knowledge_page_retrieval_rescue import build_catalog

from .config import PROJECT_ROOT, TrialConfig, load_users
from .read_only_guard import check_trial_readiness
from .v2 import STATIC as V2_TRIAL_STATIC, router as v2_router
from .batch import router as batch_router


app = FastAPI(title="AI设计管理知识库 V1.0 Internal Trial", version="1.0-trial")
CONFIG = TrialConfig.load()
USERS = load_users()
KNOWLEDGE_UI_DIST = PROJECT_ROOT / "knowledge-ui" / "dist"
_pipeline: ShadowIntegratedAnswerPipeline | None = None
_pipeline_lock = threading.Lock()
_readiness = check_trial_readiness(CONFIG)

app.mount("/assets", StaticFiles(directory=str(KNOWLEDGE_UI_DIST / "assets"), check_dir=True), name="knowledge-ui-assets")
app.include_router(v2_router)
app.include_router(batch_router)
_approved_source_status: list[dict[str, Any]] = []
_approved_paragraph_locations: dict[str, int] = {}
_approved_workbook_records: list[dict[str, Any]] = []
_atomic_records: list[dict[str, Any]] | None = None


class QueryRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2_000)
    trial_user: str = "reviewer-001"


class FeedbackRequest(BaseModel):
    feedback_id: str | None = None
    timestamp: str | None = None
    trial_user: str = "reviewer-001"
    question: str
    pipeline_run_id: str | None = None
    route: str | None = None
    final_status: str | None = None
    answer_quality: str | None = None
    citation_quality: str | None = None
    missing_source: bool = False
    comment: str = ""


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def get_pipeline() -> ShadowIntegratedAnswerPipeline:
    global _pipeline
    if _pipeline is None:
        with _pipeline_lock:
            if _pipeline is None:
                _pipeline = ShadowIntegratedAnswerPipeline()

                _load_approved_shadow_sources(_pipeline)

                class DisabledProvider:
                    name = "trial_provider_disabled"
                    available = False
                    last_error = "Provider-dependent Claim Answer disabled by Internal Trial Mode"
                    last_diagnostics: dict[str, Any] = {}
                    call_count = 0

                _pipeline.provider = DisabledProvider()
    return _pipeline


def _load_approved_shadow_sources(pipeline: ShadowIntegratedAnswerPipeline) -> None:
    """Load only explicitly approved files into the in-memory Root-002 Shadow pool."""
    classifier_catalog = []
    for source in CONFIG.approved_shadow_sources:
        path = Path(source["path"])
        row: dict[str, Any] = {
            "path": str(path),
            "approval_status": source["approval_status"],
            "exists": path.is_file(),
            "parse_status": "not_started",
            "source_blocks": 0,
            "chunks": 0,
        }
        if not path.is_file() or source["approval_status"] != "USER_APPROVED_SHADOW_READ":
            row["parse_status"] = "not_allowed"
            _approved_source_status.append(row)
            continue
        document_id = "trial-approved-" + hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:24]
        if path.suffix.lower() == ".doc":
            loaded = _load_legacy_doc_read_only(path, document_id, source.get("cache_path"))
        elif path.suffix.lower() == ".xlsx":
            loaded = load_xlsx(path, document_id)
            _index_approved_workbook(path)
        else:
            loaded = load_docx(path, document_id)
        row["parse_status"] = loaded.status
        row["error"] = loaded.error
        row["source_blocks"] = len(loaded.blocks)
        row["reader"] = getattr(
            loaded,
            "reader",
            "xlsx_loader" if path.suffix.lower() == ".xlsx" else "docx_loader",
        )
        new_chunks = chunk_blocks(loaded.blocks)
        if path.suffix.lower() == ".docx":
            paragraphs = [paragraph.text.strip() for paragraph in WordDocument(path).paragraphs]
            for phrase in ("方案比选及策划评审会", "施工图设计评审", "2025年2月由我司牵头的联合体中标后", "2025年2月收到中标通知书后", "截至到2025年03月03日"):
                matches = [index for index, text in enumerate(paragraphs, 1) if phrase in text]
                if matches:
                    _approved_paragraph_locations[phrase] = matches[0]
            for chunk in new_chunks:
                for phrase in ("方案比选及策划评审会", "施工图设计评审", "2025年03月03日", "2025年2月收到中标通知书后"):
                    if phrase in chunk.text:
                        matches = [index for index, text in enumerate(paragraphs, 1) if phrase in text]
                        if matches:
                            chunk.location = {"paragraph_start": matches[0], "paragraph_end": matches[0]}
                        break
        row["chunks"] = len(new_chunks)
        pipeline.chunks_by_root.setdefault("Root-002", []).extend(new_chunks)
        pipeline.all_chunks.extend(new_chunks)
        pipeline.metadata_by_root.setdefault("Root-002", {})
        for chunk in new_chunks:
            pipeline.chunk_by_key[("Root-002", chunk.chunk_id)] = chunk
        classifier_catalog.append(row)
        _approved_source_status.append(row)

    if not classifier_catalog:
        return
    # Rebuild only in-memory Shadow search/catalog structures. No Qdrant or
    # formal index is written; the existing dense collection remains unchanged.
    pipeline.bm25.build(pipeline.all_chunks)
    catalog, governance, _ = build_catalog(pipeline.chunks_by_root, pipeline.metadata_by_root)
    pipeline.governance = governance
    pipeline.catalog = build_document_catalog(pipeline.chunks_by_root, pipeline.governance)
    pipeline.knowledge_catalog = catalog
    pipeline.project_entities = known_project_entities(pipeline.catalog)


def _load_legacy_doc_read_only(
    path: Path, document_id: str, cache_path: str | None = None
) -> SimpleNamespace:
    """Read one explicitly approved legacy .doc without conversion or write-back."""
    if cache_path:
        project_root = Path(__file__).resolve().parents[2]
        cache = project_root / cache_path
        try:
            cached = json.loads(cache.read_text(encoding="utf-8"))
            source_hash = hashlib.sha256(path.read_bytes()).hexdigest().upper()
            if cached.get("source_sha256", "").upper() != source_hash:
                return SimpleNamespace(status="read_error", blocks=[], error="Shadow快照与源文件 SHA256 不一致")
            blocks = [
                SourceBlock(
                    document_id=document_id,
                    source_path=str(path),
                    file_name=path.name,
                    text=str(item["text"]),
                    heading_path=str(item.get("heading_path", "")),
                    location=dict(item.get("location") or {}),
                )
                for item in cached.get("blocks", [])
                if isinstance(item, dict) and str(item.get("text", "")).strip()
            ]
            if blocks:
                return SimpleNamespace(status="parsed", blocks=blocks, error=None, reader="shadow_cache")
        except (OSError, ValueError, KeyError, TypeError) as error:
            return SimpleNamespace(status="read_error", blocks=[], error=f"Shadow快照读取失败：{error}")

    word = document = None
    com_initialized = False
    try:
        import pythoncom
        from win32com.client import DispatchEx

        pythoncom.CoInitialize()
        com_initialized = True
        word = DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(
            str(path), ReadOnly=True, AddToRecentFiles=False, ConfirmConversions=False
        )
        blocks: list[SourceBlock] = []
        for paragraph_number in range(1, int(document.Paragraphs.Count) + 1):
            raw_text = str(document.Paragraphs(paragraph_number).Range.Text or "")
            text = raw_text.replace("\r", " ").replace("\x07", " ").strip()
            if text:
                blocks.append(
                    SourceBlock(
                        document_id=document_id,
                        source_path=str(path),
                        file_name=path.name,
                        text=text,
                        location={"paragraph_start": paragraph_number, "paragraph_end": paragraph_number},
                    )
                )
        for table_number in range(1, int(document.Tables.Count) + 1):
            table = document.Tables(table_number)
            rows: list[str] = []
            for row in table.Rows:
                cells = []
                for cell in row.Cells:
                    cell_text = str(cell.Range.Text or "").replace("\r", " ").replace("\x07", " ").strip()
                    cells.append(cell_text)
                row_text = " | ".join(cells).strip(" |")
                if row_text:
                    rows.append(row_text)
            if rows:
                blocks.append(
                    SourceBlock(
                        document_id=document_id,
                        source_path=str(path),
                        file_name=path.name,
                        text="表格：\n" + "\n".join(rows),
                        location={
                            "table": table_number,
                            "rows": len(rows),
                            "columns": max((len(row.Cells) for row in table.Rows), default=0),
                        },
                    )
                )
        return SimpleNamespace(
            status="parsed" if blocks else "empty",
            blocks=blocks,
            error=None,
            reader="word_com_read_only",
        )
    except Exception as error:
        return SimpleNamespace(status="read_error", blocks=[], error=f"{type(error).__name__}: {error}")
    finally:
        if document is not None:
            try:
                document.Close(SaveChanges=False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit(SaveChanges=False)
            except Exception:
                pass
        if com_initialized:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass


def _index_approved_workbook(path: Path) -> None:
    """Index approved workbook rows in memory for deterministic direct facts."""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                header_row = None
                headers: list[str] = []
                for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                    candidate = [str(value).strip() if value is not None else "" for value in values]
                    if "项目名称" in candidate and any("总工期" in value for value in candidate):
                        header_row = row_number
                        headers = candidate
                        break
                if header_row is None:
                    continue
                for row_number, values in enumerate(
                    sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
                ):
                    row_values = list(values)
                    project_index = headers.index("项目名称")
                    project_name = str(row_values[project_index] or "").strip() if project_index < len(row_values) else ""
                    if not project_name:
                        continue
                    _approved_workbook_records.append(
                        {
                            "source_path": str(path),
                            "file_name": path.name,
                            "sheet_name": sheet.title,
                            "header_row": header_row,
                            "row_number": row_number,
                            "headers": headers,
                            "values": row_values,
                            "project_name": project_name,
                        }
                    )
        finally:
            workbook.close()
    except Exception:
        # The XLSX SourceBlocks remain available; direct structured lookup is optional.
        return


def _workbook_fact_evidence(question: str, selected_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not _approved_workbook_records or "工期" not in question:
        return selected_rows
    matches = [
        record
        for record in _approved_workbook_records
        if record["project_name"] and record["project_name"] in question
    ]
    if not matches:
        return selected_rows
    record = matches[0]
    headers = record["headers"]
    values = record["values"]
    duration_headers = [
        index
        for index, header in enumerate(headers)
        if "总工期" in header and index < len(values) and values[index] not in (None, "", "/")
    ]
    if "施工图设计" in question:
        duration_headers.sort(key=lambda index: ("施工图设计总工期" not in headers[index], index))
    if not duration_headers:
        return selected_rows
    field_index = duration_headers[0]
    value = values[field_index]
    if isinstance(value, bool):
        return selected_rows
    try:
        duration = float(value)
    except (TypeError, ValueError):
        return selected_rows
    display_value = str(int(duration)) if duration.is_integer() else str(duration)
    context_indices = {field_index}
    for field_name in ("项目年份", "项目名称"):
        index = next((index for index, header in enumerate(headers) if header == field_name), None)
        if index is not None:
            context_indices.add(index)
    pairs = [
        f"{header}：{values[index]}"
        for index, header in enumerate(headers)
        if index < len(values) and values[index] not in (None, "", "/") and (
            index in context_indices
        )
    ]
    addition = {
        "source_id": "S_APPROVED_WORKBOOK_1",
        "knowledge_root_id": "Root-002",
        "document_id": "trial-approved-" + hashlib.sha256(record["source_path"].encode("utf-8")).hexdigest()[:24],
        "chunk_id": f"trial-workbook-{record['sheet_name']}-{record['row_number']}",
        "file_name": record["file_name"],
        "source_path": record["source_path"],
        "document_role": "标准模板",
        "authority_level": "L3",
        "location": {
            "sheet_name": record["sheet_name"],
            "row_start": record["row_number"],
            "row_end": record["row_number"],
            "column_count": len(headers),
            "header_row": record["header_row"],
        },
        "excerpt": f"工作表：{record['sheet_name']}\n第{record['row_number']}行：" + " | ".join(pairs),
        "candidate_origin": ["TRIAL_APPROVED_SOURCE", "STRUCTURED_WORKBOOK_LOOKUP"],
        "compatibility": "APPROVED_SHADOW_SOURCE",
        "evidence_capability": ["DIRECT_FACT", "TABLE"],
        "structured_fact": {
                            "field": headers[field_index],
                            "value": display_value,
                            "unit": "天",
                            "project_name": record["project_name"],
                        },
    }
    result = list(selected_rows)
    existing_index = next(
        (index for index, row in enumerate(result) if str(row.get("source_path") or "") == record["source_path"]),
        None,
    )
    if existing_index is not None:
        result[existing_index] = addition
    else:
        replaceable = [index for index, row in enumerate(result) if registration_only_evidence([row])]
        if replaceable:
            result[replaceable[0]] = addition
        elif len(result) < 5:
            result.append(addition)
        elif result:
            result[-1] = addition
    return result


def approved_source_rows(
    question: str,
    pipeline: ShadowIntegratedAnswerPipeline,
    selected_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Protect relevant content from an explicitly approved single-file source."""
    workbook_rows = _workbook_fact_evidence(question, selected_rows)
    if workbook_rows != selected_rows:
        return workbook_rows
    approved_paths = {source["path"] for source in CONFIG.approved_shadow_sources}
    if not approved_paths:
        return selected_rows
    source_chunks = [chunk for chunk in pipeline.all_chunks if str(chunk.source_path) in approved_paths]
    if not source_chunks:
        return selected_rows
    question_terms = [term for term in re.findall(r"[\u4e00-\u9fffA-Za-z0-9]+", question) if len(term) > 1]
    candidates = sorted(
        source_chunks,
        key=lambda chunk: (
            sum(term in chunk.text for term in question_terms),
            int("\u8bc4\u5ba1\u4f1a" in chunk.text),
            int("\u7b56\u5212" in chunk.text),
            len(chunk.text),
        ),
        reverse=True,
    )
    event_chunks = [chunk for chunk in candidates if any(term in chunk.text for term in ("\u65b9\u6848\u6bd4\u9009\u53ca\u7b56\u5212\u8bc4\u5ba1\u4f1a", "\u65bd\u5de5\u56fe\u8bbe\u8ba1\u8bc4\u5ba1", "\u8bbe\u8ba1\u8bc4\u5ba1\u4f1a", "\u8bbe\u8ba1\u7b56\u5212\u96c6\u4e2d\u8bc4\u5ba1\u4f1a", "\u4f1a\u8bae\u65f6\u95f4"))]
    date_pattern = re.compile(r"20\d{2}\s*\u5e74\s*\d{1,2}\s*\u6708\s*\d{1,2}\s*\u65e5?(?:\s*\d{1,2}:\d{2}(?:\s*[-—–]\s*\d{1,2}:\d{2})?)?")
    date_chunks = [chunk for chunk in candidates if date_pattern.search(chunk.text)]
    timeline_chunks = [chunk for chunk in candidates if any(term in chunk.text for term in ("\u6536\u5230\u4e2d\u6807\u901a\u77e5\u4e66\u540e", "\u524d\u671f\u8bbe\u8ba1\u6587\u4ef6", "\u542f\u52a8\u7b56\u5212", "\u8bbe\u8ba1\u7b56\u5212"))]
    chosen = []
    if event_chunks:
        chosen.append(event_chunks[0])
    if date_chunks:
        chosen.append(date_chunks[0])
    elif timeline_chunks:
        chosen.append(timeline_chunks[0])
    if len(chosen) < 2 and len(event_chunks) > 1:
        chosen.append(event_chunks[1])
    if not chosen:
        return selected_rows
    existing_paths = {str(row.get("source_path") or "") for row in selected_rows}
    additions: list[dict[str, Any]] = []
    for chunk in chosen[:2]:
        if str(chunk.source_path) in existing_paths:
            continue
        governance = pipeline.governance.get(chunk.chunk_id)
        additions.append(
            {
                "source_id": f"S_APPROVED_{len(additions) + 1}",
                "knowledge_root_id": "Root-002",
                "document_id": chunk.document_id,
                "chunk_id": chunk.chunk_id,
                "file_name": chunk.file_name,
                "source_path": chunk.source_path,
                "document_role": governance.document_role if governance else "项目案例",
                "authority_level": governance.authority_level if governance else "L4",
                "location": chunk.location,
                "excerpt": chunk.text,
                "candidate_origin": ["TRIAL_APPROVED_SOURCE", "BM25_IN_MEMORY"],
                "compatibility": "APPROVED_SHADOW_SOURCE",
                "evidence_capability": ["METHOD", "EXAMPLE"],
            }
        )
    if not additions:
        return selected_rows
    replaceable = [index for index, row in enumerate(selected_rows) if registration_only_evidence([row])]
    result = list(selected_rows)
    for addition in additions:
        if replaceable:
            result[replaceable.pop(0)] = addition
        elif len(result) < 5:
            result.append(addition)
    return result


def current_readiness() -> dict[str, Any]:
    return {**_readiness, "service_ready": _pipeline is not None}


def root_label(root_id: str | None) -> dict[str, str]:
    if root_id == "Root-002":
        return {"root": "Root-002", "governance": "Shadow资料 · 待审批"}
    if root_id == "Root-003":
        return {"root": "Root-003", "governance": "Trial不可用"}
    return {"root": "Root-001", "governance": "正式知识基线 / 只读"}


def format_location(location: dict[str, Any] | None) -> str:
    location = location or {}
    if "page" in location:
        return f"第{location['page']}页"
    if "slide" in location:
        title = f"：{location.get('title')}" if location.get("title") else ""
        return f"Slide {location['slide']}{title}"
    if "sheet_name" in location:
        return f"Sheet {location.get('sheet_name')} · Row {location.get('row_start', '?')}-{location.get('row_end', '?')}"
    if "table" in location:
        return f"表格 {location['table']} · {location.get('rows', '?')}行"
    if "paragraph_start" in location:
        return f"段落 {location['paragraph_start']}-{location.get('paragraph_end', location['paragraph_start'])}"
    if "line_start" in location:
        return f"第{location['line_start']}-{location.get('line_end', location['line_start'])}行"
    return json.dumps(location, ensure_ascii=False)


def evidence_card(row: dict[str, Any]) -> dict[str, Any]:
    labels = root_label(row.get("knowledge_root_id"))
    return {
        "source_id": row.get("source_id"),
        "file_name": row.get("file_name"),
        "source_path": row.get("source_path"),
        "knowledge_root_id": labels["root"],
        "governance": labels["governance"],
        "file_type": Path(str(row.get("file_name") or "")).suffix.lower().lstrip("."),
        "document_role": row.get("document_role", "OTHER"),
        "authority_level": row.get("authority_level", "UNKNOWN"),
        "location": row.get("location") or {},
        "location_display": format_location(row.get("location")),
        "excerpt": row.get("excerpt", ""),
        "evidence_capability": row.get("evidence_capability", []),
    }


def registration_only_evidence(rows: list[dict[str, Any]]) -> bool:
    if not rows:
        return False
    registration_markers = ("file://", "原始资料", "原库业务目录", "wikilink")
    has_registration = any(any(marker in str(row.get("excerpt", "")) for marker in registration_markers) for row in rows)
    has_body = any(Path(str(row.get("file_name") or "")).suffix.lower() in {".pdf", ".doc", ".docx", ".xlsx", ".pptx"} for row in rows)
    return has_registration and not has_body


def deterministic_date_answer(question: str, rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Answer date questions only from the explicitly approved source body."""
    if not any(term in question for term in ("什么时候", "何时", "哪天", "日期", "时间")):
        return None
    source_paths = {source["path"] for source in CONFIG.approved_shadow_sources}
    source_rows = [row for row in rows if str(row.get("source_path") or "") in source_paths]
    if not source_rows:
        return None
    review_terms = ("方案比选及策划评审会", "策划评审会", "施工图设计评审", "设计评审会", "设计策划集中评审会", "会议时间", "评审会")
    date_pattern = re.compile(r"20\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日?(?:\s*\d{1,2}:\d{2}(?:\s*[-—–]\s*\d{1,2}:\d{2})?)?")
    review_fragments: list[tuple[dict[str, Any], str]] = []
    dated_fragments: list[tuple[dict[str, Any], str]] = []
    for row in source_rows:
        text = str(row.get("excerpt") or "")
        fragments = [part.strip() for part in re.split(r"(?<=[。！？；])|\n", text) if part.strip()]
        for fragment in fragments:
            if any(term in fragment for term in review_terms):
                review_fragments.append((row, fragment))
            if date_pattern.search(fragment) and any(term in fragment for term in ("策划", "设计评估", "中标通知书", "前期设计文件")):
                dated_fragments.append((row, fragment))
    if not review_fragments:
        return None
    direct_dates = [date_pattern.search(fragment).group(0) for _, fragment in review_fragments if date_pattern.search(fragment)]
    if not direct_dates:
        # The legacy minutes place the event title and its "会议时间" value
        # in adjacent paragraphs. Both rows are from the same approved file.
        direct_dates = [
            date_pattern.search(str(row.get("excerpt") or "")).group(0)
            for row in source_rows
            if date_pattern.search(str(row.get("excerpt") or ""))
        ]
    claims: list[dict[str, Any]] = []
    if direct_dates:
        row, fragment = next(
            (
                item
                for item in review_fragments
                if date_pattern.search(item[1])
            ),
            next(
                (
                    (candidate, str(candidate.get("excerpt") or ""))
                    for candidate in source_rows
                    if date_pattern.search(str(candidate.get("excerpt") or ""))
                ),
                review_fragments[0],
            ),
        )
        date = date_pattern.search(fragment).group(0)
        claim_text = f"文档记载相关评审时间为：{date}。"
        claims.append({"claim_id": "C1", "claim_text": claim_text, "evidence_ids": [row["source_id"]], "support_excerpt": fragment})
        answer_text = f"结论：{claim_text} [{row['source_id']}]"
    else:
        event_row, event_fragment = review_fragments[0]
        claims.append({"claim_id": "C1", "claim_text": "现有源文件未明确记载该设计管理策划评审会的具体日期。", "evidence_ids": [event_row["source_id"]], "support_excerpt": event_fragment})
        seen_dates: set[str] = set()
        timeline_claims = []
        for row, fragment in dated_fragments:
            date_match = date_pattern.search(fragment)
            date = date_match.group(0) if date_match else ""
            if not date or date in seen_dates:
                continue
            seen_dates.add(date)
            timeline_claims.append((row, fragment))
        for index, (row, fragment) in enumerate(timeline_claims[:3], start=2):
            claims.append({"claim_id": f"C{index}", "claim_text": fragment, "evidence_ids": [row["source_id"]], "support_excerpt": fragment})
        answer_lines = [
            "结论：现有源文件未明确记载该设计管理策划评审会的具体日期。",
            "",
            "可确认的时间线：",
        ]
        answer_lines.extend(f"- {claim['claim_text']} [{claim['evidence_ids'][0]}]" for claim in claims[1:])
        answer_lines.append(f"- 文档同时记载：{event_fragment} [{event_row['source_id']}]，但该段没有给出具体日期。")
        answer_text = "\n".join(answer_lines)
    citations = [
        {
            "claim_id": claim["claim_id"],
            "evidence_id": claim["evidence_ids"][0],
            "excerpt": claim["support_excerpt"],
            "location": _claim_location(claim, rows),
        }
        for claim in claims
    ]
    return {
        "final_status": "GENERATED",
        "answer": answer_text,
        "claims": claims,
        "citations": citations,
        "provider_status": "NOT_REQUIRED",
        "generation_mode": "DETERMINISTIC_DATE_EVIDENCE_BOUNDARY",
        "date_fact_status": "EXPLICIT_DATE_NOT_FOUND" if not direct_dates else "EXPLICIT_DATE_FOUND",
    }


def deterministic_option_fallback(question: str, rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Render table-backed options when the Trial Provider is unavailable."""
    extracted = deterministic_option_answer(question, rows)
    if extracted is None:
        return None
    payload = extracted["payload"]
    evidence_bundle = rows_to_bundle(rows)
    rendered = render_option_claims(payload, evidence_bundle)
    row_by_source = {str(row.get("source_id")): row for row in rows}
    citations = [
        {
            "claim_id": claim["claim_id"],
            "evidence_id": evidence_id,
            "excerpt": row_by_source.get(evidence_id, {}).get("excerpt", ""),
            "location": row_by_source.get(evidence_id, {}).get("location", {}),
        }
        for claim in payload.get("claims", [])
        for evidence_id in claim.get("evidence_ids", [])
        if str(evidence_id) in row_by_source
    ]
    return {
        "final_status": "GENERATED",
        "answer": rendered["answer_text"],
        "claims": payload.get("claims", []),
        "citations": citations,
        "provider_status": "NOT_REQUIRED",
        "generation_mode": "DETERMINISTIC_OPTION_EVIDENCE",
        "schema_validation": extracted.get("schema_validation"),
        "citation_validation": {"valid": bool(citations), "errors": [] if citations else ["NO_VALID_CITATION"]},
    }


def deterministic_atomic_evidence_fallback(question: str) -> dict[str, Any] | None:
    """Return transparent atomic excerpts when generation cannot run."""
    global _atomic_records
    if _atomic_records is None:
        _atomic_records = load_jsonl(CONFIG.atomic_evidence_path)
    hits = search_atomic_evidence(question, _atomic_records, limit=8)
    strong_hits = [
        hit
        for hit in hits
        if not _atomic_record_is_registration(hit["record"])
        and _atomic_record_has_body(hit["record"])
        and (hit.get("phrase_matches") or hit.get("text_metric_matches") or hit.get("text_entity_matches"))
    ]
    if not strong_hits:
        return None
    evidence_rows: list[dict[str, Any]] = []
    claims: list[dict[str, Any]] = []
    for index, hit in enumerate(strong_hits[:3], start=1):
        record = hit["record"]
        source_id = f"A{index}"
        metadata = record.get("metadata") or {}
        evidence_rows.append(
            {
                "source_id": source_id,
                "knowledge_root_id": "Root-001",
                "document_id": record.get("document_id"),
                "chunk_id": record.get("evidence_id"),
                "file_name": record.get("file_name"),
                "source_path": record.get("source_path"),
                "document_role": metadata.get("document_role") or "未分类",
                "authority_level": metadata.get("authority_level") or "UNKNOWN",
                "location": record.get("location") or {},
                "excerpt": record.get("text") or "",
                "candidate_origin": ["ATOMIC_EVIDENCE_SHADOW"],
                "compatibility": "ATOMIC_EVIDENCE",
                "evidence_capability": ["ATOMIC_EVIDENCE"],
            }
        )
        claims.append(
            {
                "claim_id": f"C{index}",
                "claim_type": "EVIDENCE_EXCERPT",
                "claim_text": str(record.get("text") or "").strip(),
                "evidence_ids": [source_id],
            }
        )
    lines = [
        "当前生成模型未启用，以下仅返回与问题直接匹配的原文证据摘录，不对证据进行事实扩展：",
        "",
    ]
    lines.extend(f"- {claim['claim_text']} [{claim['evidence_ids'][0]}]" for claim in claims)
    return {
        "final_status": "EVIDENCE_ONLY",
        "answer": "\n".join(lines),
        "claims": claims,
        "citations": [
            {
                "claim_id": claim["claim_id"],
                "evidence_id": claim["evidence_ids"][0],
                "excerpt": row["excerpt"],
                "location": row["location"],
            }
            for claim, row in zip(claims, evidence_rows, strict=True)
        ],
        "provider_status": "NOT_REQUIRED",
        "generation_mode": "ATOMIC_EVIDENCE_ONLY",
        "evidence_rows": evidence_rows,
    }


def _atomic_record_is_registration(record: dict[str, Any]) -> bool:
    if record.get("file_type") != ".md":
        return False
    text = str(record.get("text") or "")
    path = str(record.get("source_path") or "").lower()
    return any(marker in text for marker in ("file://", "[[", "原始资料", "原库业务目录")) or (
        "关联内容" in text and not any(char.isdigit() for char in text)
    ) or "\\wiki\\entities\\" in path


def _atomic_record_has_body(record: dict[str, Any]) -> bool:
    text = str(record.get("text") or "").strip()
    if not text or text.startswith("#"):
        return False
    return len(text) >= 12 or any(mark in text for mark in ("：", ":", "。", "；", "|") )


def deterministic_workbook_fact_answer(question: str, rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Render a single approved workbook value without LLM calculation."""
    if "工期" not in question:
        return None
    row = next((item for item in rows if item.get("structured_fact")), None)
    if row is None:
        return None
    fact = row["structured_fact"]
    claim_text = f"{fact['project_name']}的{fact['field']}为{fact['value']}{fact['unit']}。"
    return {
        "final_status": "GENERATED",
        "answer": f"结论：{claim_text} [{row['source_id']}]",
        "claims": [
            {
                "claim_id": "C1",
                "claim_type": "DIRECT_FACT",
                "claim_text": claim_text,
                "evidence_ids": [row["source_id"]],
            }
        ],
        "citations": [
            {
                "claim_id": "C1",
                "evidence_id": row["source_id"],
                "excerpt": row.get("excerpt", ""),
                "location": row.get("location", {}),
            }
        ],
        "provider_status": "NOT_REQUIRED",
        "generation_mode": "DETERMINISTIC_WORKBOOK_FACT",
        "fact_source": row,
    }


def _claim_location(claim: dict[str, Any], rows: list[dict[str, Any]]) -> dict[str, Any]:
    support = str(claim.get("support_excerpt") or "")
    for phrase, paragraph in _approved_paragraph_locations.items():
        if phrase in support:
            return {"paragraph_start": paragraph, "paragraph_end": paragraph}
    return next(row.get("location") or {} for row in rows if row.get("source_id") == claim["evidence_ids"][0])


def state_for(result: dict[str, Any], *, source_body_missing: bool = False) -> tuple[str, str]:
    final_status = str(result.get("final_status") or "")
    provider_status = str(result.get("answer", {}).get("provider_status") or "")
    if source_body_missing:
        return "SAFE_REFUSAL", "已找到知识登记页，但目标资料正文尚未进入当前试用范围，暂不能据此生成答案。"
    if final_status in {"PROVIDER_TEMPORARY_FAILURE", "PROVIDER_PERMANENT_FAILURE", "PROVIDER_CIRCUIT_OPEN", "PROVIDER_TEST_BUDGET_EXHAUSTED", "LLM_ERROR"} or provider_status in {"PROVIDER_TEMPORARY_FAILURE", "PROVIDER_PERMANENT_FAILURE"}:
        return "GENERATION_SERVICE_UNAVAILABLE", "已找到相关资料，但当前生成服务暂不可用。"
    if final_status in {"GENERATED", "FACT_RESULT"}:
        return "NORMAL_ANSWER", "已生成答案"
    if final_status == "EVIDENCE_ONLY":
        return "NORMAL_ANSWER", "已返回可核查原文证据"
    return "SAFE_REFUSAL", "当前知识库证据不足"


def append_jsonl(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(value, ensure_ascii=False) + "\n")


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    rows: list[dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except ValueError:
            continue
        if isinstance(value, dict):
            rows.append(value)
    return rows


def feedback_summary() -> dict[str, Any]:
    audits = load_jsonl(CONFIG.audit_path)
    feedback_rows = load_jsonl(CONFIG.feedback_path / "feedback.jsonl")
    return {
        "generated_at": utc_now(),
        "questions": [
            {
                "question": row.get("question"),
                "final_status": row.get("final_status"),
                "answer_state": row.get("answer_state"),
                "pipeline_run_id": row.get("pipeline_run_id"),
            }
            for row in audits
        ],
        "final_statuses": {status: sum(row.get("final_status") == status for row in audits) for status in sorted({row.get("final_status") for row in audits})},
        "user_ratings": {rating: sum(row.get("answer_quality") == rating for row in feedback_rows) for rating in sorted({row.get("answer_quality") for row in feedback_rows})},
        "citation_ratings": {rating: sum(row.get("citation_quality") == rating for row in feedback_rows) for rating in sorted({row.get("citation_quality") for row in feedback_rows})},
        "missing_sources": [row for row in feedback_rows if row.get("missing_source")],
        "failed_questions": [
            {"question": row.get("question"), "final_status": row.get("final_status"), "answer_state": row.get("answer_state")}
            for row in audits
            if row.get("answer_state") != "NORMAL_ANSWER"
        ],
    }


def audit_record(result: dict[str, Any], trial_user: str, answer_state: str) -> dict[str, Any]:
    answer = result.get("answer") or {}
    return {
        "timestamp": utc_now(),
        "trial_user": trial_user,
        "pipeline_run_id": result.get("pipeline_run_id"),
        "question": result.get("question"),
        "route": result.get("route"),
        "retrieval_status": "RETRIEVAL_COMPLETE" if result.get("retrieval", {}).get("rrf_count", 0) else "RETRIEVAL_EMPTY",
        "selected_evidence": [evidence_card(row) for row in result.get("selected_evidence", [])],
        "preflight_status": result.get("claim_preflight", {}).get("claim_preflight_status"),
        "provider_status": answer.get("provider_status"),
        "final_status": result.get("final_status"),
        "answer_state": answer_state,
        "citation": answer.get("citations", []),
        "feedback": None,
    }


@app.get("/", response_class=FileResponse)
def home() -> FileResponse:
    return FileResponse(str(KNOWLEDGE_UI_DIST / "index.html"), headers={"Cache-Control": "no-store, max-age=0"})


@app.get("/knowledge-os", response_class=FileResponse)
def knowledge_os() -> FileResponse:
    return FileResponse(str(KNOWLEDGE_UI_DIST / "index.html"), headers={"Cache-Control": "no-store, max-age=0"})


@app.get("/trial", response_class=FileResponse)
def trial_home() -> str:
    return str(Path(__file__).parent / "static" / "index.html")


@app.get("/v2-trial", response_class=FileResponse)
def v2_trial_home() -> FileResponse:
    return FileResponse(str(V2_TRIAL_STATIC), headers={"Cache-Control": "no-store, max-age=0"})


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"service": "AI设计管理知识库 V1.0 Internal Trial", "status": "ok", "readiness": current_readiness()}


@app.get("/api/readiness")
def readiness() -> dict[str, Any]:
    return current_readiness()


@app.get("/api/stats")
def stats() -> dict[str, Any]:
    audits = load_jsonl(CONFIG.audit_path)
    feedback = load_jsonl(CONFIG.feedback_path / "feedback.jsonl")
    return {
        "total_questions": len(audits),
        "normal_answers": sum(item.get("answer_state") == "NORMAL_ANSWER" for item in audits),
        "safe_refusals": sum(item.get("answer_state") == "SAFE_REFUSAL" for item in audits),
        "generation_service_unavailable": sum(item.get("answer_state") == "GENERATION_SERVICE_UNAVAILABLE" for item in audits),
        "deterministic_answers": sum(item.get("final_status") == "FACT_RESULT" for item in audits),
        "feedback_correct": sum(item.get("answer_quality") == "正确" for item in feedback),
        "feedback_partial": sum(item.get("answer_quality") == "部分正确" for item in feedback),
        "feedback_error": sum(item.get("answer_quality") == "错误" for item in feedback),
        "feedback_missing_source": sum(bool(item.get("missing_source")) for item in feedback),
        "audit_logging_enabled": CONFIG.get("audit_logging_enabled") is True,
    }


@app.get("/api/feedback-summary")
def feedback_summary_endpoint() -> dict[str, Any]:
    return feedback_summary()


@app.post("/api/query")
def query(request: QueryRequest) -> dict[str, Any]:
    if request.trial_user not in USERS:
        raise HTTPException(status_code=403, detail="试用用户不在白名单中")
    if not _readiness.get("ready"):
        raise HTTPException(status_code=503, detail="STARTUP_BLOCKED")
    result = get_pipeline().run(request.trial_user, request.question)
    selected_rows = approved_source_rows(request.question, get_pipeline(), result.get("selected_evidence", []))
    result["selected_evidence"] = selected_rows
    source_body_missing = registration_only_evidence(selected_rows)
    answer = result.get("answer") or {}
    deterministic = None if source_body_missing else deterministic_date_answer(request.question, selected_rows)
    if deterministic is None and not source_body_missing:
        deterministic = deterministic_workbook_fact_answer(request.question, selected_rows)
    provider_failed = result.get("final_status") in {
        "PROVIDER_TEMPORARY_FAILURE",
        "PROVIDER_PERMANENT_FAILURE",
        "PROVIDER_CIRCUIT_OPEN",
        "LLM_ERROR",
    }
    if deterministic is None and not source_body_missing and provider_failed:
        deterministic = deterministic_option_fallback(request.question, selected_rows)
    if deterministic is None and (
        provider_failed or result.get("final_status") in {"NO_EVIDENCE", "STRUCTURE_INVALID"}
    ):
        atomic_fallback = deterministic_atomic_evidence_fallback(request.question)
        if atomic_fallback is not None:
            selected_rows = [*selected_rows, *atomic_fallback.pop("evidence_rows", [])]
            source_body_missing = False
            deterministic = atomic_fallback
    if deterministic is not None:
        answer = deterministic
        result["answer"] = deterministic
        result["final_status"] = deterministic["final_status"]
        result["answer_path"] = deterministic["generation_mode"]
    answer_state, user_message = state_for(result, source_body_missing=source_body_missing)
    selected = [evidence_card(row) for row in selected_rows]
    fact_authority = answer.get("citations", []) if result.get("final_status") == "FACT_RESULT" else []
    response = {
        "question": request.question,
        "answer_state": answer_state,
        "user_message": user_message,
        "answer": "" if source_body_missing else answer.get("answer", ""),
        "final_status": result.get("final_status"),
        "route": result.get("route"),
        "answer_path": result.get("answer_path"),
        "pipeline_run_id": result.get("pipeline_run_id"),
        "retrieval_status": "RETRIEVAL_COMPLETE" if result.get("retrieval", {}).get("rrf_count", 0) else "RETRIEVAL_EMPTY",
        "preflight_status": result.get("claim_preflight", {}).get("claim_preflight_status"),
        "provider_status": answer.get("provider_status", "NOT_REQUIRED"),
        "trial_diagnostic": "SOURCE_BODY_MISSING" if source_body_missing else None,
        "approved_shadow_sources": _approved_source_status,
        "evidence": selected,
        "retrieval_context": selected if result.get("final_status") == "FACT_RESULT" else [],
        "fact_authority_source": fact_authority,
        "citation": answer.get("citations", []),
        "trial_notice": "当前为受控只读试用环境，部分生成式问答功能暂未开放，回答请以引用原文为准。",
    }
    append_jsonl(CONFIG.audit_path, audit_record(result, request.trial_user, answer_state))
    return response


@app.post("/api/feedback")
def feedback(request: FeedbackRequest) -> dict[str, Any]:
    if request.trial_user not in USERS:
        raise HTTPException(status_code=403, detail="试用用户不在白名单中")
    value = request.model_dump()
    value["feedback_id"] = value.get("feedback_id") or f"feedback-{uuid.uuid4().hex}"
    value["timestamp"] = value.get("timestamp") or utc_now()
    append_jsonl(CONFIG.feedback_path / "feedback.jsonl", value)
    return {"saved": True, "feedback_id": value["feedback_id"], "path": "data/trial_feedback/feedback.jsonl"}


@app.get("/{client_path:path}", response_class=FileResponse, include_in_schema=False)
def knowledge_ui_route(client_path: str) -> str:
    """Serve BrowserRouter routes without intercepting declared API or asset paths."""
    return str(KNOWLEDGE_UI_DIST / "index.html")
