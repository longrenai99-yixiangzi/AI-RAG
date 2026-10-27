import base64
import hashlib
import json
import time
import uuid
from pathlib import Path

import numpy as np
import pytest
from fastapi.testclient import TestClient

from app.answer_engine.llm.llm_provider import GenerationResult
from app.trial.platform_app import create_app
from app.trial.platform_generation import selection_packet, render_selection
from app.trial.platform_backup import create_backup, restore_backup
from app.trial.platform_store import PlatformStore
from test_platform_existing_content import ExistingFixture, HEADERS, ask, feedback, publish


class FeatureFixture(ExistingFixture):
    def features(self,texts,question=''):
        if question:
            return {'model_digest':'test-model','scores':[2.0 if any(w in t for w in ('责任人','期限','版本','橙色')) else -2.0 for t in texts]}
        vectors=[]
        for text in texts:
            v=np.zeros(1024);v[0]=.1
            for c in text:v[ord(c)%1024]+=1
            vectors.append((v/np.linalg.norm(v)).tolist())
        return {'model_digest':'test-model','vectors':vectors}

    def retrieve(self,question,request_id):
        if '《' in question:
            self.calls+=1
            return {'answer_status':'EVIDENCE_INSUFFICIENT','citations':[]}
        return super().retrieve(question,request_id)

    def source_snapshot(self,source_id):
        source=self.catalog_source(source_id)
        text=self.path.read_text(encoding='utf-8')
        return {'model_digest':'test-model','records':[{'evidence_id':'E1','source_version':source['current_hash'],'text':text,'heading_path':'第一条','location':{'line':1}}],
                'vectors':self.features([text])['vectors']}


class FakeProvider:
    count=0
    invalid=False
    def __init__(self,packet,approval,out,max_requests):
        assert max_requests==1 and json.loads(Path(approval).read_text())['approved']
        self.body=json.loads(Path(packet).read_text(encoding='utf-8'))[0]['body']
    def generate(self,system,user,**kwargs):
        FakeProvider.count+=1
        assert self.body['messages']==[{'role':'system','content':system},{'role':'user','content':user}]
        units=json.loads(user)['units']
        ids=['U999'] if self.invalid else [u['unit_id'] for u in units if not any(v.get('context_unit_id')==u['unit_id'] for v in units)]
        return GenerationResult(json.dumps({'selected_unit_ids':ids,'gaps':[]}),None,0)


@pytest.fixture
def complete(tmp_path):
    existing=FeatureFixture(tmp_path/'原文.md')
    FakeProvider.count=0;FakeProvider.invalid=False
    with TestClient(create_app(tmp_path/'state.json',existing=existing,provider_factory=FakeProvider)) as client:
        yield client,existing


def job_wait(client,job_id):
    for _ in range(300):
        job=next(j for j in client.get('/api/platform/state').json()['jobs'] if j['job_id']==job_id)
        if job['status'] not in {'QUEUED','PROCESSING'}:return job
        time.sleep(.01)
    raise AssertionError('task did not finish')


def upload(client,name,text):
    response=client.post('/api/platform/imports',headers=HEADERS,json={'file_name':name,'content_base64':base64.b64encode(text.encode()).decode()})
    assert response.status_code==200,response.text
    return job_wait(client,response.json()['job_id'])


def activate(client,job,action='activate'):
    response=client.post(f'/api/platform/imports/{job["job_id"]}/actions',headers=HEADERS,json={'action':action})
    assert response.status_code==200,response.text
    return response.json()


def test_incremental_ingest_replace_withdraw_persists_and_deduplicates(complete):
    client,existing=complete
    job=upload(client,'新增资料.md','# 橙色流程\n责任人须记录期限，版本一。')
    assert job['status']=='READY'
    assert ask(client,'《新增资料.md》责任人和期限怎么记录？')['result']['answer_status']=='EVIDENCE_INSUFFICIENT'
    activate(client,job)
    assert '版本一' in ask(client,'《新增资料.md》责任人和期限怎么记录？')['result']['answer']
    assert client.get('/api/platform/library',params={'q':'新增资料'}).json()['total']==2  # fixture base catalog returns one
    same=upload(client,'新增资料.md','# 橙色流程\n责任人须记录期限，版本一。')
    assert same['job_id']==job['job_id'] and same['attempts']==1
    newer=upload(client,'新增资料.md','# 橙色流程\n责任人须记录期限，版本二。')
    activate(client,newer)
    result=ask(client,'《新增资料.md》责任人和期限怎么记录？')['result']['answer']
    assert '版本二' in result and '版本一' not in result
    activate(client,newer,'withdraw')
    assert ask(client,'《新增资料.md》责任人和期限怎么记录？')['result']['answer_status']=='EVIDENCE_INSUFFICIENT'
    activate(client,job)
    assert '版本一' in ask(client,'《新增资料.md》责任人和期限怎么记录？')['result']['answer']


def test_existing_snapshot_reuse_never_calls_parser(complete,monkeypatch):
    client,existing=complete
    monkeypatch.setattr('app.trial.platform_lifecycle.parse_file',lambda *a:(_ for _ in ()).throw(AssertionError('must reuse')))
    response=client.post('/api/platform/imports/from-existing',headers=HEADERS,json={'source_id':'SRC_fixture'})
    assert response.status_code==200
    job=job_wait(client,response.json()['job_id'])
    assert job['status']=='READY' and job['method']=='REUSE_EXISTING_CHUNKS'
    assert hashlib.sha256(existing.path.read_bytes()).hexdigest()==existing.hash


def test_reactivating_old_material_allows_new_review_but_never_revives_old_knowledge(complete):
    client,_=complete
    text='# 橙色流程\n责任人须记录期限，橙色版本一。'
    old=upload(client,'合成来源回退.md',text);activate(client,old)
    question='《合成来源回退.md》责任人和期限怎么记录？';run=ask(client,question)
    fb,_=feedback(client,run,expected_answer='责任人须记录期限，橙色版本一。',required_terms=['责任人','期限'],applicability='仅合成测试')
    assert fb.status_code==200
    published=publish(client,fb.json()['candidate_id'])
    newer=upload(client,'合成来源回退.md',text.replace('版本一','版本二'));activate(client,newer)
    activate(client,old)
    state=client.get('/api/platform/state').json();source=next(s for s in state['sources'] if s['source_id']==old['source_id'])
    assert source['status']=='REGISTERED'
    knowledge=next(k for k in state['knowledge'] if k['knowledge_id']==published['knowledge_id'])
    assert knowledge['status']=='REVIEW_REQUIRED'
    new_run=ask(client,question);assert new_run['result'].get('generation_mode')!='HUMAN_REVIEWED_EXACT_QUESTION'
    next_feedback,_=feedback(client,new_run,expected_answer='责任人须记录期限，橙色版本一。',required_terms=['责任人','期限'])
    assert next_feedback.status_code==200


def test_bare_link_navigation_is_not_accepted_as_a_new_body(complete):
    client,_=complete
    job=upload(client,'项目机制登记.md','[机制原件](file:///D:/演示/机制.docx)\n## ## 关联内容\n- [[业务/机制]]')
    assert job['status']=='FAILED' and '仅链接登记卡' in job['error']


def test_existing_snapshot_reuse_cannot_bypass_navigation_admission(complete,monkeypatch):
    client,existing=complete
    existing.path.write_text('[原件](file:///D:/演示/机制.docx)\n## 关联内容\n- [[业务/机制]]',encoding='utf-8')
    existing.hash=hashlib.sha256(existing.path.read_bytes()).hexdigest();existing.source['current_hash']=existing.hash
    monkeypatch.setattr(existing,'source_snapshot',lambda *a:(_ for _ in ()).throw(AssertionError('must reject before snapshot')))
    response=client.post('/api/platform/imports/from-existing',headers=HEADERS,json={'source_id':'SRC_fixture'})
    assert response.status_code==200
    job=job_wait(client,response.json()['job_id'])
    assert job['status']=='FAILED' and '仅链接登记卡' in job['error']


def test_parse_failure_retry_cancel_and_path_guards(complete):
    client,_=complete
    failed=upload(client,'坏文件.docx','not a zip')
    assert failed['status']=='FAILED'
    activate(client,failed,'retry')
    failed=job_wait(client,failed['job_id'])
    assert failed['status']=='FAILED' and failed['attempts']==2
    activate(client,failed,'cancel')
    for name in ('../outside.md','CON.md','name.md.','x.exe'):
        assert client.post('/api/platform/imports',headers=HEADERS,json={'file_name':name,'content_base64':'YQ=='}).status_code==422
    assert client.post('/api/platform/imports',headers=HEADERS,json={'file_name':'x.md','content_base64':'!!!'}).status_code==422


def test_approved_knowledge_is_retrieved_for_nonidentical_question_then_withdrawn(complete):
    client,_=complete
    row=ask(client)
    fb,_=feedback(client,row)
    candidate=publish(client,fb.json()['candidate_id'])
    result=ask(client,'责任人和期限这两项如何留痕？')['result']
    assert any(c.get('knowledge_id')==candidate['knowledge_id'] for c in result['citations'])
    knowledge=client.get('/api/platform/state').json()['knowledge'][0]
    client.post(f'/api/platform/knowledge/{knowledge["knowledge_id"]}/actions',headers=HEADERS,json={'action':'withdraw','revision':knowledge['revision'],'note':'test'})
    assert not any(c.get('knowledge_id') for c in ask(client,'责任人和期限这两项如何留痕？')['result']['citations'])


def test_incremental_tamper_does_not_use_stale_chunks(complete):
    client,_=complete
    job=upload(client,'篡改.md','# 版本\n责任人和期限。')
    activate(client,job)
    Path(job['input_path']).write_text('new',encoding='utf-8')
    result=ask(client,'《篡改.md》责任人如何记录？')['result']
    assert result['answer_status']=='EVIDENCE_INSUFFICIENT'
    current=next(j for j in client.get('/api/platform/state').json()['jobs'] if j['job_id']==job['job_id'])
    assert current['status']=='FAILED'


def test_generation_previews_require_new_explicit_approval_and_never_retry(complete):
    client,_=complete
    row=ask(client)
    response=client.post('/api/platform/generation-preview',headers=HEADERS,json={'query_run_id':row['query_run_id']})
    assert response.status_code==200,response.text
    preview=response.json()
    assert FakeProvider.count==0
    path=f'/api/platform/generation/{preview["preview_id"]}/run'
    body={'packet_sha256':preview['packet_sha256'],'approved':False}
    assert client.post(path,headers=HEADERS,json=body).status_code==403
    body['approved']=True
    generated=client.post(path,headers=HEADERS,json=body)
    assert generated.status_code==200 and generated.json()['result']['answer_status']=='MODEL_SOURCE_ORGANIZED'
    assert FakeProvider.count==1
    assert client.post(path,headers=HEADERS,json=body).status_code==422
    assert FakeProvider.count==1


def test_invalid_model_ids_fail_without_free_text_or_second_call(complete):
    client,_=complete
    preview=client.post('/api/platform/generation-preview',headers=HEADERS,json={'query_run_id':ask(client)['query_run_id']}).json()
    FakeProvider.invalid=True
    result=client.post(f'/api/platform/generation/{preview["preview_id"]}/run',headers=HEADERS,json={'packet_sha256':preview['packet_sha256'],'approved':True}).json()
    assert result['status']=='ERROR' and result['result']['answer_status']=='MODEL_VALIDATION_FAILED'
    assert FakeProvider.count==1


def test_context_and_original_text_are_retained_by_unit_selection():
    cards=[{'citation_id':'S1','file_name':'f.md','excerpt':'只有审批通过后适用。第一，责任人记录期限。第二，不得改变总人数。','answer_usage':'PRIMARY'}]
    body,units=selection_packet('如何记录？',cards)
    result=render_selection(json.dumps({'selected_unit_ids':['U2'],'gaps':[]}),units,cards)
    assert '只有审批通过后适用' in result['answer'] and '第一，责任人记录期限。' in result['answer']
    for malformed in ({'selected_unit_ids':['U2','U2'],'gaps':[]},{'selected_unit_ids':['U999'],'gaps':[]},{'selected_unit_ids':['U2'],'gaps':[],'answer':'invented'}):
        with pytest.raises(ValueError):render_selection(json.dumps(malformed),units,cards)
    normalized=render_selection(json.dumps({'type':'json_object','selected_unit_ids':['U1','U2','U3'],'gaps':[]}),units,cards)
    assert normalized['format_normalization']=='REDUNDANT_JSON_OBJECT_TYPE_REMOVED'
    assert '3. ' not in normalized['answer'] and '2. 第二' in normalized['answer']
    with pytest.raises(ValueError):render_selection(json.dumps({'type':'other','selected_unit_ids':['U2'],'gaps':[]}),units,cards)


def test_backup_restore_preserves_records_and_never_resets_generation_spend(complete):
    client,_=complete
    job=upload(client,'备份.md','# 版本\n橙色责任人记录期限。')
    activate(client,job)
    run=ask(client)
    backup=client.post('/api/platform/backups',headers=HEADERS,json={}).json()
    preview=client.post('/api/platform/generation-preview',headers=HEADERS,json={'query_run_id':run['query_run_id']}).json()
    generated=client.post(f'/api/platform/generation/{preview["preview_id"]}/run',headers=HEADERS,json={'packet_sha256':preview['packet_sha256'],'approved':True}).json()
    assert FakeProvider.count==1
    activate(client,job,'withdraw')
    response=client.post(f'/api/platform/backups/{backup["backup_id"]}/restore',headers=HEADERS,json={'sha256':backup['sha256'],'confirmed':True})
    assert response.status_code==200,response.text
    assert response.json()['paid_call_ledgers_restored'] is False
    assert client.get('/api/platform/state').json()['jobs'][0]['status']=='ACTIVE'
    assert client.post(f'/api/platform/generation/{preview["preview_id"]}/run',headers=HEADERS,json={'packet_sha256':preview['packet_sha256'],'approved':True}).status_code==422
    assert FakeProvider.count==1
    assert client.get('/api/platform/runs/'+run['query_run_id']).status_code==200
    assert client.get('/api/platform/runs/'+generated['query_run_id']).status_code==200


def test_corrupt_backup_and_unconfirmed_restore_do_not_change_state(complete):
    client,_=complete
    backup=client.post('/api/platform/backups',headers=HEADERS,json={}).json()
    before=client.app.state.platform_store.path.read_bytes()
    assert client.post(f'/api/platform/backups/{backup["backup_id"]}/restore',headers=HEADERS,json={'sha256':backup['sha256']}).status_code==403
    Path(backup['path']).write_bytes(b'corrupted')
    assert client.post(f'/api/platform/backups/{backup["backup_id"]}/restore',headers=HEADERS,json={'sha256':backup['sha256'],'confirmed':True}).status_code==422
    assert client.app.state.platform_store.path.read_bytes()==before


@pytest.mark.parametrize('field',['package_path','input_path'])
def test_valid_backup_restores_missing_managed_file_and_flags_partial_safety_snapshot(complete,field):
    client,_=complete
    job=upload(client,'恢复缺失文件.md','# 原文\n责任人和期限。')
    activate(client,job)
    job=client.get('/api/platform/state').json()['jobs'][0]
    path=Path(job[field]);original=path.read_bytes()
    backup=client.post('/api/platform/backups',headers=HEADERS,json={}).json()
    assert path.resolve().is_relative_to((client.app.state.platform_store.path.parent/'runtime').resolve())
    path.unlink()
    assert client.post('/api/platform/backups',headers=HEADERS,json={}).status_code==503
    response=client.post(f'/api/platform/backups/{backup["backup_id"]}/restore',headers=HEADERS,json={'sha256':backup['sha256'],'confirmed':True})
    assert response.status_code==200,response.text
    assert response.json()['safety_missing_files'] and path.read_bytes()==original
    safety=response.json()['safety_backup_id']
    import zipfile
    with zipfile.ZipFile(client.app.state.platform_store.path.parent/'backups'/(safety+'.zip')) as archive:
        saved=json.loads(archive.read('state.json'))
    assert saved['jobs'][job['job_id']]['status']=='FAILED'
    assert '缺失' in saved['jobs'][job['job_id']]['error']


def test_restart_marks_inflight_operations_interrupted(tmp_path):
    store=PlatformStore(tmp_path/'state.json')
    row,_=store.begin_run('a'*32,'未完成的问题','fixture')
    store._mutate(lambda s:s['generations'].__setitem__('g',{'status':'EXECUTING'}))
    with TestClient(create_app(store.path,existing=FeatureFixture(tmp_path/'source.md'))) as client:
        state=client.get('/api/platform/state').json()
        assert state['generations'][0]['status']=='INTERRUPTED_SPENT'
        assert client.get('/api/platform/runs/'+row['query_run_id']).json()['status']=='INTERRUPTED'
        result=client.post('/api/platform/query',headers=HEADERS,json={'request_id':'a'*32,'question':'未完成的问题'}).json()
        assert result['status']=='INTERRUPTED'


def test_restore_is_blocked_while_model_call_is_inflight(complete):
    client,_=complete
    backup=client.post('/api/platform/backups',headers=HEADERS,json={}).json()
    store=client.app.state.platform_store
    store._mutate(lambda s:s['generations'].__setitem__('g',{'status':'EXECUTING'}))
    assert client.post(f'/api/platform/backups/{backup["backup_id"]}/restore',headers=HEADERS,json={'sha256':backup['sha256'],'confirmed':True}).status_code==409


def test_pricing_update_requires_confirmation_and_invalidates_old_preview(complete):
    client,_=complete
    preview=client.post('/api/platform/generation-preview',headers=HEADERS,json={'query_run_id':ask(client)['query_run_id']}).json()
    prices={'input_usd_per_million':.4,'output_usd_per_million':2,'fx_ceiling':10}
    assert client.post('/api/platform/generation-pricing',headers=HEADERS,json=prices).status_code==403
    assert client.post('/api/platform/generation-pricing',headers=HEADERS,json={**prices,'confirmed':True}).status_code==200
    assert client.post(f'/api/platform/generation/{preview["preview_id"]}/run',headers=HEADERS,json={'packet_sha256':preview['packet_sha256'],'approved':True}).status_code==422
    assert FakeProvider.count==0


def test_invalid_zip_with_matching_external_digest_is_rejected(complete):
    client,_=complete
    backup=client.post('/api/platform/backups',headers=HEADERS,json={}).json()
    Path(backup['path']).write_bytes(b'not a valid zip')
    digest=hashlib.sha256(Path(backup['path']).read_bytes()).hexdigest()
    before=client.app.state.platform_store.path.read_bytes()
    result=client.post(f'/api/platform/backups/{backup["backup_id"]}/restore',headers=HEADERS,json={'sha256':digest,'confirmed':True})
    assert result.status_code==422 and client.app.state.platform_store.path.read_bytes()==before


@pytest.mark.parametrize('extension',['.docx','.xlsx','.pptx','.pdf'])
def test_native_file_formats_enter_ready_queue(complete,tmp_path,extension):
    client,_=complete
    path=tmp_path/('原生样本'+extension)
    if extension=='.docx':
        from docx import Document
        document=Document();document.add_heading('橙色流程',1);document.add_paragraph('责任人应记录期限。');document.save(path)
    elif extension=='.xlsx':
        from openpyxl import Workbook
        document=Workbook();document.active.append(['责任人','期限']);document.active.append(['橙色负责人','明天']);document.save(path)
    elif extension=='.pptx':
        from pptx import Presentation
        from pptx.util import Inches
        document=Presentation();slide=document.slides.add_slide(document.slide_layouts[5]);slide.shapes.title.text='橙色流程'
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(2)).text='责任人应记录期限。';document.save(path)
    else:
        import pymupdf
        document=pymupdf.open();page=document.new_page();page.insert_text((72,72),'Orange owner records deadline.');document.save(path);document.close()
    response=client.post('/api/platform/imports',headers=HEADERS,json={'file_name':path.name,'content_base64':base64.b64encode(path.read_bytes()).decode()})
    assert response.status_code==200
    job=job_wait(client,response.json()['job_id'])
    assert job['status']=='READY' and job['chunk_count']>0,job


def test_pdf_without_readable_text_does_not_claim_success(complete,tmp_path):
    import pymupdf
    client,_=complete
    path=tmp_path/'无正文.pdf'
    document=pymupdf.open();document.new_page();document.save(path);document.close()
    response=client.post('/api/platform/imports',headers=HEADERS,json={'file_name':path.name,'content_base64':base64.b64encode(path.read_bytes()).decode()})
    job=job_wait(client,response.json()['job_id'])
    assert job['status']=='FAILED' and 'OCR' in job['error']
