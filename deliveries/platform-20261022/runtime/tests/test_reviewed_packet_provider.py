import hashlib,json
import pytest
from app.answer_engine.llm.reviewed_packet_provider import ReviewedPacketProvider,wire,sha
from app.llm import LLMError


def packet(tmp_path,count=2):
    tasks=[]
    for i in range(count):
        body={'model':'fixed-model','messages':[{'role':'system','content':'schema'},{'role':'user','content':str(i)}],'temperature':0.1,'max_tokens':4096,'thinking':{'type':'disabled'}}
        tasks.append({'id':f'T{i}','body':body,'request_sha256':hashlib.sha256(wire(body).encode()).hexdigest()})
    p=tmp_path/'packet.json';p.write_text(json.dumps(tasks),encoding='utf-8');a=tmp_path/'approval.json'
    return p,a


def approve(p,a,count=2):a.write_text(json.dumps({'approved':True,'packet_sha256':sha(p),'max_requests':count,'max_budget_cny':1}),encoding='utf-8')


def test_no_approval_never_invokes_transport(tmp_path):
    p,a=packet(tmp_path);provider=ReviewedPacketProvider(p,a,tmp_path/'run',transport=lambda _:pytest.fail('must not call'))
    assert not provider.available
    with pytest.raises(LLMError):provider.generate('schema','0')


def test_exact_requests_once_across_restart_and_no_repair_calls(tmp_path):
    p,a=packet(tmp_path);approve(p,a);calls=[]
    def post(body):calls.append(body);return {'id':'test','choices':[{'message':{'content':'{}'},'finish_reason':'stop'}]}
    provider=ReviewedPacketProvider(p,a,tmp_path/'run',transport=post)
    provider.generate('schema','0')
    with pytest.raises(LLMError):provider.generate('schema','0')
    with pytest.raises(LLMError):provider.generate('repair','new')
    provider=ReviewedPacketProvider(p,a,tmp_path/'run',transport=post)
    with pytest.raises(LLMError):provider.generate('schema','0')
    provider.generate('schema','1')
    assert len(calls)==2 and provider.budget.snapshot()['used_real_requests']==2


def test_failed_request_consumes_budget_and_cannot_be_automatically_retried(tmp_path):
    p,a=packet(tmp_path);approve(p,a);calls=[]
    def post(body):calls.append(body);raise RuntimeError('offline simulated transport failure')
    provider=ReviewedPacketProvider(p,a,tmp_path/'run',transport=post)
    for _ in range(3):
        with pytest.raises(LLMError):provider.generate('schema','0')
    assert len(calls)==1 and provider.budget.snapshot()['used_real_requests']==1


def test_reviewed_batch_preflight_and_tampered_price_are_offline(tmp_path):
    from evaluation.unified_colleague_flow_20261019.run_reviewed_batch import preflight, retrieval
    from datetime import datetime, timezone
    from app.answer_engine.answer_policy import policy_for_intent
    from app.answer_engine.llm.prompt_builder import build_prompt
    from app.trial.assisted_chat import assemble
    source = tmp_path/'synthetic.md'
    source.write_text('合成演示规定：保留会议纪要。', encoding='utf-8')
    tasks = []
    for index in range(2):
        task = {'id':f'S{index}', 'question':'合成演示要求保留什么？', 'private_binding':[{'source_id':'demo', 'source_path':str(source), 'source_version':sha(source), 'evidence_id':'paragraph', 'excerpt':source.read_text(encoding='utf-8'), 'location':{'section':'合成演示'}}]}
        evidence,_ = assemble(retrieval(task))
        prompt = build_prompt(task['question'], policy_for_intent('GENERAL_QUERY'), evidence)
        task['body'] = {'model':'deepseek-v4-flash','messages':[{'role':'system','content':prompt.system_prompt},{'role':'user','content':prompt.user_prompt}], 'temperature':0.1,'max_tokens':4096,'response_format':{'type':'json_object'},'thinking':{'type':'disabled'}}
        task['request_sha256'] = hashlib.sha256(wire(task['body']).encode()).hexdigest()
        tasks.append(task)
    packet_path = tmp_path/'synthetic-packet.json'
    packet_path.write_text(json.dumps(tasks),encoding='utf-8')
    prices = {'request_packet_sha256':sha(packet_path),'verified_at_utc':datetime.now(timezone.utc).isoformat(),'peak_cny_per_million':{'input_miss':2,'output':8},'max_batch_budget_cny':1}
    fresh = tmp_path/'price.json'
    fresh.write_text(json.dumps(prices),encoding='utf-8')
    _,receipt = preflight(packet_path,fresh)
    assert receipt['new_real_calls']==0 and receipt['estimated_max_cny']<1
    prices['request_packet_sha256']='0'*64
    fresh.write_text(json.dumps(prices),encoding='utf-8')
    with pytest.raises(AssertionError,match='different packet'):
        preflight(packet_path,fresh)


def test_word_preflight_checks_native_quote_and_detects_source_change(tmp_path):
    from datetime import datetime,timezone
    from docx import Document
    from app.answer_engine.answer_policy import policy_for_intent
    from app.answer_engine.llm.prompt_builder import build_prompt
    from app.trial.assisted_chat import assemble
    from evaluation.unified_colleague_flow_20261019.run_reviewed_batch import preflight,retrieval
    path=tmp_path/'项目原件.docx';doc=Document();doc.add_paragraph('本项目配置设计负责人1人。');doc.save(path)
    task={'id':'W01','question':'配置多少人？','private_binding':[{'source_id':'doc','source_path':str(path),'source_version':sha(path),'evidence_id':'paragraph','excerpt':'本项目配置设计负责人1人。','location':{'section':'组织'}}]}
    evidence,_=assemble(retrieval(task));prompt=build_prompt(task['question'],policy_for_intent('GENERAL_QUERY'),evidence)
    task['body']={'model':'deepseek-v4-flash','messages':[{'role':'system','content':prompt.system_prompt},{'role':'user','content':prompt.user_prompt}],'temperature':0.1,'max_tokens':4096,'response_format':{'type':'json_object'},'thinking':{'type':'disabled'}}
    task['request_sha256']=hashlib.sha256(wire(task['body']).encode()).hexdigest()
    packet=tmp_path/'word-packet.json';packet.write_text(json.dumps([task]),encoding='utf-8')
    price=tmp_path/'price.json';price.write_text(json.dumps({'request_packet_sha256':sha(packet),'verified_at_utc':datetime.now(timezone.utc).isoformat(),'peak_cny_per_million':{'input_miss':2,'output':8},'max_batch_budget_cny':1}),encoding='utf-8')
    assert preflight(packet,price,expected_count=1)[1]['new_real_calls']==0
    doc.add_paragraph('原件已变更。');doc.save(path)
    with pytest.raises(AssertionError,match='source missing or source version changed'):preflight(packet,price,expected_count=1)


def test_new_four_call_batch_requires_its_exact_count_and_separate_approval(tmp_path):
    p,a=packet(tmp_path,4);approve(p,a,4);calls=[]
    assert not ReviewedPacketProvider(p,a,tmp_path/'wrong-count').available
    provider=ReviewedPacketProvider(p,a,tmp_path/'four',max_requests=4,transport=lambda body:calls.append(body) or {'choices':[{'message':{'content':'{}'},'finish_reason':'stop'}]})
    for i in range(4):provider.generate('schema',str(i))
    with pytest.raises(LLMError):provider.generate('schema','0')
    assert len(calls)==4 and provider.budget.snapshot()['remaining_real_requests']==0
