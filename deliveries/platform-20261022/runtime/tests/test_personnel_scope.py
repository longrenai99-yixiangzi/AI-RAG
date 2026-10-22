import json,hashlib
from pathlib import Path
import pytest
from docx import Document
from app.document_intelligence.v2 import DocumentIntelligenceV2Builder
from app.retrieval.personnel_scope import build_records,PersonnelIndex,request_scope
from app.retrieval.named_source_scope import path_key


def native_index(tmp_path):
    path=tmp_path/'原件.docx';doc=Document();doc.add_heading('2.2 组织架构',level=1)
    doc.add_paragraph('本项目配置技术总监7人，负责统筹管理技术和设计工作。')
    table=doc.add_table(rows=1,cols=4)
    for c,t in zip(table.rows[0].cells,['序号','分项','具体内容','数量']):c.text=t
    for values in [['1','项目组织结构模式','A模式',''],['2','项目领导班子成员','项目经理',''],['3','部门','工程部',''],['合计','合计','合计','52人']]:
        for c,t in zip(table.add_row().cells,values):c.text=t
    table=doc.add_table(rows=1,cols=5)
    for c,t in zip(table.rows[0].cells,['序号','阶段','岗位','人员配置数量','专职/兼职']):c.text=t
    for values in [['1','施工图设计阶段','设计支持中心驻场人员','8','4人专职4人兼职'],['2','专项设计阶段','设计支持中心驻场人员','3','专职'],['合计','合计','合计','11','跨阶段原表合计，人员可能复用']]:
        for c,t in zip(table.add_row().cells,values):c.text=t
    doc.save(path);parsed=DocumentIntelligenceV2Builder(tmp_path).build([path]);records=build_records(parsed)
    pack=tmp_path/'index';pack.mkdir();file=pack/'records.json';file.write_text(json.dumps(records,ensure_ascii=False),encoding='utf-8')
    digest=hashlib.sha256(path.read_bytes()).hexdigest();(pack/'manifest.json').write_text(json.dumps({'base_candidate_hash':'base','records_sha256':hashlib.sha256(file.read_bytes()).hexdigest()}),encoding='utf-8')
    return PersonnelIndex(pack,'base',{(path_key(path),digest):'canonical-source'}),path,pack


@pytest.mark.parametrize('q,value,table,row',[
    ('谁统筹技术和设计，配置多少人？','7人',None,None),
    ('组织架构表中谁统筹技术和设计，配置多少人？','7人',None,None),
    ('组织架构表中项目管理人员合计多少人？','52人',1,5),
    ('施工图设计阶段设计支持中心驻场人员配置多少人？','8',2,2),
    ('专项设计阶段驻场人员配置多少人？','3',2,3),
])
def test_native_role_total_and_phase_are_not_substituted(tmp_path,q,value,table,row):
    index,path,_=native_index(tmp_path);r=index.lookup(q,[path_key(path)]);assert r['bundle_status']=='PRIMARY_SECTION_CANDIDATE'
    c=r['candidate_evidence'][0];assert value in c['text'] and c['source_id']=='canonical-source'
    if table:assert c['location']['table']==table and c['location']['row']==row


@pytest.mark.parametrize('q',[
    '仅设计岗位的设计管理团队总人数是多少？',
    '整个团队总人数是多少？',
    '配置多少人？',
    '施工图设计阶段专职人数合计多少人？',
    '初步设计阶段驻场人员配置多少人？',
    '设计总监配置多少人？',
    '组织架构表中领导班子合计多少人？',
    '负责人是谁以及团队总人数是多少？',
])
def test_unproven_totals_subgroups_and_other_roles_stop(tmp_path,q):
    index,path,_=native_index(tmp_path);r=index.lookup(q,[path_key(path)])
    assert r['bundle_status']=='SOURCE_SCOPE_MISSING' and all(c['answer_usage']=='BACKGROUND' for c in r['candidate_evidence'])


def test_scope_requires_one_current_source_and_tampered_index_stops(tmp_path):
    index,path,pack=native_index(tmp_path);q='谁统筹，配置多少人？'
    assert index.lookup(q,None)['bundle_status']=='SOURCE_SCOPE_MISSING'
    assert index.lookup(q,[path_key(path),'another'])['bundle_status']=='SOURCE_SCOPE_MISSING'
    path.write_bytes(b'new version');assert index.lookup(q,[path_key(path)])['bundle_status']=='SOURCE_SCOPE_MISSING'
    with pytest.raises(ValueError,match='another candidate'):PersonnelIndex(pack,'wrong',{})
    (pack/'records.json').write_text('[]')
    with pytest.raises(ValueError,match='evidence changed'):PersonnelIndex(pack,'base',{})


def test_non_count_query_keeps_the_existing_path():
    assert request_scope('设计评估有几个方面？') is None
    assert request_scope('专项设计管理有哪些要求？') is None


@pytest.mark.parametrize('q',[
    '施工图设计阶段商务总监配置多少人？',
    '施工图设计阶段机电驻场人员配置多少人？',
    '施工图设计阶段安装驻场人员配置多少人？',
    '2025年技术总监配置多少人？',
    '机电技术总监配置多少人？',
    '机电负责人配置多少人？',
    '施工图技术总监配置多少人？',
    '专项设计阶段驻场专职人员配置多少人？',
    '施工图设计阶段驻场人员中兼职人员配置多少人？',
    '项目组织架构表总人数中设计岗位有多少人？',
    '项目组织架构表项目管理人员一共多少人，不含技术部？',
    '施工图设计阶段设计支持中心和安装事业部合计多少人？',
    '专项设计阶段驻场人员配置多少人，其中兼职多少人？',
])
def test_unproven_restrictions_cannot_be_dropped(tmp_path,q):
    index,path,_=native_index(tmp_path);result=index.lookup(q,[path_key(path)])
    assert result['bundle_status']=='SOURCE_SCOPE_MISSING'
    assert all(row['answer_usage']=='BACKGROUND' for row in result['candidate_evidence'])


@pytest.mark.parametrize('q,value',[
    ('技术总监配置几位？','7人'),
    ('谁统筹技术和设计，需要几个人？','7人'),
    ('统筹技术和设计的岗位及人数是多少？','7人'),
    ('统筹技术和设计的岗位和人数是多少？','7人'),
    ('统筹技术和设计的岗位与人数是多少？','7人'),
    ('项目管理人员一共几名？','52人'),
    ('专项阶段驻场人员需要几人？',' | 3 | 专职'),
    ('施工图设计阶段驻场人员配置多少人，专兼职如何？','4人专职4人兼职'),
])
def test_supported_count_wording_keeps_native_range(tmp_path,q,value):
    index,path,_=native_index(tmp_path);result=index.lookup(q,[path_key(path)])
    assert result['bundle_status']=='PRIMARY_SECTION_CANDIDATE' and value in result['candidate_evidence'][0]['text']


@pytest.mark.parametrize('available',[False,True])
@pytest.mark.parametrize('q,value',[
    ('谁统筹技术和设计，配置多少人？','技术总监7人'),
    ('项目管理人员合计多少人？','52人'),
    ('施工图设计阶段驻场人员配置多少人？','4人专职4人兼职'),
    ('专项设计阶段驻场人员配置多少人？','专项设计阶段 | 设计支持中心驻场人员 | 3 | 专职'),
])
def test_scoped_answers_quote_native_units_without_model(tmp_path,q,value,available):
    from app.trial.assisted_chat import AssistedQuery,run_flow
    from app.answer_engine.llm.answer_generator import ShadowAnswerGenerator
    class NeverProvider:
        def generate(self,*a,**kw):raise AssertionError('Personnel extraction must not call model')
    provider=NeverProvider();provider.available=available
    index,path,_=native_index(tmp_path);bundle=index.lookup(q,[path_key(path)]);saved={}
    result=list(run_flow(AssistedQuery(question=q),lambda _:bundle,ShadowAnswerGenerator(provider),lambda k,v:saved.update({k:v})))[-2]['result']
    assert result['answer_status']=='SOURCE_EXTRACT_DRAFT' and value in result['answer']
    assert result['claims'][0]['claim_text']==result['citations'][0]['excerpt']==bundle['candidate_evidence'][0]['text']
    assert result['claims'][0]['claim_type']=='SOURCE_EXTRACT' and result['support_check']=='EXACT_NATIVE_UNIT'
    assert saved[result['query_run_id']]['result']==result


@pytest.mark.parametrize('change',['text','row','header','cells','version','kind','target','eligible','duplicate','stale','mode','gap'])
def test_invalid_or_source_only_personnel_never_generates(tmp_path,change):
    from app.trial.assisted_chat import AssistedQuery,run_flow
    from app.answer_engine.llm.answer_generator import ShadowAnswerGenerator
    index,path,_=native_index(tmp_path);q='专项设计阶段驻场人员配置多少人？'
    if change=='gap':q='仅设计岗位的设计管理团队总人数是多少？'
    bundle=index.lookup(q,[path_key(path)]);row=bundle['candidate_evidence'][0]
    if change=='text':row['text']=row['text'].replace(' | 3 | ',' | 99 | ')
    if change=='row':row['location']['row']=2
    if change=='header':row['location']['header_row']=2
    if change=='cells':row['table_cells'][3]='99'
    if change=='version':row['source_version']='0'*64
    if change=='kind':row['personnel']['kind']='LEAD_ROLE'
    if change=='target':bundle['personnel_scope_audit']['target']='LEAD_ROLE'
    if change=='eligible':bundle['personnel_scope_audit']['eligible_rows']=2
    if change=='duplicate':bundle['candidate_evidence'].append(dict(row))
    if change=='stale':path.write_bytes(b'changed')
    class NeverProvider:
        available=True
        def generate(self,*a,**kw):raise AssertionError('Personnel extraction must not call model')
    # A provider exception would be caught by run_flow; assert the result as well.
    result=list(run_flow(AssistedQuery(question=q,mode='sources' if change=='mode' else 'answer'),lambda _:bundle,ShadowAnswerGenerator(NeverProvider()),lambda *a:None))[-2]['result']
    assert result['answer_status']==('SOURCES_ONLY' if change=='mode' else 'EVIDENCE_INSUFFICIENT')
    assert not result['claims']


@pytest.mark.parametrize('q,kind',[
    ('技术总监配置几位？','LEAD_ROLE_COUNT'),
    ('项目管理人员一共几名？','PROJECT_ORGANIZATION_COUNT'),
    ('专项设计阶段驻场人员配置多少人？','PHASE_POSITION_COUNT'),
    ('施工图阶段设计总监岗位人数是多少？','PHASE_POSITION_COUNT'),
    ('施工图阶段驻场人员的总人数是多少？','PHASE_POSITION_COUNT'),
    ('项目组织架构表项目管理人员合计多少人？','PROJECT_ORGANIZATION_COUNT'),
])
def test_explicit_count_scope_is_not_displayed_as_unknown(q,kind):
    from app.answer_engine.question_scope import question_scope
    assert question_scope(q)['kind']==kind and '尚不明确' not in question_scope(q)['label']


@pytest.mark.parametrize('change',[None,'number','row','header'])
def test_personnel_table_preflight_requires_exact_native_cells(tmp_path,change):
    from datetime import datetime,timezone
    from app.trial.assisted_chat import assemble
    from app.answer_engine.answer_policy import policy_for_intent
    from app.answer_engine.llm.prompt_builder import build_prompt_v2
    from app.answer_engine.llm.reviewed_packet_provider import wire,sha
    from evaluation.unified_colleague_flow_20261019.run_reviewed_batch import preflight,retrieval
    index,path,_=native_index(tmp_path);bundle=index.lookup('项目组织架构表项目管理人员合计多少人？',[path_key(path)]);_,cards=assemble(bundle)
    if change=='number':cards[0]['excerpt']=cards[0]['excerpt'].replace('52人','99人')
    if change=='row':cards[0]['location']['row']=4
    if change=='header':cards[0]['location']['header_row']=2
    task={'id':'T1','question':'项目组织架构表项目管理人员合计多少人？','private_binding':cards}
    e,_=assemble(retrieval(task));prompt=build_prompt_v2(task['question'],policy_for_intent('GENERAL_QUERY'),e)
    task['body']={'model':'deepseek-v4-flash','messages':[{'role':'system','content':prompt.system_prompt},{'role':'user','content':prompt.user_prompt}],'temperature':0.1,'max_tokens':4096,'response_format':{'type':'json_object'},'thinking':{'type':'disabled'}}
    task['request_sha256']=hashlib.sha256(wire(task['body']).encode()).hexdigest()
    packet=tmp_path/'packet.json';packet.write_text(json.dumps([task]),encoding='utf-8');price=tmp_path/'price.json';price.write_text(json.dumps({'request_packet_sha256':sha(packet),'verified_at_utc':datetime.now(timezone.utc).isoformat(),'peak_cny_per_million':{'input_miss':2,'output':8},'max_batch_budget_cny':1}),encoding='utf-8')
    if change:
        with pytest.raises(AssertionError,match='native table'):preflight(packet,price,expected_count=1,protocol_version='V2')
    else:assert preflight(packet,price,expected_count=1,protocol_version='V2')[1]['new_real_calls']==0
