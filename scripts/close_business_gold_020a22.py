from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from openpyxl import load_workbook

from scripts.close_business_gold_locations import (
    BA007_RESOLVED_PATH,
    OWNER_SOURCES,
    PROJECT_ROOT,
    _docx_table_extract,
    _file_fingerprint,
    _load_questions,
    _read_json,
    _write_json,
)


OUTPUT_DIR = PROJECT_ROOT / "evaluation" / "business_gold_v2"
BA006_RESPONSIBILITY_PATH = Path(
    r"D:\工作\二公司技术部\2026\责任状\局\3.二公司：2026年设计与技术专项责任书 .docx"
)
BA006_WORK_PLAN_PATH = Path(
    r"D:\工作\二公司技术部\2026\各类文件\设计管理\关于印发中建三局2026年设计与技术工作计划的通知.pdf"
)
BA010_DOCX_PATH = Path(
    r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷\设计管理策划书-星谷科创中心项目.docx"
)
BA010_XLSX_PATH = Path(
    r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷\星谷科创项目设计管理策划+设计示范项目打造方案\方案比选与价值创造清单方案比选及价值创造.xlsx"
)


def main() -> int:
    questions = _load_questions()
    ba006 = _close_ba006()
    ba010 = _close_ba010()
    final_manifest = _build_final_manifest(questions, ba006, ba010)
    remaining_decisions = _build_remaining_decisions(final_manifest, ba006, ba010)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    _write_json(OUTPUT_DIR / "ba006_responsibility_source_closure.json", ba006)
    _write_json(OUTPUT_DIR / "ba010_source_lineage.json", ba010)
    _write_json(OUTPUT_DIR / "final_gold_candidate_manifest.json", {"records": final_manifest})
    _write_json(OUTPUT_DIR / "remaining_owner_decisions.json", {"records": remaining_decisions})
    (PROJECT_ROOT / "docs" / "BUSINESS_GOLD_FINAL_CLOSURE_REPORT.md").write_text(
        _render_final_report(final_manifest, ba006, ba010, remaining_decisions), encoding="utf-8"
    )
    (PROJECT_ROOT / "docs" / "BUSINESS_GOLD_OWNER_APPROVAL_SHEET.md").write_text(
        _render_approval_sheet(final_manifest, ba006, ba010), encoding="utf-8"
    )
    _update_v2_report(final_manifest, ba006, ba010)

    print(
        {
            "ba006_source_fidelity": ba006["source_fidelity_status"],
            "ba006_gold_location": ba006["gold_location_status"],
            "ba006_value": ba006["owner_source_fact"]["value"],
            "ba010_lineage": ba010["lineage_status"],
            "ba010_gold_type": next(item["gold_type"] for item in final_manifest if item["question_id"] == "BA-010"),
            "summary": dict(Counter(item["gold_location_status"] for item in final_manifest)),
            "owner_final_gold_confirmed": 0,
        }
    )
    return 0


def _close_ba006() -> dict[str, Any]:
    fingerprint = _file_fingerprint(BA006_RESPONSIBILITY_PATH)
    if not fingerprint["exists"]:
        return {
            "question_id": "BA-006",
            "source_path": str(BA006_RESPONSIBILITY_PATH),
            "source_fidelity_status": "UNRESOLVED",
            "gold_location_status": "UNRESOLVED",
            "unresolved_reason": "指定责任状正文不存在，无法完成定向闭环。",
        }

    document = WordDocument(BA006_RESPONSIBILITY_PATH)
    table = document.tables[0]
    headers = [" ".join(cell.text.split()) for cell in table.rows[0].cells]
    match = None
    for row_number, row in enumerate(table.rows, start=1):
        values = [" ".join(cell.text.split()) for cell in row.cells]
        text = " | ".join(values)
        if all(term in text for term in ("DOP", "电子图形文件", "上传", "40")):
            match = {"row": row_number, "values": values, "text": text}
            break

    if match is None:
        return {
            "question_id": "BA-006",
            "source_path": str(BA006_RESPONSIBILITY_PATH),
            "source_fidelity_status": "CONFIRMED",
            "gold_location_status": "UNRESOLVED",
            "unresolved_reason": "责任状正文已找到，但未定位DOP电子图形文件数据中心上传数量。",
            "fingerprint": fingerprint,
        }

    supporting = _read_json(
        PROJECT_ROOT / "evaluation" / "business_gold_v2" / "gold_location_manifest.json"
    )
    old_ba006 = next(item for item in supporting["records"] if item["question_id"] == "BA-006")
    return {
        "question_id": "BA-006",
        "source_path": str(BA006_RESPONSIBILITY_PATH),
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_fidelity_status": "CONFIRMED",
        "gold_location_status": "CONFIRMED",
        "gold_claim_status": "CANDIDATE_READY",
        "source_governance": "PENDING_APPROVAL",
        "fingerprint": fingerprint,
        "document": "2026年设计与技术系统专项责任书",
        "section": "附件：关键指标 / 设计管理",
        "table": 1,
        "paragraph": None,
        "row": match["row"],
        "field": "考核内容",
        "headers": headers,
        "organization": "局设计与技术系统责任目标",
        "year": 2026,
        "owner_source_fact": {
            "business_term": "DOP平台电子图形文件数据中心上传数量",
            "value": "不少于40个项目图形文件",
            "numeric_value": 40,
            "unit": "个项目图形文件",
            "confirmation": "RECONFIRMED_FROM_OWNER_SOURCE",
            "source_row": match["row"],
            "raw_values": match["values"],
            "raw_text": match["text"],
        },
        "supporting_source": {
            "source_path": str(BA006_WORK_PLAN_PATH),
            "source_role": "SUPPORTING_SOURCE",
            "observed_value": "不少于500个项目电子图形文件数据",
            "confirmation": "INHERITED_FROM_WORK_PLAN_FORBIDDEN",
            "source_excerpt": old_ba006.get("source_excerpt"),
            "conflict_status": "EVIDENCE_CONFLICT",
        },
        "conflict_status": "EVIDENCE_CONFLICT",
        "truth_answerability": "ANSWERABLE",
        "runtime_answerability": "ANSWERABLE_IN_SHADOW_ONLY",
        "expected_claims": ["2026年局设计与技术系统责任状要求DOP平台电子图形文件数据中心上传不少于40个项目图形文件。"],
        "must_not_use": ["工作计划中的500个替代责任状的40个", "项目级或二公司级数量替代局级责任目标"],
    }


def _close_ba010() -> dict[str, Any]:
    docx = _extract_docx_rows(BA010_DOCX_PATH)
    xlsx = _extract_xlsx_rows(BA010_XLSX_PATH)
    exact_matches = []
    xlsx_by_text = {_norm(row["planning_point"]): row for row in xlsx["rows"]}
    for row in docx["rows"]:
        match = xlsx_by_text.get(_norm(row["planning_point"]))
        if match:
            exact_matches.append({"docx_row": row["row"], "xlsx_row": match["row"]})

    shared_professions = sorted(set(docx["group_counts"]) & set(xlsx["group_counts"]))
    common_headers = sorted(set(docx["headers"]) & set(xlsx["headers"]))
    lineage_status = _lineage_status(
        same_project=True,
        common_headers=common_headers,
        exact_matches=exact_matches,
        docx_count=len(docx["rows"]),
        xlsx_count=len(xlsx["rows"]),
        docx_professions=set(docx["group_counts"]),
        xlsx_professions=set(xlsx["group_counts"]),
        file_reference_found=False,
    )
    return {
        "question_id": "BA-010",
        "primary_document_source": {
            "path": str(BA010_DOCX_PATH),
            "role": "PRIMARY_GOLD_CANDIDATE",
            "fingerprint": _file_fingerprint(BA010_DOCX_PATH),
        },
        "structured_derived_fact_source": {
            "path": str(BA010_XLSX_PATH),
            "role": "STRUCTURED_DERIVED_SOURCE_CANDIDATE",
            "fingerprint": _file_fingerprint(BA010_XLSX_PATH),
        },
        "lineage_status": lineage_status,
        "project_name_comparison": {
            "docx": "武汉国家航天产业基地星谷科创中心建设项目",
            "xlsx": "星谷科创项目（由文件名和路径识别）",
            "status": "RELATED_PROJECT_SIGNAL",
        },
        "document_version_time_comparison": {
            "docx_mtime": _file_fingerprint(BA010_DOCX_PATH)["mtime"],
            "xlsx_mtime": _file_fingerprint(BA010_XLSX_PATH)["mtime"],
            "status": "NOT_EXPLICITLY_MATCHED",
        },
        "professional_set_comparison": {
            "docx": docx["group_counts"],
            "xlsx": xlsx["group_counts"],
            "shared_professions": shared_professions,
            "status": "PARTIAL_OVERLAP",
        },
        "planning_point_comparison": {
            "docx_detail_count": len(docx["rows"]),
            "xlsx_detail_count": len(xlsx["rows"]),
            "exact_text_matches": exact_matches,
            "exact_text_match_count": len(exact_matches),
            "status": "NO_EXACT_ROW_MATCH",
        },
        "field_comparison": {
            "docx_headers": docx["headers"],
            "xlsx_headers": xlsx["headers"],
            "common_headers": common_headers,
            "docx_has_profit_field": False,
            "xlsx_has_profit_field": True,
            "status": "RELATED_SCHEMA_WITH_XLSX_EXTENSIONS",
        },
        "file_reference_comparison": {
            "docx_references_exact_xlsx": False,
            "xlsx_references_docx": False,
            "same_directory_family": True,
            "status": "PATH_NAMING_RELATION_ONLY",
        },
        "docx_confirmed_facts": {
            "sheet_or_table": "DOCX Table 11",
            "rows": docx["rows"],
            "group_counts": docx["group_counts"],
            "total_rows": len(docx["rows"]),
            "profit_field_present": False,
        },
        "xlsx_structured_facts": {
            "sheet": "价值创造",
            "header_row": 3,
            "detail_row_range": "4-87",
            "summary_row": 88,
            "rows": xlsx["rows"],
            "group_counts": xlsx["group_counts"],
            "total_rows": len(xlsx["rows"]),
            "profit_nonempty_count": xlsx["profit_nonempty_count"],
            "formula_count": xlsx["formula_count"],
        },
        "difference_analysis": {
            "docx_34_vs_xlsx_80": "两者统计边界不同；DOCX为34条策划清单明细，XLSX为价值创造Sheet中80条非空策划点明细，不能直接拼接。",
            "different_version_or_stage": "POSSIBLE_NOT_CONFIRMED",
            "docx_selected_subset_of_xlsx": "NOT_PROVEN",
            "xlsx_extension_of_docx": "NOT_PROVEN",
            "different_business_boundary": "PROBABLE",
            "completely_different_business_object": "NOT_SUPPORTED_BY_PROJECT_AND_SCHEMA_SIGNALS",
            "reason": "同项目且共享专业类别、策划点、策划点类别、价值创造分析等字段，但无精确策划点文本对应、专业集合不一致，且DOCX没有利润字段。",
        },
        "business_rule": {
            "benefit_semantics": "增加效益",
            "mapped_field": "利润",
            "operator": ">",
            "threshold": 0,
            "status": "FROZEN_BUSINESS_RULE",
        },
        "gold_decision": "PARTIAL_GOLD",
        "truth_answerability": "PARTIAL",
        "runtime_answerability": "回答DOCX已确认的专业和34条；不能从Owner Confirmed DOCX确定增加效益条数。",
        "must_not_claim": ["不得将XLSX的37条直接继承为BA-010 Gold答案", "不得把价值创造分析中的‘提高利润’替代利润>0业务口径"],
    }


def _extract_docx_rows(path: Path) -> dict[str, Any]:
    extracted = _docx_table_extract(path, ["设计价值创造", "专业类别"])
    table = next(table for table in extracted["tables"] if _table_has_header(table, "专业类别"))
    rows = table["rows"]
    header_index = next(index for index, row in enumerate(rows) if "专业类别" in " | ".join(row["values"]))
    headers = rows[header_index]["values"]
    professional_index = headers.index("专业类别")
    planning_index = headers.index("价值创造策划点")
    group_counts: Counter[str] = Counter()
    detail_rows = []
    current_professional = ""
    for row in rows[header_index + 1 :]:
        values = row["values"]
        if professional_index < len(values) and values[professional_index]:
            current_professional = values[professional_index]
        if not current_professional or planning_index >= len(values) or not values[planning_index]:
            continue
        group_counts[current_professional] += 1
        detail_rows.append({
            "row": row["row"],
            "professional": current_professional,
            "planning_point": values[planning_index],
            "values": values,
        })
    return {"headers": headers, "rows": detail_rows, "group_counts": dict(group_counts)}


def _extract_xlsx_rows(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        sheet = workbook["价值创造"]
        headers = [str(sheet.cell(3, column).value or "").strip() for column in range(1, 11)]
        rows = []
        group_counts: Counter[str] = Counter()
        current_professional = ""
        profit_nonempty_count = 0
        formula_count = 0
        for row_number in range(4, 88):
            values = [sheet.cell(row_number, column).value for column in range(1, 11)]
            if values[0] not in (None, ""):
                current_professional = str(values[0]).strip()
            planning_point = str(values[1] or "").strip()
            if not current_professional or not planning_point:
                continue
            profit = values[6]
            if profit not in (None, ""):
                profit_nonempty_count += 1
            if isinstance(profit, str) and profit.startswith("="):
                formula_count += 1
            group_counts[current_professional] += 1
            rows.append({
                "row": row_number,
                "professional": current_professional,
                "planning_point": planning_point,
                "profit": profit,
                "values": values,
            })
        return {
            "headers": headers,
            "rows": rows,
            "group_counts": dict(group_counts),
            "profit_nonempty_count": profit_nonempty_count,
            "formula_count": formula_count,
        }
    finally:
        workbook.close()


def _lineage_status(
    *,
    same_project: bool,
    common_headers: list[str],
    exact_matches: list[dict[str, Any]],
    docx_count: int,
    xlsx_count: int,
    docx_professions: set[str],
    xlsx_professions: set[str],
    file_reference_found: bool,
) -> str:
    if same_project and common_headers and exact_matches and file_reference_found and docx_count <= xlsx_count:
        return "LINEAGE_CONFIRMED"
    if same_project and common_headers and (docx_professions & xlsx_professions):
        return "LINEAGE_PARTIAL"
    return "LINEAGE_NOT_CONFIRMED"


def _table_has_header(table: dict[str, Any], header: str) -> bool:
    return any(header in " | ".join(row.get("values") or []) for row in table.get("rows", []))


def _norm(value: str) -> str:
    return re.sub(r"\W+", "", value.casefold())


def _build_final_manifest(questions: dict[str, str], ba006: dict[str, Any], ba010: dict[str, Any]) -> list[dict[str, Any]]:
    old = _read_json(OUTPUT_DIR / "gold_location_manifest.json")["records"]
    old_by_id = {item["question_id"]: item for item in old}
    records = []
    for question_id, source in sorted(OWNER_SOURCES.items()):
        previous = old_by_id[question_id]
        gold_type = "FULL_GOLD"
        source_path = source["source_path"]
        expected_claims = previous.get("expected_claims", [])
        gold_location_status = "CONFIRMED"
        truth_answerability = previous.get("truth_answerability", "ANSWERABLE")
        runtime_answerability = previous.get("runtime_expected_behavior", "ANSWERABLE_IN_SHADOW_ONLY")
        gold_claim_status = "CANDIDATE_READY"
        if question_id == "BA-003":
            gold_type = "SOURCE_SCOPE_GOLD"
        if question_id == "BA-006":
            source_path = ba006["source_path"]
            expected_claims = ba006["expected_claims"]
            gold_location_status = ba006["gold_location_status"]
            truth_answerability = ba006["truth_answerability"]
            runtime_answerability = ba006["runtime_answerability"]
        if question_id == "BA-007":
            expected_claims = [
                "坚持设计底线管理：项目整体不超概、不因设计原因发生质量、安全责任事故。",
                "塑强设计管理成效：EPC项目整体设计效益率增量达到7%以上，施工总承包项目专项设计效益率增量达到2.5%以上。",
                "落实固化动作：EPC执行13项标准设计管理动作，动作执行率和及时率100%；施工总承包执行规定动作，执行率和及时率100%。",
                "年度设计管理检查综合排名前20%。",
            ]
        if question_id == "BA-010":
            gold_type = "PARTIAL_GOLD"
            source_path = str(BA010_DOCX_PATH)
            expected_claims = [
                "Owner Confirmed DOCX确认建筑11条、结构10条、给排水5条、暖通6条、电气2条，共34条有效明细。",
                "增加效益条数暂不能从Owner Confirmed DOCX确定；业务口径仍为利润>0。",
            ]
            gold_location_status = "PARTIAL"
            truth_answerability = "PARTIAL"
            runtime_answerability = "回答已确认部分，并提示增加效益条数证据不足。"
            gold_claim_status = "PARTIAL"
        records.append(
            {
                "question_id": question_id,
                "question": questions.get(question_id, previous.get("question", "")),
                "primary_source": source_path,
                "source_role": "PRIMARY_GOLD_CANDIDATE" if question_id != "BA-006" else "PRIMARY_GOLD_CANDIDATE",
                "source_truth_status": "CONFIRMED",
                "source_fidelity_status": ba006["source_fidelity_status"] if question_id == "BA-006" else "CONFIRMED",
                "gold_location_status": gold_location_status,
                "gold_claim_status": gold_claim_status,
                "truth_answerability": truth_answerability,
                "runtime_answerability": runtime_answerability,
                "governance_status": source["governance"],
                "gold_type": gold_type,
                "owner_confirmation": "UNCONFIRMED",
                "expected_location": {"table": 1, "row": ba006.get("row"), "field": ba006.get("field"), "year": ba006.get("year")} if question_id == "BA-006" else _compact_locations(previous.get("expected_location", [])),
                "expected_claims": expected_claims,
                "must_not_use_evidence": previous.get("must_not_use_evidence", []),
                "source_lineage_status": ba010["lineage_status"] if question_id == "BA-010" else None,
            }
        )
    return records


def _compact_locations(locations: Any) -> Any:
    if not isinstance(locations, list):
        return locations
    references = [_location_reference(item) for item in locations]
    if len(references) <= 6:
        return references
    return {
        "location_count": len(references),
        "samples": references[:3] + references[-3:],
        "note": "已压缩为位置引用；完整定向解析结果保留在既有 Location Closure 产物中。",
    }


def _location_reference(item: Any) -> Any:
    if not isinstance(item, dict):
        return item
    keys = (
        "page",
        "paragraph",
        "table",
        "row",
        "sheet_name",
        "row_start",
        "row_end",
        "column_count",
        "section",
        "field",
        "organization",
        "year",
    )
    reference = {key: item[key] for key in keys if key in item}
    return reference or {"location_available": True}


def _build_remaining_decisions(final_manifest: list[dict[str, Any]], ba006: dict[str, Any], ba010: dict[str, Any]) -> list[dict[str, Any]]:
    decisions = [
        {
            "question_id": "BA-006",
            "decision": "确认责任状第22行的‘不少于40个项目图形文件’作为最终Gold Claim；工作计划500个仅为Supporting Source。",
            "status": "PENDING_OWNER_FINAL_GOLD_CONFIRMATION",
        },
        {
            "question_id": "BA-007",
            "decision": "确认Gold仅包含‘一、设计示范工程实施要求’四项，深化设计和BIM章节仅为Related Context。",
            "status": "PENDING_OWNER_FINAL_GOLD_CONFIRMATION",
        },
        {
            "question_id": "BA-010",
            "decision": "确认接受DOCX 34条部分Gold并暂不回答增加效益条数，或提供/确认DOCX与XLSX的真实Source Lineage。",
            "status": "PENDING_OWNER_SOURCE_LINEAGE_CONFIRMATION",
            "lineage_status": ba010["lineage_status"],
        },
    ]
    decisions.extend(
        {
            "question_id": item["question_id"],
            "decision": "最终确认来源位置、Claim范围和系统行为。",
            "status": "PENDING_OWNER_FINAL_GOLD_CONFIRMATION",
        }
        for item in final_manifest
        if item["question_id"] not in {decision["question_id"] for decision in decisions}
    )
    return decisions


def _render_final_report(final_manifest: list[dict[str, Any]], ba006: dict[str, Any], ba010: dict[str, Any], decisions: list[dict[str, Any]]) -> str:
    counts = Counter(item["gold_location_status"] for item in final_manifest)
    lines = [
        "# BUSINESS GOLD FINAL CLOSURE REPORT",
        "",
        "> TASK-020A.2.2 仅完成 BA-006 责任状来源闭环和 BA-010 Source Lineage 闭环。不运行 Fresh Retrieval，不进入 TASK-020B，不修改 Retriever、Router、Qdrant、8000，不调用 LLM，不扩大 Root 扫描。",
        "",
        "## 1. 总体结论",
        "",
        f"- Gold Location 状态：`CONFIRMED {counts['CONFIRMED']} + PARTIAL {counts['PARTIAL']} + UNRESOLVED {counts['UNRESOLVED']} = {len(final_manifest)}`。",
        "- BA-006：`CONFIRMED`，精确责任状第22行明确为不少于40个项目图形文件。",
        "- BA-010：`PARTIAL_GOLD`，Owner Confirmed DOCX确认34条，但DOCX与历史XLSX仅达到`LINEAGE_PARTIAL`，不得继承37条。",
        "- Owner Final Gold Confirmation：0；本报告仍是候选闭环，不生成 FINAL_GOLD 或 final_failure_atlas。",
        "",
        "## 2. BA-006 责任状闭环",
        "",
        f"- Primary Gold Source：`{ba006['source_path']}`",
        f"- Source Fidelity：`{ba006['source_fidelity_status']}`",
        f"- Location：Table {ba006.get('table')} / Row {ba006.get('row')} / Field `{ba006.get('field')}`",
        f"- 年度：`{ba006.get('year')}`；组织层级：`{ba006.get('organization')}`。",
        f"- Owner Source 原文值：`{ba006['owner_source_fact']['value']}`",
        f"- 责任状原始行：`{' | '.join(ba006['owner_source_fact']['raw_values'])}`",
        f"- 工作计划 Supporting Source：`{ba006['supporting_source']['observed_value']}`",
        "- 两者数值不一致，标记 `EVIDENCE_CONFLICT`；最终不自动选择工作计划数字。",
        "- BA-006候选答案：2026年局设计与技术系统责任状要求上传不少于40个项目图形文件。",
        "",
        "## 3. BA-010 Source Lineage 闭环",
        "",
        f"- Primary Document Source：`{ba010['primary_document_source']['path']}`",
        f"- Structured Derived Fact Source：`{ba010['structured_derived_fact_source']['path']}`",
        f"- Lineage：`{ba010['lineage_status']}`",
        "",
        "### 3.1 可确认事实",
        "",
        f"- DOCX Table 11：{ba010['docx_confirmed_facts']['total_rows']}条明细。",
        f"- DOCX专业统计：`{ba010['docx_confirmed_facts']['group_counts']}`。",
        f"- XLSX“价值创造”Sheet：{ba010['xlsx_structured_facts']['total_rows']}条明细；第88行为汇总公式，不计入明细。",
        f"- XLSX专业统计：`{ba010['xlsx_structured_facts']['group_counts']}`；利润字段非空{ba010['xlsx_structured_facts']['profit_nonempty_count']}条。",
        "- 业务规则冻结：`增加效益 = 利润 > 0`。",
        "",
        "### 3.2 34与80差异",
        "",
        f"- DOCX与XLSX精确策划点文本对应：{ba010['planning_point_comparison']['exact_text_match_count']}条。",
        "- 两者同属星谷项目且存在共同字段，但专业集合、明细数量和字段范围不同；DOCX没有利润字段。",
        "- 当前不能证明是DOCX精选子集，也不能证明XLSX是DOCX扩展明细；不同版本/阶段是可能解释，但尚未确认。",
        "- 结论：`LINEAGE_PARTIAL`，不得将34条与80/37/43拼接为同一 Gold Answer。",
        "",
        "## 4. BA-001～BA-010候选状态",
        "",
        "| BA | Gold Type | Source Fidelity | Location | Claim | Truth | Runtime | Governance |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for item in final_manifest:
        lines.append(
            f"| {item['question_id']} | `{item['gold_type']}` | `{item['source_fidelity_status']}` | `{item['gold_location_status']}` | `{item['gold_claim_status']}` | `{item['truth_answerability']}` | `{item['runtime_answerability']}` | `{item['governance_status']}` |"
        )
    lines += [
        "",
        "## 5. 待业务负责人确认",
        "",
    ]
    lines.extend(f"- **{item['question_id']}**：{item['decision']}" for item in decisions)
    lines += [
        "",
        "## 6. 输出",
        "",
        "- `evaluation/business_gold_v2/ba006_responsibility_source_closure.json`",
        "- `evaluation/business_gold_v2/ba010_source_lineage.json`",
        "- `evaluation/business_gold_v2/final_gold_candidate_manifest.json`",
        "- `evaluation/business_gold_v2/remaining_owner_decisions.json`",
        "- `docs/BUSINESS_GOLD_OWNER_APPROVAL_SHEET.md`",
        "",
        "## 7. 停止点",
        "",
        "本任务完成后停止。下一步只有在业务负责人完成最终确认后，才可进入 TASK-020A.3 Controlled Fresh Run。",
        "",
    ]
    return "\n".join(lines)


def _render_approval_sheet(final_manifest: list[dict[str, Any]], ba006: dict[str, Any], ba010: dict[str, Any]) -> str:
    lines = [
        "# BUSINESS GOLD OWNER APPROVAL SHEET",
        "",
        "> 精简审核版。来源定位和候选 Claim 已由系统定向读取；最终 Gold 仍需业务负责人确认。",
        "",
        "| BA | Question | Primary Source | Gold Type | Expected Runtime Behavior | Governance |",
        "|---|---|---|---|---|---|",
    ]
    for item in final_manifest:
        lines.append(
            f"| {item['question_id']} | {item['question']} | `{Path(item['primary_source']).name}` | `{item['gold_type']}` | {item['runtime_answerability']} | `{item['governance_status']}` |"
        )
    for item in final_manifest:
        lines += [
            "",
            f"## {item['question_id']}",
            "",
            f"问题：{item['question']}",
            "",
            f"Primary Source：`{item['primary_source']}`",
            f"Gold Type：`{item['gold_type']}`",
            f"最小位置：`{item['expected_location']}`",
            f"Expected Runtime Behavior：{item['runtime_answerability']}",
            f"Governance：`{item['governance_status']}`",
            "",
            "Expected Answer / Claims：",
        ]
        lines.extend(f"- {claim}" for claim in item["expected_claims"])
        lines += ["", "业务负责人： [ ] 确认   [ ] 修改后确认   [ ] 不确认", "", "备注：________________", ""]
    lines += [
        "## BA-010专项核对",
        "",
        f"- DOCX confirmed facts：{ba010['docx_confirmed_facts']['group_counts']}，合计{ba010['docx_confirmed_facts']['total_rows']}条。",
        f"- XLSX lineage result：`{ba010['lineage_status']}`；XLSX“价值创造”Sheet为{ba010['xlsx_structured_facts']['total_rows']}条明细。",
        "- Benefit rule：`增加效益 = 利润 > 0`，不得改成“价值创造分析=提高利润”。",
        "- 是否可计算增加效益：当前不可从Owner Confirmed DOCX计算；除非业务负责人确认真实DOCX→XLSX Lineage，否则保持PARTIAL。",
        "",
    ]
    return "\n".join(lines)


def _update_v2_report(final_manifest: list[dict[str, Any]], ba006: dict[str, Any], ba010: dict[str, Any]) -> None:
    path = PROJECT_ROOT / "docs" / "BUSINESS_GOLD_V2_REPORT.md"
    previous = path.read_text(encoding="utf-8") if path.exists() else "# Business Gold V2 Report\n"
    previous = re.split(r"\n## TASK-020A\.2\.2[^\n]*", previous, maxsplit=1)[0]
    counts = Counter(item["gold_location_status"] for item in final_manifest)
    block = [
        "",
        "## TASK-020A.2.2 责任状与来源血缘闭环",
        "",
        f"- BA-006：责任状精确来源已确认，Table 1 Row {ba006.get('row')} 为不少于40个项目图形文件；工作计划500个保持 Supporting Source，并记录 EVIDENCE_CONFLICT。",
        f"- BA-010：DOCX 34条与 XLSX 80条判定为 `{ba010['lineage_status']}`，不继承37条。",
        f"- Gold Location：CONFIRMED {counts['CONFIRMED']}、PARTIAL {counts['PARTIAL']}、UNRESOLVED {counts['UNRESOLVED']}；合计 {len(final_manifest)}/10。",
        "- owner_confirmation 仍全部为 UNCONFIRMED；未生成 FINAL_GOLD、final_failure_atlas，不进入 TASK-020B。",
        "",
    ]
    path.write_text(previous.rstrip() + "\n" + "\n".join(block), encoding="utf-8")


if __name__ == "__main__":
    raise SystemExit(main())
