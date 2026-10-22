from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import pytest
from app.retrieval.anchored_sections import sections_for,expand_candidates,effective_outline


def fixture(tmp_path):
    doc=Document();items=[('1、甲专业','Heading 1'),('1.1方案阶段',None),('方案要求。',None),('1.2初步设计阶段',None),('初设要求甲。',None),('初设要求乙。',None),('1.3施工图阶段',None),('施工图要求。',None)]
    lines=[]
    for i,(text,style) in enumerate(items,1):
        p=doc.add_paragraph(text,style);p._p.set(qn('w14:paraId'),f'{i:08X}')
        lines.append(f'[/body/p[@paraId={i:08X}]] '+text)
    native=tmp_path/'source.docx';doc.save(native)
    md=tmp_path/'source.md';md.write_text('\n'.join(lines),encoding='utf-8')
    return md,native


def test_declared_parent_numbered_siblings_and_exact_section_bounds(tmp_path):
    md,native=fixture(tmp_path);sections=sections_for(md,native)
    section=next(s for s in sections if s['title']=='1.2初步设计阶段')
    assert section['heading_path']=='1、甲专业 > 1.2初步设计阶段'
    assert section['paragraph_anchors']==['00000004','00000005','00000006']
    assert '初设要求乙' in section['text'] and '施工图' not in section['text'] and '方案要求' not in section['text']
    assert section['heading_origin'].startswith('NUMBERED_SIBLINGS_INFERRED')
    assert md.read_text(encoding='utf-8').splitlines()[section['location']['line_start']-1].endswith('1.2初步设计阶段')


def test_no_numbered_heading_guess_outside_declared_matching_parent(tmp_path):
    md,native=fixture(tmp_path);doc=Document(native);doc.paragraphs[0].style=doc.styles['Normal'];doc.save(native)
    assert sections_for(md,native)==[]


def test_modified_export_fails_instead_of_borrowing_native_outline(tmp_path):
    md,native=fixture(tmp_path);md.write_text(md.read_text(encoding='utf-8').replace('要求甲','改过的要求'),encoding='utf-8')
    with pytest.raises(ValueError,match='does not match'):sections_for(md,native)


def test_expand_hit_to_full_section_and_preserve_other_source(tmp_path):
    md,native=fixture(tmp_path);ss=sections_for(md,native);s=next(s for s in ss if s['title'].startswith('1.2'))
    hit={'evidence_id':'child','source_path':str(md),'source_version':s['source_version'],'text':'[/body/p[@paraId=00000005]] 初设要求甲。','role':'DIRECT'}
    other={**hit,'evidence_id':'other','source_path':str(tmp_path/'other.md')}
    result=expand_candidates([hit,hit,other],ss)
    assert len(result)==2 and result[1]==other
    assert '初设要求乙' in result[0]['text'] and '施工图' not in result[0]['text']
    assert result[0]['location']['paragraph_anchor']=='00000004'


def test_explicit_body_outline_overrides_inherited_heading(tmp_path):
    from docx.oxml import OxmlElement
    doc=Document();p=doc.add_paragraph('正文','Heading 1');n=OxmlElement('w:outlineLvl');n.set(qn('w:val'),'9');p._p.get_or_add_pPr().append(n)
    assert effective_outline(p) is None


def test_bounded_section_flows_through_existing_query_without_generation(tmp_path):
    from app.trial.assisted_chat import AssistedQuery,run_flow
    from app.answer_engine.llm.answer_generator import ShadowAnswerGenerator
    from app.answer_engine.llm.llm_provider import UnavailableLLMProvider
    md,native=fixture(tmp_path);sections=sections_for(md,native)
    hit={'evidence_id':'child','document_id':'doc','source_id':'source','source_path':str(md),'source_version':sections[0]['source_version'],'text':'[/body/p[@paraId=00000005]] 初设要求甲。','role':'DIRECT'}
    bundle={'candidate_evidence':expand_candidates([hit],sections)}
    events=list(run_flow(AssistedQuery(question='初设有哪些要求？',mode='sources'),lambda _:bundle,ShadowAnswerGenerator(UnavailableLLMProvider('disabled')),lambda *_:None))
    result=next(e['result'] for e in events if e['type']=='result')
    assert result['answer_status']=='SOURCES_ONLY'
    card=result['citations'][0]
    assert card['location']['paragraph_anchor']=='00000004' and '初设要求乙' in card['excerpt'] and '施工图' not in card['excerpt']
