import hashlib

import pytest
from docx import Document
from fastapi import HTTPException

from app.ingestion.loaders.docx_loader import load_docx
from app.retrieval.query_planner_v1 import (
    annual_plan_kind, design_deliverable_stages, design_depth_stages, plan_query,
)
from app.trial.project_navigation import record_in_project, resolve_project_question
from app.trial.runtime_reporting import (
    feedback_fingerprint, feedback_view, knowledge_view, source_counts, traffic_kind,
)


def test_word_body_and_table_keep_original_order(tmp_path):
    document = Document()
    document.add_heading('演示章节', level=1)
    document.add_paragraph('表前内容')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '字段'
    table.cell(0, 1).text = '数值'
    table.cell(1, 0).text = '数量'
    table.cell(1, 1).text = '7'
    document.add_paragraph('表后内容')
    path = tmp_path / 'synthetic.docx'
    document.save(path)
    result = load_docx(path, 'demo')
    assert result.status == 'parsed'
    assert len(result.blocks) == 3
    assert '表前内容' in result.blocks[0].text
    assert '数量 | 7' in result.blocks[1].text
    assert '表后内容' in result.blocks[2].text
    assert all(block.heading_path == '演示章节' for block in result.blocks)


def test_two_annual_plan_scopes_are_separate():
    assert annual_plan_kind('公司年度工作计划如何检查执行情况？') == 'CORPORATE_WORK_PLAN'
    assert annual_plan_kind('项目年度设计进度计划如何检查与纠偏？') == 'PROJECT_DESIGN_SCHEDULE'
    plan = plan_query('公司年度工作计划和项目年度设计进度计划分别如何检查执行情况？')
    assert len(plan.subquestions) == 2


def test_design_depth_is_not_deliverable_or_schedule():
    assert design_depth_stages('方案和施工图的设计深度要求是什么？') == ['方案设计', '施工图设计']
    assert design_depth_stages('施工图交付物清单有哪些？') == []
    assert design_depth_stages('施工图深度审批时间是多少？') == []
    assert design_deliverable_stages('初步设计和施工图分别提交哪些经济文件？') == ['初步设计', '施工图设计']


def test_generic_question_does_not_require_a_named_project():
    assert plan_query('如何编制项目设计任务书？').project == []


def test_source_stats_require_active_version_and_exclude_withdrawn():
    rows = [
        {'index_status': 'INDEXED', 'source_version': 'v1'},
        {'index_status': 'INDEXED'},
        {'index_status': 'INDEXED', 'source_version': 'v2', 'withdrawn': True},
        {'index_status': 'NOT_READY', 'source_version': 'v3'},
    ]
    assert source_counts(rows) == {
        'sources': 4, 'indexed_sources': 1, 'unavailable_sources': 3,
        'source_count_scope': 'CURRENT_RUNTIME_CATALOG_SOURCE_VERSIONS',
    }


def test_answered_is_not_evidence_of_real_use():
    assert traffic_kind({'status': 'ANSWERED', 'conversation_id': 'random-id'}) == 'UNCLASSIFIED'
    assert traffic_kind({'conversation_id': 'DEMO_CANARY_1'}) == 'TECHNICAL'
    assert traffic_kind({'traffic_kind': 'REAL'}) == 'REAL'


def test_feedback_annotation_requires_candidate_and_unchanged_record():
    row = {'feedback_id': 'demo', 'question': '演示问题', 'status': 'RECEIVED'}
    annotation = {'candidate_hash': 'v1', 'records': {'demo': {
        'feedback_fingerprint': feedback_fingerprint(row),
        'technical_fix_verified': True, 'business_review_status': 'ACCEPTED',
    }}}
    assert feedback_view([row], annotation, 'v1')[0]['processing_status'] == 'CLOSED'
    assert feedback_view([row], annotation, 'v2')[0]['processing_status'] == 'PENDING_REVIEW'
    assert feedback_view([{**row, 'question': '已变更'}], annotation, 'v1')[0]['processing_status'] == 'PENDING_REVIEW'


def test_knowledge_overlay_requires_content_and_source_version():
    item = {'knowledge_id': 'demo', 'content': '合成内容', 'source_version': 'v1', 'node_id': 'old', 'title': '旧标题'}
    annotation = {'candidate_hash': 'candidate', 'records': {'demo': {
        'source_version': 'v1', 'content_sha256': hashlib.sha256(item['content'].encode()).hexdigest(), 'title': '新标题',
    }}}
    assert knowledge_view([item], annotation, 'candidate')[0]['title'] == '新标题'
    assert knowledge_view([{**item, 'source_version': 'v2'}], annotation, 'candidate')[0]['title'] == '旧标题'


def test_project_navigation_rejects_other_project_and_version():
    project = {'id': 'a', 'name': '示例甲项目', 'aliases': ['示例甲'], 'sources': [{'source_id': 's', 'source_version': 'v1'}]}
    other = {'id': 'b', 'name': '示例乙项目', 'aliases': ['示例乙'], 'sources': []}
    navigation = {'projects': [project, other]}
    assert record_in_project({'source_id': 's', 'source_version': 'v1'}, project)
    assert not record_in_project({'source_id': 's', 'source_version': 'v2'}, project)
    assert resolve_project_question('合同界面是什么？', 'a', navigation, {}) == '示例甲项目：合同界面是什么？'
    with pytest.raises(HTTPException):
        resolve_project_question('示例乙的合同界面是什么？', 'a', navigation, {})
