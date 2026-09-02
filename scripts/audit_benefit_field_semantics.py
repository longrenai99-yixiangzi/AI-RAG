from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

TARGET_WORKBOOK = (
    Path(r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷")
    / "星谷科创项目设计管理策划+设计示范项目打造方案"
    / "方案比选与价值创造清单方案比选及价值创造.xlsx"
)
PROJECT_DOCX = Path(r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷\设计管理策划书-星谷科创中心项目.docx")
REVIEW_XLSX = Path(r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷\设计策划评审意见表(1).xlsx")
STAGING = PROJECT_ROOT / "data" / "shadow" / "root002_import" / "pipeline_staging.jsonl"
REPORT = PROJECT_ROOT / "docs" / "BENEFIT_FIELD_SEMANTICS_AUDIT.md"

TERMS = (
    "效益",
    "增加效益",
    "效益增量",
    "创效",
    "创效金额",
    "利润",
    "收入",
    "涉及金额",
    "成本",
    "收益",
    "价值创造",
)


def value_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def numeric(value: Any) -> tuple[float | None, str]:
    text = value_text(value)
    if not text or text in {"/", "-", "—", "无", "NA", "N/A"}:
        return None, "empty_or_placeholder"
    match = re.search(r"[-+]?\d+(?:,\d{3})*(?:\.\d+)?", text)
    if not match:
        return None, "text_or_unparseable"
    try:
        return float(match.group(0).replace(",", "")), "numeric"
    except ValueError:
        return None, "numeric_parse_error"


def workbook_audit(path: Path) -> dict[str, Any]:
    formulas: list[dict[str, Any]] = []
    comments: list[dict[str, Any]] = []
    term_hits: dict[str, list[dict[str, Any]]] = {term: [] for term in TERMS}
    sheets: list[dict[str, Any]] = []
    workbook = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    try:
        for worksheet in workbook.worksheets:
            cached_sheet = cached[worksheet.title]
            nonempty = []
            for row in worksheet.iter_rows():
                row_values = [value_text(cell.value) for cell in row if not isinstance(cell, MergedCell)]
                if any(row_values):
                    nonempty.append(row_values)
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    value = value_text(cell.value)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas.append(
                            {
                                "sheet": worksheet.title,
                                "cell": cell.coordinate,
                                "formula": cell.value,
                                "cached_value": cached_sheet[cell.coordinate].value,
                            }
                        )
                    if cell.comment:
                        comments.append(
                            {
                                "sheet": worksheet.title,
                                "cell": cell.coordinate,
                                "author": cell.comment.author,
                                "text": cell.comment.text,
                            }
                        )
                    for term in TERMS:
                        if term in value and len(term_hits[term]) < 20:
                            term_hits[term].append(
                                {
                                    "sheet": worksheet.title,
                                    "cell": cell.coordinate,
                                    "value": value[:300],
                                    "is_formula": isinstance(cell.value, str) and cell.value.startswith("="),
                                }
                            )
            summary_rows = []
            for row_number in range(max(1, worksheet.max_row - 4), worksheet.max_row + 1):
                values = [value_text(worksheet.cell(row_number, column).value) for column in range(1, worksheet.max_column + 1)]
                if any(values):
                    summary_rows.append({"row": row_number, "values": values})
            sheets.append(
                {
                    "sheet": worksheet.title,
                    "state": worksheet.sheet_state,
                    "max_row": worksheet.max_row,
                    "max_column": worksheet.max_column,
                    "nonempty_rows": len(nonempty),
                    "merged_ranges_count": len(worksheet.merged_cells.ranges),
                    "merged_ranges_sample": [str(value) for value in list(worksheet.merged_cells.ranges)[:25]],
                    "first_rows": [
                        [value_text(worksheet.cell(row, column).value) for column in range(1, min(worksheet.max_column, 16) + 1)]
                        for row in range(1, min(worksheet.max_row, 6) + 1)
                    ],
                    "summary_rows": summary_rows,
                }
            )
    finally:
        workbook.close()
        cached.close()
    return {
        "path": str(path),
        "exists": path.exists(),
        "sheets": sheets,
        "valid_sheet_count": sum(sheet["state"] == "visible" and sheet["nonempty_rows"] > 0 for sheet in sheets),
        "formula_count": len(formulas),
        "formula_examples": formulas[:40],
        "comment_count": len(comments),
        "comment_examples": comments[:20],
        "term_hits": term_hits,
    }


def candidate_field_stats(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        sheet = workbook["价值创造"]
        headers = [value_text(sheet.cell(3, column).value) for column in range(1, sheet.max_column + 1)]
        field_stats: dict[str, dict[str, Any]] = {}
        for field in ("涉及金额", "利润", "效益对比（万元）", "效益对比", "增加效益"):
            if field not in headers:
                field_stats[field] = {"available": False}
                continue
            column = headers.index(field) + 1
            rows: list[int] = []
            positive_rows: list[int] = []
            parse_failures: list[dict[str, Any]] = []
            for row in range(4, sheet.max_row + 1):
                value = sheet.cell(row, column).value
                text = value_text(value)
                if not text:
                    continue
                rows.append(row)
                parsed, reason = numeric(value)
                if parsed is not None and parsed > 0:
                    positive_rows.append(row)
                elif parsed is None and reason not in {"empty_or_placeholder"}:
                    parse_failures.append({"row": row, "value": text, "reason": reason})
            field_stats[field] = {
                "available": True,
                "column": column,
                "nonempty_rows": len(rows),
                "numeric_positive_rows": len(positive_rows),
                "positive_row_numbers": positive_rows,
                "parse_failures": parse_failures,
            }
        return {"headers_row_3": headers, "field_stats": field_stats}
    finally:
        workbook.close()


def docx_hits(path: Path) -> dict[str, Any]:
    document = Document(path)
    hits: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text and any(term in text for term in TERMS):
            hits.append({"location": f"paragraph {index}", "text": text[:600]})
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for column_index, cell in enumerate(row.cells, start=1):
                text = " ".join(paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip())
                if text and any(term in text for term in TERMS):
                    hits.append({"location": f"table {table_index}, row {row_index}, col {column_index}", "text": text[:600]})
    return {
        "path": str(path),
        "exists": path.exists(),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "hits": hits[:80],
    }


def staging_hits() -> list[dict[str, Any]]:
    hits: list[dict[str, Any]] = []
    for line in STAGING.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        document = json.loads(line)
        if Path(document.get("path", "")).name not in {PROJECT_DOCX.name, REVIEW_XLSX.name}:
            continue
        for block in document.get("source_blocks", []):
            for line_number, text in enumerate(str(block.get("text", "")).splitlines(), start=1):
                if any(term in text for term in TERMS):
                    hits.append(
                        {
                            "file": document.get("file_name"),
                            "path": document.get("path"),
                            "sheet_or_heading": block.get("heading_path"),
                            "source_location": block.get("location"),
                            "line": line_number,
                            "text": text[:600],
                        }
                    )
    return hits[:120]


def render_report(audit: dict[str, Any], stats: dict[str, Any], project_doc: dict[str, Any], review_audit: dict[str, Any], staging: list[dict[str, Any]]) -> str:
    target_sheet = next(sheet for sheet in audit["sheets"] if sheet["sheet"] == "价值创造")
    other_fields = {term: hits for term, hits in audit["term_hits"].items() if term in {"收入", "效益", "成本", "收益", "创效", "创效金额", "效益增量", "增加效益"}}
    lines = [
        "# Benefit Field Semantics Audit",
        "",
        "> TASK-016E-2A.1：只读审计 BA-010“增加效益”字段语义，不修改 Fact Aggregator、Answer Engine、Retriever、正式 Qdrant 或原始文件。",
        "",
        "## 1. 结论",
        "",
        "**SEMANTICS_AMBIGUOUS**",
        "",
        "审计发现多个可能的金额/效益字段，但没有企业业务资料明确规定“增加效益”唯一对应哪个字段或计算规则。因此本次不得计算“增加效益条数”。",
        "",
        "候选字段：",
        "",
        "- `涉及金额`：价值创造 Sheet 存在，当前非空并可解析为正数的行已统计；但没有资料证明它等同于“增加效益”。",
        "- `利润`：价值创造 Sheet 存在，当前非空并可解析为正数的行已统计；但没有资料证明它等同于“增加效益”。",
        "- `收入` / `效益`：出现在同 Workbook 的“方案比选”Sheet，不应直接替代“价值创造”Sheet 的业务字段。",
        "- `效益对比（万元）` / `增加效益`：目标 Workbook 中未找到明确字段。",
        "",
        "## 2. 权威 Workbook",
        "",
        f"- 文件：`{audit['path']}`",
        f"- 有效 Sheet：`{audit['valid_sheet_count']}` / `{len(audit['sheets'])}`",
        f"- Workbook 公式总数：`{audit['formula_count']}`",
        f"- Workbook 批注数：`{audit['comment_count']}`",
        "",
        "### 有效 Sheet",
        "",
        "| Sheet | 行数 | 列数 | 非空行 | 合并范围数 | 公式数（Workbook级示例） |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    formula_by_sheet = Counter(item["sheet"] for item in audit["formula_examples"])
    for sheet in audit["sheets"]:
        if sheet["state"] == "visible" and sheet["nonempty_rows"] > 0:
            lines.append(f"| {sheet['sheet']} | {sheet['max_row']} | {sheet['max_column']} | {sheet['nonempty_rows']} | {sheet['merged_ranges_count']} | {formula_by_sheet.get(sheet['sheet'], 0)} |")
    lines += [
        "",
        "## 3. Workbook 全部字段与结构检查",
        "",
        "### 3.1 关键字段命中",
        "",
        "| 字段 | 命中情况 | 说明/位置示例 |",
        "|---|---|---|",
    ]
    for term in TERMS:
        hits = audit["term_hits"].get(term, [])
        example = "；".join(f"{hit['sheet']}!{hit['cell']}={hit['value'][:100]}" for hit in hits[:3])
        lines.append(f"| {term} | {'存在' if hits else '未找到'} | {example or '-'} |")
    lines += [
        "",
        "### 3.2 价值创造 Sheet 表头",
        "",
        f"- 第 3 行是语义表头：`{json.dumps(stats['headers_row_3'], ensure_ascii=False)}`",
        "- 目标字段明确存在：`涉及金额`、`利润`。",
        "- 目标字段明确不存在：`效益对比（万元）`、`增加效益`。",
        "",
        "### 3.3 合并表头、说明行、汇总行、公式、批注",
        "",
        f"- 价值创造 Sheet 合并范围示例：`{', '.join(target_sheet['merged_ranges_sample'][:12])}`",
        f"- 价值创造 Sheet 汇总/末尾行：`{json.dumps(target_sheet['summary_rows'], ensure_ascii=False)}`",
        f"- 公式示例：`{json.dumps([item for item in audit['formula_examples'] if item['sheet'] == '价值创造'][:20], ensure_ascii=False)}`",
        f"- 批注：`{json.dumps(audit['comment_examples'], ensure_ascii=False)}`",
        "",
        "## 4. 候选字段统计（只统计，不映射语义）",
        "",
        "| 字段 | 是否存在 | 非空行 | numeric > 0 行 | 解析失败 | 是否认定为增加效益 |",
        "|---|---|---:|---:|---:|---|",
    ]
    for field in ("涉及金额", "利润", "效益对比（万元）", "效益对比", "增加效益"):
        info = stats["field_stats"][field]
        lines.append(f"| {field} | {info.get('available')} | {info.get('nonempty_rows', '-')} | {info.get('numeric_positive_rows', '-')} | {len(info.get('parse_failures', [])) if info.get('available') else '-'} | 否 |")
    lines += [
        "",
        "- `涉及金额` 和 `利润` 的正数统计仅说明数据格式可解析，不代表业务语义已经确认。",
        "- 由于问题没有要求总金额，本次不执行 SUM。",
        "",
        "## 5. 关联 Root-002 资料语义检查",
        "",
        f"### 5.1 {project_doc['path']}",
        "",
    ]
    for hit in project_doc["hits"][:30]:
        lines.append(f"- `{hit['location']}`：{hit['text']}")
    lines += ["", f"### 5.2 {review_audit['path']}", ""]
    for hit in review_audit["term_hits"][:40]:
        lines.append(f"- `{hit['sheet']}!{hit['cell']}`：{hit['value']}")
    lines += [
        "",
        "### 5.3 语义证据判断",
        "",
        "- 项目策划书说明了“设计价值创造率：EPC项目设计阶段效益提升5%”，这是目标/比例，不是“增加效益条数”的字段定义。",
        "- 评审意见指出“设计效益增量未测算”“设计价值创造点清单与设计限额目标中的‘调整金额’不对应”，说明口径尚未闭合。",
        "- 评审意见还要求“按施工界面划分，分开体现不同单位的效益情况”，但没有定义“涉及金额”和“利润”哪一个对应增加效益。",
        "- 没有发现明确规则：`增加效益 = 涉及金额 > 0`，也没有发现明确规则：`增加效益 = 利润 > 0`。",
        "",
        "## 6. Provenance",
        "",
        "- Workbook：`方案比选与价值创造清单方案比选及价值创造.xlsx` / `价值创造` Sheet / 表头第 3 行 / 字段 F-G。",
        "- 项目策划书：`设计管理策划书-星谷科创中心项目.docx` / 2.3 设计阶段效益目标、7. 设计价值创造策划及相关表格。",
        "- 评审表：`设计策划评审意见表(1).xlsx` / Sheet1 / 设计价值创造清单相关意见行。",
        "- 所有候选字段统计均来自权威 Workbook 的真实单元格；未使用其他项目价值创造案例补数字。",
        "",
        "## 7. 业务确认事项",
        "",
        "请业务负责人明确以下一项：",
        "",
        "1. “增加效益”是否等于 `利润 > 0`；",
        "2. “增加效益”是否等于 `涉及金额 > 0`；",
        "3. 是否应使用 Workbook 其他 Sheet 的 `效益` 字段；",
        "4. 是否存在未进入当前 Shadow 的正式业务定义或计算规则。",
        "",
        "在业务负责人确认前，Fact Aggregator 必须保持 `PARTIAL_AGGREGATION`，不得输出“增加效益条数”。",
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    audit = workbook_audit(TARGET_WORKBOOK)
    stats = candidate_field_stats(TARGET_WORKBOOK)
    project_doc = docx_hits(PROJECT_DOCX)
    review_audit = workbook_audit(REVIEW_XLSX)
    review_hits = []
    for term, hits in review_audit["term_hits"].items():
        for hit in hits:
            review_hits.append({"term": term, **hit})
    review_audit = {"path": review_audit["path"], "term_hits": review_hits}
    report = render_report(audit, stats, project_doc, review_audit, [])
    REPORT.write_text(report, encoding="utf-8")
    print(json.dumps({"report": str(REPORT.resolve()), "conclusion": "SEMANTICS_AMBIGUOUS", "valid_sheets": audit["valid_sheet_count"], "formula_count": audit["formula_count"], "comment_count": audit["comment_count"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
