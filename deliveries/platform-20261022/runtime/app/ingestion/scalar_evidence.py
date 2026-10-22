"""Source-bound scalar facts for an offline trial, preserving original units and cells."""
from __future__ import annotations

from decimal import Decimal
import hashlib
from pathlib import Path
import re

from docx import Document
import pymupdf

from app.ingestion.layout_evidence import compact, identity, pptx_evidence, record

METRICS = {
    '初设概算总投资': ('初设概算总投资', '初设总投资'),
    '初设概算费用': ('初设概算费用', '初步设计概算费用'),
    '合同建安工程费': ('合同建安工程费',),
    '建安工程费': ('建安工程费',),
    '工程建设其他费用': ('工程建设其他费用',),
    '合同金额': ('合同金额', '合同额', '合同总价'),
    '工程造价': ('工程造价',),
    '工程费': ('工程费',),
    '预备费': ('预备费',),
    '总投资': ('总投资',),
    '创效金额': ('创效金额', '创效', '创造效益'),
    '提前天数': ('提前',),
}
ALIASES = {alias: metric for metric, aliases in METRICS.items() for alias in aliases}
LABEL = '|'.join(re.escape(v) for v in sorted(ALIASES, key=len, reverse=True))
VALUE = r'[0-9]+(?:,[0-9]{3})*(?:\.[0-9]+)?'
MONEY = re.compile(rf'(?P<label>{LABEL})(?:共计|合计|为|是|约|[：:\s])*(?P<value>{VALUE})\s*(?P<unit>亿元|万元|元|W)(?![A-Za-z])')


def requested_metrics(question):
    return list(dict.fromkeys(ALIASES[m[0]] for m in re.finditer(LABEL, compact(question))))


def project_key(text):
    text=re.sub(r'^(?:请问|请告诉我|请说明|关于|请给出|请)+','',text)
    return re.sub(r'EPC|设计采购施工|工程总承包|总承包|项目|工程|[一二三四五六七八九十]+期|[（()）\s、，,：:—-]', '', text, flags=re.I)


def phases(text):
    return set(re.findall(r'[一二三四五六七八九十]+期', text))


def _admitted(location, config):
    return any(all(location.get(k) == v for k, v in target.items()) for target in config['admitted_locations'])


def _fact(parent, config, metric, value, unit, *, subject='', raw_metric='', location=None, unit_confirmation=None):
    loc = location or parent['location']
    output = {**parent, 'evidence_id': identity(parent['evidence_id'], metric, value, unit, subject, str(loc)),
              'parent_evidence_id': parent['evidence_id'], 'location': loc,
              'source_scope': config['scope'],
              'scalar': {'metric': metric, 'value': str(Decimal(value.replace(',', ''))), 'raw_value': value,
                         'unit': unit_confirmation['unit'] if unit_confirmation else unit,
                         'raw_unit': unit, 'raw_metric': raw_metric or metric, 'subject': subject,
                         'unit_status': 'OWNER_CONFIRMED_EXACT_CELL' if unit_confirmation else ('SOURCE_EXPLICIT' if unit in {'亿元','万元','元','天'} else 'UNRESOLVED'),
                         'unit_confirmation': unit_confirmation,
                         'source_context': parent['text']}}
    output['heading_path'] = config['scope']['project']+' > '+(subject+' > ' if subject else '')+metric
    output['retrieval_text'] = output['heading_path']+' '+value+unit
    return output


def _key_value_table_facts(parent,config,node):
    facts=[]
    for ri,cells in enumerate(node['matrix'],1):
        for ci,text in enumerate(cells,1):
            for match in MONEY.finditer(compact(text)):
                if ALIASES[match['label']]=='创效金额':continue
                loc={**parent['location'],'shape_id':node['shape_id'],'bbox':node['bbox'],'row':ri,'column':ci}
                facts.append(_fact(parent,config,ALIASES[match['label']],match['value'],match['unit'],raw_metric=match['label'],location=loc))
        for ci,label in enumerate(cells[:-1]):
            metrics=requested_metrics(label)
            if len(metrics)!=1 or len(compact(label))>24:
                continue
            # Only a standalone metric label followed by its own value cell is a key/value row.
            normalized=compact(label).strip('：:')
            if normalized not in ALIASES:
                continue
            value_cell=next(((j,v) for j,v in enumerate(cells[ci+1:],ci+1) if compact(v)),None)
            if not value_cell:continue
            column,text=value_cell
            match=re.fullmatch(rf'\s*({VALUE})\s*(亿元|万元|元|亿|万|W)\s*([（(].*[）)])?\s*',text,re.S)
            if not match:continue
            raw_unit=match[2];unit={'亿':'亿元','万':'万元'}.get(raw_unit,raw_unit)
            if raw_unit in {'亿','万'} and metrics[0] not in {'工程造价','合同金额','总投资','初设概算费用','初设概算总投资'}:continue
            loc={**parent['location'],'shape_id':node['shape_id'],'bbox':node['bbox'],'row':ri,'column':column+1,'label_column':ci+1}
            fact=_fact(parent,config,metrics[0],match[1],unit,raw_metric=label,location=loc)
            fact['scalar']['raw_unit']=raw_unit
            fact['scalar']['qualifier']=(match[3] or '').strip('（）()')
            facts.append(fact)
    return facts


def _docx_parents(path, digest):
    document = Document(path)
    parents = [record(path,digest,{'paragraph':i},f'Paragraph {i}',p.text)
               for i,p in enumerate(document.paragraphs,1) if p.text.strip()]
    for ti,table in enumerate(document.tables,1):
        matrix = [[cell.text for cell in row.cells] for row in table.rows]
        for ri,cells in enumerate(matrix,1):
            parent = record(path,digest,{'table':ti,'row':ri},f'Table {ti} Row {ri}',
                            '表头：'+' | '.join(matrix[0])+'\n行：'+' | '.join(cells))
            parent['table_cells'] = cells; parent['table_headers'] = matrix[0]
            parents.append(parent)
    return parents


def extract_scalars(config):
    path = Path(config['path']); digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != config['sha256']:
        raise RuntimeError('Source version changed: '+str(path))
    if path.suffix.lower() == '.pptx':
        parents = [r for r in pptx_evidence(path) if (r.get('layout') or {}).get('kind') == 'slide_parent']
    elif path.suffix.lower() == '.docx':
        parents = _docx_parents(path,digest)
    elif path.suffix.lower() == '.pdf':
        with pymupdf.open(path) as document:
            parents = [record(path,digest,{'page':i},f'Page {i}',p.get_text()) for i,p in enumerate(document,1)]
    else:
        raise ValueError('Unsupported native scalar input: '+str(path))
    facts = []
    for parent in parents:
        if not _admitted(parent['location'],config):
            continue
        cells = parent.get('table_cells'); headers = parent.get('table_headers')
        if cells:
            if parent['location']['row'] == 1:
                continue
            object_columns = [i for i,h in enumerate(headers) if compact(h) in {'创效方法','优化事项','优化对象','比选项目'}]
            for ci,header in enumerate(headers):
                metric = next((v for v in requested_metrics(header) if v != '提前天数'),None)
                match = re.fullmatch(rf'\s*({VALUE})\s*(亿元|万元|元|W)?\s*',cells[ci])
                if not metric or not match or not object_columns:
                    continue
                loc = {**parent['location'],'column':ci+1,'subject_column':object_columns[0]+1}
                raw_unit = match[2] or next((u for u in ['亿元','万元','元'] if u in header),'')
                confirmed = next((u for u in config.get('unit_confirmations',[]) if u['source_sha256'] == digest and u['location'] == {k:loc[k] for k in ('table','row','column')} and u['raw_unit'] == raw_unit),None)
                facts.append(_fact(parent,config,metric,match[1],raw_unit,subject=compact(cells[object_columns[0]]),raw_metric=header,location=loc,unit_confirmation=confirmed))
            continue
        # shortcut: native labelled amounts and one-line milestone/date/advance groups only; OCR and cross-page binding need a layout adapter.
        nodes = (parent.get('layout') or {}).get('nodes') or [{'text':parent['text']}]
        for node in nodes:
            if node.get('kind') == 'table':
                facts.extend(_key_value_table_facts(parent,config,node))
                continue
            for match in MONEY.finditer(compact(node['text'])):
                metric = ALIASES[match['label']]
                loc = {**parent['location']}
                if node.get('shape_id'):loc.update({'shape_id':node['shape_id'],'bbox':node['bbox']})
                facts.append(_fact(parent,config,metric,match['value'],match['unit'],raw_metric=match['label'],location=loc))
        timeline = list(re.finditer(rf'(?m)^\s*\d+[.、．]\s*(?P<object>[^\n\d]+?)\s+(?P<date>\d{{4}}/\d{{1,2}}/\d{{1,2}})\s*(?:整体)?提前(?P<value>{VALUE})天',parent['text']))
        for match in timeline:
            facts.append(_fact(parent,config,'提前天数',match['value'],'天',subject=compact(match['object']),location={**parent['location'],'milestone':compact(match['object']),'date':match['date']}))
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    return parents, list({r['evidence_id']:r for r in facts}.values())
