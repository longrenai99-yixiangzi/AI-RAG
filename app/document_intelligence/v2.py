from __future__ import annotations

import hashlib
import json
import re
import uuid
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import unquote, urlparse

import pymupdf
from docx import Document as WordDocument
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation

from app.ingestion.metadata.classifier import MetadataClassifier
from app.ingestion.metadata.governance import GovernanceClassifier
from app.ingestion.metadata.schema import MetadataRecord
from app.ingestion.loaders.markdown_loader import MarkdownLoader


SCHEMA_VERSION = "document_intelligence.v2"
SUPPORTED_EXTENSIONS = {".md", ".markdown", ".pdf", ".docx", ".xlsx", ".pptx"}
HEADING_RE = re.compile(r"^(#{1,6})[ \t]+(.+?)\s*$")
YEAR_RE = re.compile(r"20\d{2}")


def stable_id(kind: str, *parts: object) -> str:
    value = "|".join([SCHEMA_VERSION, kind, *(str(part) for part in parts)])
    return str(uuid.uuid5(uuid.NAMESPACE_URL, value))


def document_id(path: Path) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, str(path.resolve()).lower()))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1_048_576), b""):
                digest.update(block)
    except OSError:
        return ""
    return digest.hexdigest()


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def clean_text(value: Any) -> str:
    return " ".join(str(json_value(value) or "").split())


def profile_field(value: Any, source: str, confidence: float) -> dict[str, Any]:
    return {"value": value, "source": source, "confidence": round(float(confidence), 3)}


class DocumentIntelligenceV2Builder:
    """Builds an isolated Document -> Section -> Evidence structure."""

    def __init__(self, root: Path) -> None:
        self.root = root.resolve()
        self.metadata_classifier = MetadataClassifier()
        self.governance_classifier = GovernanceClassifier()

    def build(self, paths: Iterable[Path], atomic_records: list[dict[str, Any]] | None = None) -> dict[str, list[dict[str, Any]]]:
        bundle: dict[str, list[dict[str, Any]]] = {
            "documents": [],
            "headings": [],
            "sections": [],
            "paragraphs": [],
            "tables": [],
            "table_rows": [],
            "lineage": [],
            "metadata_conflicts": [],
            "atomic_evidence": [],
            "legacy_evidence_id_map": [],
        }
        for path in sorted({Path(item) for item in paths}, key=lambda item: str(item).casefold()):
            if path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
                continue
            result = self._build_document(path)
            for key in ("documents", "headings", "sections", "paragraphs", "tables", "table_rows", "lineage", "metadata_conflicts"):
                bundle[key].extend(result[key])
        self._attach_parent_links(bundle, atomic_records or [])
        return bundle

    def _build_document(self, path: Path) -> dict[str, list[dict[str, Any]]]:
        did = document_id(path)
        content_hash = sha256(path)
        parsed = self._empty_result()
        try:
            parsed = getattr(self, f"_parse_{path.suffix.casefold().lstrip('.')}")(path, did)
        except Exception as error:  # Keep a Document Object even when one parser fails.
            parsed["parse_status"] = "read_error"
            parsed["error"] = f"{type(error).__name__}: {error}"

        text = parsed.get("text", "")
        title = parsed.get("title", "")
        metadata = self._classify_metadata_without_rehash(
            path,
            parse_status=str(parsed.get("parse_status") or "read_error"),
            title=title,
            headers=[item["heading_text"] for item in parsed["headings"]],
            text=text[:4_000],
            front_matter=parsed.get("front_matter") or {},
            content_hash=content_hash,
        )
        governance = self.governance_classifier.classify(
            file_name=path.name,
            source_path=str(path),
            heading_path=" > ".join(item["heading_text"] for item in parsed["headings"][:5]),
            text=text,
            metadata=metadata,
        )
        profile = self._build_profile(path, text, title, metadata, governance)
        lineage, lineage_records = self._build_lineage(path, text, did)
        for section in parsed["sections"]:
            section["section_text"] = self._section_text(section, parsed["paragraphs"], parsed["tables"], parsed["table_rows"])
        document = {
            "schema_version": SCHEMA_VERSION,
            "document_id": did,
            "knowledge_root_id": "Root-001" if self._inside_root(path) else "UNKNOWN",
            "source_path": str(path),
            "file_name": path.name,
            "file_type": path.suffix.casefold(),
            "content_hash": content_hash,
            "parse_status": parsed.get("parse_status", "read_error"),
            "error": parsed.get("error"),
            "document_profile": profile,
            "source_lineage": lineage,
            "structure": {
                "title": title,
                "heading_tree": [item["heading_id"] for item in parsed["headings"]],
                "sections": [item["section_id"] for item in parsed["sections"]],
                "tables": [item["table_id"] for item in parsed["tables"]],
                "appendices": [],
                "heading_candidates": parsed.get("heading_candidates", []),
            },
            "metadata_quality": self._metadata_quality(profile, metadata),
        }
        for item in parsed["headings"]:
            item["schema_version"] = SCHEMA_VERSION
            item["document_id"] = did
        for item in parsed["sections"]:
            item["schema_version"] = SCHEMA_VERSION
            item["document_id"] = did
        for item in parsed["paragraphs"]:
            item["schema_version"] = SCHEMA_VERSION
            item["document_id"] = did
        for item in parsed["tables"]:
            item["schema_version"] = SCHEMA_VERSION
            item["document_id"] = did
        for item in parsed["table_rows"]:
            item["schema_version"] = SCHEMA_VERSION
            item["document_id"] = did
        conflicts = self._detect_conflicts(document, path, text)
        return {
            "documents": [document],
            "headings": parsed["headings"],
            "sections": parsed["sections"],
            "paragraphs": parsed["paragraphs"],
            "tables": parsed["tables"],
            "table_rows": parsed["table_rows"],
            "lineage": lineage_records,
            "metadata_conflicts": conflicts,
        }

    @staticmethod
    def _empty_result() -> dict[str, Any]:
        return {
            "parse_status": "read_error",
            "error": None,
            "title": "",
            "text": "",
            "front_matter": {},
            "headings": [],
            "heading_candidates": [],
            "sections": [],
            "paragraphs": [],
            "tables": [],
            "table_rows": [],
        }

    def _classify_metadata_without_rehash(
        self,
        path: Path,
        *,
        parse_status: str,
        title: str,
        headers: list[str],
        text: str,
        front_matter: dict[str, Any],
        content_hash: str,
    ) -> dict[str, Any]:
        evidence = {
            "path": self.metadata_classifier._relative_path(path, self.root),
            "file_name": path.name,
            "title": title,
            "headers": " ".join(headers),
            "body": text,
        }
        values: dict[str, str | None] = {}
        sources: dict[str, str] = {}
        rules: dict[str, str] = {}
        confidence: dict[str, float] = {}
        needs_review = parse_status not in {"parsed", "empty"}
        for field_name, field_rules in self.metadata_classifier.rules.get("fields", {}).items():
            value, source, rule_id, score = self.metadata_classifier._classify_field(field_name, field_rules, evidence, front_matter)
            values[field_name] = value
            sources[field_name] = source
            rules[field_name] = rule_id
            confidence[field_name] = score
            if value is None or score < self.metadata_classifier.auto_min_confidence:
                needs_review = True
        return MetadataRecord(
            file_type=path.suffix.casefold(),
            file_name=path.name,
            source_path=str(path),
            sha256=content_hash,
            parse_status=parse_status,
            board=values.get("board"),
            knowledge_type=values.get("knowledge_type"),
            discipline=values.get("discipline"),
            metadata_source=sources,
            metadata_rule=rules,
            metadata_confidence=confidence,
            metadata_review_status="NEEDS_REVIEW" if needs_review else "AUTO",
        ).to_dict()

    def _parse_md(self, path: Path, did: str) -> dict[str, Any]:
        loaded = MarkdownLoader().load(path, did)
        result = self._empty_result()
        result.update({"parse_status": loaded.status, "error": loaded.error, "front_matter": loaded.front_matter})
        if loaded.status != "parsed":
            return result
        lines = path.read_text(encoding="utf-8-sig").splitlines()
        body_start, front_matter, front_matter_error = MarkdownLoader._parse_front_matter(lines)
        if front_matter_error:
            result.update({"parse_status": "front_matter_error", "error": front_matter_error, "front_matter": front_matter})
            return result
        result["front_matter"] = front_matter
        heading_specs: list[tuple[int, int, str]] = []
        table_ranges: list[tuple[int, int]] = []
        in_fence = False
        index = body_start
        while index < len(lines):
            line = lines[index]
            if line.strip().startswith(("```", "~~~")):
                in_fence = not in_fence
            match = HEADING_RE.match(line) if not in_fence else None
            if match:
                heading_specs.append((index + 1, len(match.group(1)), re.sub(r"[ \t]+#+[ \t]*$", "", match.group(2)).strip()))
            if "|" in line and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
                end = index + 2
                while end < len(lines) and "|" in lines[end] and lines[end].strip():
                    end += 1
                table_ranges.append((index + 1, end))
                index = end
                continue
            index += 1
        headings, sections = _make_heading_sections(did, heading_specs, len(lines) or 1, "markdown")
        result["headings"], result["sections"] = headings, sections
        for line_number, line in enumerate(lines[body_start:], start=body_start + 1):
            if not line.strip() or HEADING_RE.match(line) or _in_ranges(line_number, table_ranges):
                continue
            section = _section_for_location(sections, line_number)
            result["paragraphs"].append({
                "paragraph_id": stable_id("paragraph", did, line_number, line),
                "section_id": section["section_id"],
                "text": line.strip(),
                "order": line_number,
                "location": {"line_start": line_number, "line_end": line_number},
                "style": "markdown",
                "source_block_id": stable_id("source-block", did, line_number),
            })
        for start, end in table_ranges:
            raw = [lines[i - 1] for i in range(start, end)]
            header = _split_pipe(raw[0]) if raw else []
            section = _section_for_location(sections, start)
            table_id = stable_id("table", did, start, end)
            result["tables"].append({
                "table_id": table_id,
                "section_id": section["section_id"],
                "sheet_name": None,
                "table_title": "",
                "header": header,
                "row_count": max(0, len(raw) - 2),
                "column_count": len(header),
                "source_location": {"line_start": start, "line_end": end},
                "semantic": {"title": "", "header": header, "semantic_summary": "Markdown table", "business_fields": header, "representative_rows": []},
                "structured": {"columns": header, "rows": [], "cells": [], "formulas": [], "cached_values": [], "source_row_numbers": list(range(start, end))},
            })
            for offset, row_text in enumerate(raw[2:], start=2):
                values = _split_pipe(row_text)
                result["table_rows"].append(_table_row(table_id, section["section_id"], start + offset, values, header, {"line_start": start + offset, "line_end": start + offset}))
        result["title"] = heading_specs[0][2] if heading_specs else path.stem
        result["text"] = "\n".join(lines)
        return result

    def _parse_docx(self, path: Path, did: str) -> dict[str, Any]:
        result = self._empty_result()
        document = WordDocument(path)
        specs: list[tuple[int, int, str]] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = clean_text(paragraph.text)
            level = _heading_level(paragraph.style.name if paragraph.style else "")
            if text and level:
                specs.append((index, level, text))
        max_location = max(len(document.paragraphs), 1)
        headings, sections = _make_heading_sections(did, specs, max_location, "docx")
        result["headings"], result["sections"] = headings, sections
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = clean_text(paragraph.text)
            if not text or _heading_level(paragraph.style.name if paragraph.style else ""):
                continue
            section = _section_for_location(sections, index)
            result["paragraphs"].append({
                "paragraph_id": stable_id("paragraph", did, index),
                "section_id": section["section_id"],
                "text": text,
                "order": index,
                "location": {"paragraph_start": index, "paragraph_end": index},
                "style": paragraph.style.name if paragraph.style else "Normal",
                "source_block_id": stable_id("source-block", did, index),
            })
        for table_index, table in enumerate(document.tables, start=1):
            values = [[clean_text(cell.text) for cell in row.cells] for row in table.rows]
            if not values:
                continue
            table_text = "\n".join(" | ".join(row) for row in values)
            matching = [section for section in sections if section["heading"] and section["heading"] in table_text]
            section = matching[-1] if matching else sections[-1]
            header = values[0]
            table_id = stable_id("table", did, table_index)
            table_record = {
                "table_id": table_id,
                "section_id": section["section_id"],
                "sheet_name": None,
                "table_title": next((p["text"] for p in result["paragraphs"] if "清单" in p["text"] or "表" in p["text"]), ""),
                "header": header,
                "row_count": len(values),
                "column_count": max((len(row) for row in values), default=0),
                "source_location": {"table": table_index},
                "semantic": {"title": "", "header": header, "semantic_summary": table_text[:500], "business_fields": header, "representative_rows": values[1:3]},
                "structured": {"columns": header, "rows": [], "cells": [], "formulas": [], "cached_values": [], "source_row_numbers": list(range(1, len(values) + 1))},
            }
            result["tables"].append(table_record)
            for row_index, row in enumerate(values, start=1):
                result["table_rows"].append(_table_row(table_id, section["section_id"], row_index, row, header, {"table": table_index, "row": row_index}))
        result["title"] = next((p.text for p in document.paragraphs if clean_text(p.text)), path.stem)
        result["text"] = "\n".join([p["text"] for p in result["paragraphs"]] + [t["semantic"]["semantic_summary"] for t in result["tables"]])
        result["parse_status"] = "parsed"
        return result

    def _parse_pdf(self, path: Path, did: str) -> dict[str, Any]:
        result = self._empty_result()
        with pymupdf.open(path) as pdf:
            text_parts: list[str] = []
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text") or ""
                lines = [line.strip() for line in text.splitlines() if line.strip()]
                text_parts.extend(lines)
                section_id = stable_id("section", did, "page", page_number)
                result["sections"].append({
                    "section_id": section_id,
                    "heading_id": None,
                    "heading": f"Page {page_number}",
                    "heading_path": "",
                    "parent_section_id": None,
                    "depth": 0,
                    "section_order": page_number,
                    "section_text": "",
                    "location_start": {"page": page_number},
                    "location_end": {"page": page_number},
                    "paragraph_ids": [],
                    "table_ids": [],
                    "synthetic": True,
                })
                for line_number, line in enumerate(lines, start=1):
                    paragraph_id = stable_id("paragraph", did, page_number, line_number)
                    result["paragraphs"].append({
                        "paragraph_id": paragraph_id,
                        "section_id": section_id,
                        "text": line,
                        "order": len(result["paragraphs"]) + 1,
                        "location": {"page": page_number, "text_line_start": line_number, "text_line_end": line_number},
                        "style": "pdf_text_block",
                        "source_block_id": stable_id("source-block", did, page_number, line_number),
                    })
                    if _pdf_heading_candidate(line):
                        result["heading_candidates"].append({"text": line, "page": page_number, "line": line_number})
            result["text"] = "\n".join(text_parts)
            result["title"] = next(iter(text_parts), path.stem)
            result["parse_status"] = "parsed" if text_parts else "empty"
        return result

    def _parse_pptx(self, path: Path, did: str) -> dict[str, Any]:
        result = self._empty_result()
        presentation = Presentation(path)
        text_parts: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            texts = []
            title = ""
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = clean_text(shape.text)
                    if value:
                        texts.append(value)
                        if not title and "title" in str(getattr(shape, "name", "")).casefold():
                            title = value
                if getattr(shape, "has_table", False):
                    table_id = stable_id("table", did, slide_number, len(result["tables"]))
                    rows = [[clean_text(cell.text) for cell in row.cells] for row in shape.table.rows]
                    header = rows[0] if rows else []
                    section = _ensure_section(result["sections"], did, slide_number, title or f"Slide {slide_number}")
                    result["tables"].append({
                        "table_id": table_id,
                        "section_id": section["section_id"],
                        "sheet_name": None,
                        "table_title": title,
                        "header": header,
                        "row_count": len(rows),
                        "column_count": len(header),
                        "source_location": {"slide": slide_number},
                        "semantic": {"title": title, "header": header, "semantic_summary": "PPTX table", "business_fields": header, "representative_rows": rows[1:3]},
                        "structured": {"columns": header, "rows": [], "cells": [], "formulas": [], "cached_values": [], "source_row_numbers": list(range(1, len(rows) + 1))},
                    })
                    for row_number, row in enumerate(rows, start=1):
                        result["table_rows"].append(_table_row(table_id, section["section_id"], row_number, row, header, {"slide": slide_number, "row": row_number}))
            section = _ensure_section(result["sections"], did, slide_number, title or f"Slide {slide_number}")
            if title:
                heading_id = stable_id("heading", did, "slide", slide_number)
                result["headings"].append({"heading_id": heading_id, "document_id": did, "heading_text": title, "heading_level": 1, "heading_path": title, "parent_heading_id": None, "order": slide_number, "location": {"slide": slide_number}})
                section["heading_id"] = heading_id
                section["heading_path"] = title
            for text_index, value in enumerate(texts, start=1):
                result["paragraphs"].append({
                    "paragraph_id": stable_id("paragraph", did, slide_number, text_index),
                    "section_id": section["section_id"],
                    "text": value,
                    "order": len(result["paragraphs"]) + 1,
                    "location": {"slide": slide_number, "text_index": text_index},
                    "style": "slide_text",
                    "source_block_id": stable_id("source-block", did, slide_number, text_index),
                })
            text_parts.extend(texts)
        result["title"] = next((h["heading_text"] for h in result["headings"]), path.stem)
        result["text"] = "\n".join(text_parts)
        result["parse_status"] = "parsed" if text_parts or result["tables"] else "empty"
        return result

    def _parse_xlsx(self, path: Path, did: str) -> dict[str, Any]:
        result = self._empty_result()
        formulas = load_workbook(path, data_only=False, read_only=False)
        cached = load_workbook(path, data_only=True, read_only=False)
        text_parts: list[str] = []
        try:
            for sheet in formulas.worksheets:
                value_sheet = cached[sheet.title]
                nonempty = []
                for row_number in range(1, sheet.max_row + 1):
                    cells = [_cell_with_merge(sheet, value_sheet, row_number, column) for column in range(1, sheet.max_column + 1)]
                    if any(cell["display"] not in (None, "") for cell in cells):
                        nonempty.append((row_number, cells))
                heading_id = stable_id("heading", did, "sheet", sheet.title)
                section_id = stable_id("section", did, "sheet", sheet.title)
                result["headings"].append({"heading_id": heading_id, "document_id": did, "heading_text": sheet.title, "heading_level": 1, "heading_path": sheet.title, "parent_heading_id": None, "order": len(result["headings"]) + 1, "location": {"sheet_name": sheet.title}, "synthetic": True})
                result["sections"].append({"section_id": section_id, "heading_id": heading_id, "heading": sheet.title, "heading_path": sheet.title, "parent_section_id": None, "depth": 0, "section_order": len(result["sections"]) + 1, "section_text": "", "location_start": {"sheet_name": sheet.title}, "location_end": {"sheet_name": sheet.title}, "paragraph_ids": [], "table_ids": [], "synthetic": True})
                if not nonempty:
                    continue
                header_number, header_cells = _find_xlsx_header(nonempty)
                headers = [clean_text(cell["display"]) or f"列{index}" for index, cell in enumerate(header_cells, start=1)]
                carry_index = next((index for index, header in enumerate(headers) if any(term in header for term in ("专业", "类别", "单位", "分公司"))), None)
                carry_value = None
                table_id = stable_id("table", did, sheet.title, header_number)
                last_row = nonempty[-1][0]
                table_record = {
                    "table_id": table_id,
                    "section_id": section_id,
                    "sheet_name": sheet.title,
                    "table_title": sheet.title,
                    "header": headers,
                    "row_count": len([item for item in nonempty if item[0] >= header_number]),
                    "column_count": len(headers),
                    "source_location": {"sheet_name": sheet.title, "row_start": header_number, "row_end": last_row, "column_start": 1, "column_end": len(headers)},
                    "semantic": {"title": sheet.title, "header": headers, "semantic_summary": "Workbook sheet/region", "business_fields": headers, "representative_rows": []},
                    "structured": {"columns": headers, "rows": [], "cells": [], "formulas": [], "cached_values": [], "source_row_numbers": []},
                    "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                }
                result["tables"].append(table_record)
                result["sections"][-1]["table_ids"].append(table_id)
                for row_number, cells in nonempty:
                    if row_number < header_number:
                        continue
                    if carry_index is not None:
                        if cells[carry_index]["display"] in (None, "") and carry_value not in (None, ""):
                            cells[carry_index]["display"] = carry_value
                            cells[carry_index]["inherited_from_blank"] = f"row:{row_number - 1}"
                        elif cells[carry_index]["display"] not in (None, ""):
                            carry_value = cells[carry_index]["display"]
                    values = [cell["display"] for cell in cells[: len(headers)]]
                    row_record = _table_row(table_id, section_id, row_number, values, headers, {"sheet_name": sheet.title, "row_start": row_number, "row_end": row_number, "column_count": len(headers), "header_row": header_number})
                    row_record["is_header"] = row_number == header_number
                    row_record["is_summary"] = _is_summary_row(values, cells)
                    row_record["cells"] = [
                        {"column": get_column_letter(index + 1), "header": headers[index], "value": json_value(cell["display"]), "formula": json_value(cell["formula"]), "cached_value": json_value(cell["cached"]), "inherited_from_merged": cell.get("inherited_from_merged"), "inherited_from_blank": cell.get("inherited_from_blank")}
                        for index, cell in enumerate(cells[: len(headers)])
                    ]
                    result["table_rows"].append(row_record)
                    table_record["structured"]["rows"].append(row_record["row_number"])
                    table_record["structured"]["source_row_numbers"].append(row_number)
                    table_record["semantic"]["representative_rows"].append(values[: min(len(values), 5)])
                    text_parts.append(f"工作表：{sheet.title} 第{row_number}行：" + " | ".join(f"{headers[i]}：{clean_text(values[i])}" for i in range(min(len(headers), len(values))) if values[i] not in (None, "")))
            result["title"] = path.stem
            result["text"] = "\n".join(text_parts)
            result["parse_status"] = "parsed" if result["tables"] else "empty"
        finally:
            formulas.close()
            cached.close()
        return result

    def _build_profile(self, path: Path, text: str, title: str, metadata: dict[str, Any], governance: Any) -> dict[str, Any]:
        path_text = str(path)
        years = list(dict.fromkeys(YEAR_RE.findall(f"{path_text} {title} {text[:4_000]}")))
        projects = _project_candidates(path_text, title, text)
        organizations = _organization_candidates(path_text, text)
        specialties = [value for value in (metadata.get("discipline"),) if value]
        profile = {
            "document_type": profile_field(_document_type(path, text, metadata), "RULE+PATH+FILENAME", 0.92),
            "document_role": profile_field(governance.document_role, governance.source, governance.confidence),
            "business_domain": profile_field(metadata.get("board"), metadata.get("metadata_source", {}).get("board", "METADATA_RULE"), _field_confidence(metadata, "board")),
            "organization": profile_field(organizations, "PATH+BODY", 0.88 if organizations else 0.0),
            "project": profile_field(projects, "PATH+FILENAME+BODY", 0.88 if projects else 0.0),
            "specialty": profile_field(specialties, "METADATA+BODY", 0.82 if specialties else 0.0),
            "year": profile_field(years, "PATH+FILENAME+BODY", 0.85 if years else 0.0),
            "version": profile_field(_version_candidates(path_text, text), "FILENAME+BODY", 0.78 if _version_candidates(path_text, text) else 0.0),
            "status": profile_field("parsed", "PARSER", 1.0),
            "authority_level": profile_field(governance.authority_level, "GOVERNANCE_RULE", 0.95),
            "applicable_scope": profile_field({"organization": organizations, "project": projects, "year": years, "specialty": specialties}, "PROFILE_FIELDS", 0.82),
            "entities": profile_field(list(dict.fromkeys(organizations + projects + specialties)), "PATH+BODY+METADATA", 0.8),
            "metrics": profile_field(_metric_candidates(text), "BODY+TABLE_HEADER", 0.75),
            "profile_text": f"{title or path.name}；角色={governance.document_role}；组织={','.join(organizations)}；项目={','.join(projects)}",
        }
        return profile

    @staticmethod
    def _metadata_quality(profile: dict[str, Any], metadata: dict[str, Any]) -> dict[str, Any]:
        fields = ["document_type", "document_role", "business_domain", "organization", "project", "specialty", "year", "version", "authority_level", "applicable_scope"]
        known = sum(1 for field in fields if profile.get(field, {}).get("value") not in (None, [], "", "UNKNOWN"))
        confidences = [float(profile[field].get("confidence", 0.0)) for field in fields if profile.get(field, {}).get("value") not in (None, [], "")]
        return {
            "completeness": round(known / len(fields), 3),
            "confidence": round(sum(confidences) / len(confidences), 3) if confidences else 0.0,
            "source": "rules+metadata+governance",
            "review_status": metadata.get("metadata_review_status", "NEEDS_REVIEW"),
        }

    def _build_lineage(self, path: Path, text: str, did: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        links = _extract_links(text)
        relation_records = []
        for index, target in enumerate(links, start=1):
            resolved = _resolve_link(target, path, self.root)
            relation_records.append({
                "schema_version": SCHEMA_VERSION,
                "lineage_id": stable_id("lineage", did, index, target),
                "source_document_id": did,
                "source_path": str(path),
                "relation": "REGISTER_POINTS_TO" if "wiki" in str(path).casefold() else "DERIVED_FROM",
                "target": target,
                "resolved_path": str(resolved) if resolved else None,
                "lineage_status": "LINEAGE_CONFIRMED" if resolved and resolved.is_file() else "LINEAGE_PARTIAL",
                "evidence": {"link_text": target, "source": "BODY_OR_FRONT_MATTER"},
            })
        status = "LINEAGE_NOT_APPLICABLE" if not relation_records else ("LINEAGE_CONFIRMED" if all(item["lineage_status"] == "LINEAGE_CONFIRMED" for item in relation_records) else "LINEAGE_PARTIAL")
        lineage = {
            "original_source": str(path),
            "derived_from": [],
            "supporting_sources": [],
            "registration_page": str(path) if "\\wiki\\sources\\" in str(path).casefold() else None,
            "supersedes": [],
            "superseded_by": [],
            "lineage_status": status,
            "lineage_evidence": [item["lineage_id"] for item in relation_records],
        }
        return lineage, relation_records

    def _detect_conflicts(self, document: dict[str, Any], path: Path, text: str) -> list[dict[str, Any]]:
        profile = document["document_profile"]
        path_years = set(YEAR_RE.findall(str(path)))
        body_years = set(YEAR_RE.findall(text[:4_000]))
        if path_years and body_years and path_years.isdisjoint(body_years):
            return [{
                "schema_version": SCHEMA_VERSION,
                "conflict_id": stable_id("metadata-conflict", document["document_id"], "year"),
                "document_id": document["document_id"],
                "field": "year",
                "values": {"path": sorted(path_years), "body": sorted(body_years)},
                "status": "METADATA_CONFLICT",
                "resolution": "NEEDS_REVIEW",
            }]
        return []

    @staticmethod
    def _inside_root(path: Path, root: Path | None = None) -> bool:
        try:
            path.resolve().relative_to((root or path.anchor).resolve() if isinstance(root or path.anchor, Path) else Path(root or path.anchor).resolve())
            return True
        except (ValueError, OSError):
            return False

    def _section_text(self, section: dict[str, Any], paragraphs: list[dict[str, Any]], tables: list[dict[str, Any]], rows: list[dict[str, Any]]) -> str:
        values = [item["text"] for item in paragraphs if item.get("section_id") == section["section_id"]]
        values.extend(item.get("semantic", {}).get("semantic_summary", "") for item in tables if item.get("section_id") == section["section_id"])
        section["paragraph_ids"] = [item["paragraph_id"] for item in paragraphs if item.get("section_id") == section["section_id"]]
        section["table_ids"] = [item["table_id"] for item in tables if item.get("section_id") == section["section_id"]]
        return "\n".join(value for value in values if value)

    def _attach_parent_links(self, bundle: dict[str, list[dict[str, Any]]], atomic_records: list[dict[str, Any]]) -> None:
        sections = bundle["sections"]
        paragraphs = bundle["paragraphs"]
        tables = bundle["tables"]
        by_doc = defaultdict(list)
        for section in sections:
            by_doc[section["document_id"]].append(section)
        for record in atomic_records:
            did = str(record.get("document_id") or "")
            candidates = by_doc.get(did, [])
            heading_path = str(record.get("heading_path") or "")
            section = next((item for item in candidates if item.get("heading_path") == heading_path), None)
            if section is None:
                section = next((item for item in candidates if item.get("heading_path") and heading_path.startswith(item["heading_path"])), None)
            if section is None and candidates:
                section = next((item for item in candidates if item.get("synthetic")), candidates[0])
            enriched = dict(record)
            enriched.update({
                "schema_version": SCHEMA_VERSION,
                "section_id": section["section_id"] if section else None,
                "paragraph_id": _parent_paragraph(record, paragraphs),
                "table_id": _parent_table(record, tables),
                "parent_heading_path": section.get("heading_path", "") if section else heading_path,
                "parent_mapping_status": "MAPPED" if section else "UNMAPPED",
            })
            bundle["atomic_evidence"].append(enriched)
            bundle["legacy_evidence_id_map"].append({
                "schema_version": SCHEMA_VERSION,
                "legacy_evidence_id": record.get("evidence_id"),
                "v2_evidence_id": record.get("evidence_id"),
                "status": "UNCHANGED",
            })


def _make_heading_sections(did: str, specs: list[tuple[int, int, str]], max_location: int, kind: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    headings: list[dict[str, Any]] = []
    sections: list[dict[str, Any]] = []
    stack: list[tuple[int, dict[str, Any]]] = []
    for order, (location, level, text) in enumerate(specs, start=1):
        while stack and stack[-1][0] >= level:
            stack.pop()
        parent = stack[-1][1] if stack else None
        path = " > ".join([item["heading_text"] for _, item in stack] + [text])
        heading_id = stable_id("heading", did, kind, order, location, text)
        section_id = stable_id("section", did, kind, order, location, text)
        heading = {
            "heading_id": heading_id,
            "document_id": did,
            "heading_text": text,
            "heading_level": level,
            "heading_path": path,
            "parent_heading_id": parent["heading_id"] if parent else None,
            "order": order,
            "location": {"line_start": location, "line_end": location} if kind == "markdown" else {"paragraph": location},
        }
        section = {
            "section_id": section_id,
            "heading_id": heading_id,
            "heading": text,
            "heading_path": path,
            "parent_section_id": next((item["section_id"] for _, item in reversed(stack) if item.get("heading_level") < level), None),
            "depth": level - 1,
            "section_order": order,
            "section_text": "",
            "location_start": heading["location"],
            "location_end": {},
            "paragraph_ids": [],
            "table_ids": [],
        }
        headings.append(heading)
        sections.append(section)
        stack.append((level, {**heading, "section_id": section_id}))
    for index, section in enumerate(sections):
        next_start = sections[index + 1]["location_start"] if index + 1 < len(sections) else None
        if kind == "markdown":
            end = (next_start.get("line_start", max_location) - 1) if next_start else max_location
            section["location_end"] = {"line_end": max(end, section["location_start"].get("line_start", 1))}
        else:
            end = (next_start.get("paragraph", max_location) - 1) if next_start else max_location
            section["location_end"] = {"paragraph_end": max(end, section["location_start"].get("paragraph", 1))}
    if not sections:
        section_id = stable_id("section", did, kind, "body")
        sections.append({
            "section_id": section_id,
            "heading_id": None,
            "heading": "Document Body",
            "heading_path": "",
            "parent_section_id": None,
            "depth": 0,
            "section_order": 1,
            "section_text": "",
            "location_start": {"start": 1},
            "location_end": {"end": max_location},
            "paragraph_ids": [],
            "table_ids": [],
            "synthetic": True,
        })
    return headings, sections


def _section_for_location(sections: list[dict[str, Any]], location: int) -> dict[str, Any]:
    eligible = []
    for section in sections:
        start = section.get("location_start", {})
        end = section.get("location_end", {})
        start_value = start.get("line_start", start.get("paragraph", start.get("start", 1)))
        end_value = end.get("line_end", end.get("paragraph_end", end.get("end", 10**9)))
        if start_value <= location <= end_value:
            eligible.append(section)
    return eligible[-1] if eligible else sections[0]


def _ensure_section(sections: list[dict[str, Any]], did: str, slide: int, title: str) -> dict[str, Any]:
    for section in sections:
        if section.get("section_order") == slide:
            return section
    section = {
        "section_id": stable_id("section", did, "slide", slide),
        "heading_id": None,
        "heading": title,
        "heading_path": title if title.startswith("Slide ") is False else "",
        "parent_section_id": None,
        "depth": 0,
        "section_order": slide,
        "section_text": "",
        "location_start": {"slide": slide},
        "location_end": {"slide": slide},
        "paragraph_ids": [],
        "table_ids": [],
        "synthetic": True,
    }
    sections.append(section)
    return section


def _heading_level(style: str) -> int | None:
    match = re.search(r"(?:Heading|标题)\s*([1-9])", style or "", re.IGNORECASE)
    return int(match.group(1)) if match else None


def _split_pipe(value: str) -> list[str]:
    value = value.strip().strip("|")
    return [item.strip() for item in value.split("|")]


def _in_ranges(value: int, ranges: list[tuple[int, int]]) -> bool:
    return any(start <= value < end for start, end in ranges)


def _table_row(table_id: str, section_id: str, row_number: int, values: list[Any], headers: list[str], location: dict[str, Any]) -> dict[str, Any]:
    return {
        "table_row_id": stable_id("table-row", table_id, row_number),
        "table_id": table_id,
        "section_id": section_id,
        "row_number": row_number,
        "values": [json_value(value) for value in values],
        "cells": [{"column": index + 1, "header": headers[index] if index < len(headers) else f"列{index + 1}", "value": json_value(value)} for index, value in enumerate(values)],
        "source_location": location,
    }


def _cell_with_merge(formula_sheet: Any, value_sheet: Any, row: int, column: int) -> dict[str, Any]:
    formula = formula_sheet.cell(row, column).value
    cached = value_sheet.cell(row, column).value
    inherited = None
    if formula in (None, "") and cached in (None, ""):
        for merged in formula_sheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                top_formula = formula_sheet.cell(merged.min_row, merged.min_col).value
                top_cached = value_sheet.cell(merged.min_row, merged.min_col).value
                formula, cached = top_formula, top_cached
                inherited = f"{merged}"
                break
    display = cached if isinstance(formula, str) and formula.startswith("=") and cached is not None else formula
    return {"formula": formula, "cached": cached, "display": display, "inherited_from_merged": inherited}


def _find_xlsx_header(rows: list[tuple[int, list[dict[str, Any]]]]) -> tuple[int, list[dict[str, Any]]]:
    preferred_terms = ("专业", "序号", "项目名称", "工作任务", "审查内容", "字段", "名称", "类别")
    for row_number, cells in rows[:20]:
        values = [clean_text(cell["display"]) for cell in cells]
        nonempty = [value for value in values if value]
        if len(nonempty) >= 2 and len(set(nonempty)) >= 2 and any(any(term in value for term in preferred_terms) for value in nonempty):
            return row_number, cells
    return rows[0]


def _is_summary_row(values: list[Any], cells: list[dict[str, Any]]) -> bool:
    text = " ".join(clean_text(value) for value in values)
    return any(term in text for term in ("合计", "总计", "汇总")) or any(isinstance(cell.get("formula"), str) and "SUM" in cell["formula"].upper() for cell in cells)


def _pdf_heading_candidate(value: str) -> bool:
    return len(value) <= 80 and (bool(re.match(r"^\d+(?:\.\d+)*\s*\S+", value)) or value.endswith((":", "：")))


def _field_confidence(metadata: dict[str, Any], field: str) -> float:
    values = metadata.get("metadata_confidence", {}).get(field, 0.0)
    return float(values or 0.0)


def _project_candidates(path_text: str, title: str, text: str) -> list[str]:
    source = f"{path_text} {title} {text[:2_000]}"
    values = []
    for part in re.split(r"[\\/\n，。；：()（）]", source):
        part = part.strip()
        if "项目" in part and 2 <= len(part) <= 80:
            values.append(part.rsplit(".", 1)[0])
    return list(dict.fromkeys(values))[:8]


def _organization_candidates(path_text: str, text: str) -> list[str]:
    values = []
    for term in ("中国建筑第三工程局有限公司", "中建三局第二建设工程有限责任公司", "第二建设公司", "二公司", "局"):
        if term in path_text or term in text[:2_000]:
            values.append(term)
    return values[:5]


def _version_candidates(path_text: str, text: str) -> list[str]:
    return list(dict.fromkeys(re.findall(r"(?:V\d+(?:\.\d+)*|\d+(?:\.\d+)?版|定稿|最终版|试行)", f"{path_text} {text[:2_000]}", flags=re.IGNORECASE)))[:5]


def _metric_candidates(text: str) -> list[str]:
    terms = ("金额", "效益", "利润", "数量", "比例", "率", "成本", "工期", "面积", "上传")
    return [term for term in terms if term in text][:12]


def _document_type(path: Path, text: str, metadata: dict[str, Any]) -> str:
    haystack = f"{path} {path.name} {text[:2_000]}"
    path_text = str(path).replace("/", "\\").casefold()
    if "\\wiki\\sources\\" in path_text or "资料登记" in path.name or "登记页" in path.name:
        return "REGISTER_PAGE"
    if "\\wiki\\queries\\" in path_text:
        return "QUERY_PAGE"
    rules = (
        ("RESPONSIBILITY_CONTRACT", ("责任书", "责任状", "专项责任")),
        ("DESIGN_TASK_BOOK", ("任务书",)),
        ("REVIEW_RECORD", ("评审", "会议纪要", "督办")),
        ("TRAINING", ("培训", "课件", "讲义", "题库")),
        ("TABLE_LEDGER", ("台账", "清单", "统计表")),
        ("RETROSPECTIVE", ("复盘", "经验总结", "年度总结", "半年总结")),
        ("PROJECT_CASE", ("案例", "项目实践", "项目总结")),
        ("PROJECT_PLAN", ("策划", "方案")),
        ("TEMPLATE", ("模板", "范本", "表单", "示范文本")),
        ("WORK_PLAN", ("工作计划", "计划")),
        ("POLICY", ("制度", "办法", "细则", "规定")),
        ("MANAGEMENT_GUIDE", ("指南", "手册", "管理方法", "规范")),
    )
    for document_type, terms in rules:
        if any(term in haystack for term in terms):
            return document_type
    return "OTHER"


def _extract_links(text: str) -> list[str]:
    values = []
    values.extend(match.group(1) for match in re.finditer(r"\[[^\]]*\]\(([^)]+)\)", text))
    values.extend(match.group(1).strip() for match in re.finditer(r"\[\[([^\]|]+)(?:\|[^\]]+)?\]\]", text))
    values.extend(match.group(1) for match in re.finditer(r"(?:file://)?([A-Za-z]:[^\s)\]]+)", text))
    return list(dict.fromkeys(values))


def _resolve_link(target: str, source_path: Path, root: Path) -> Path | None:
    value = unquote(target).strip().replace("/", "\\")
    if value.startswith("file:\\") or value.startswith("file://"):
        value = urlparse(target).path.replace("/", "\\")
        if re.match(r"^\\[A-Za-z]:", value):
            value = value[1:]
    candidate = Path(value)
    if not candidate.is_absolute():
        candidate = source_path.parent / candidate
    try:
        resolved = candidate.resolve()
    except OSError:
        return None
    try:
        resolved.relative_to(root)
    except ValueError:
        return None
    return resolved if resolved.is_file() else None


def _parent_paragraph(record: dict[str, Any], paragraphs: list[dict[str, Any]]) -> str | None:
    location = record.get("location") or {}
    did = record.get("document_id")
    for paragraph in paragraphs:
        if paragraph.get("document_id") != did:
            continue
        p_location = paragraph.get("location") or {}
        if "line_start" in location and p_location.get("line_start") == location.get("line_start"):
            return paragraph["paragraph_id"]
        if "paragraph_start" in location and p_location.get("paragraph_start") == location.get("paragraph_start"):
            return paragraph["paragraph_id"]
        if "page" in location and p_location.get("page") == location.get("page") and p_location.get("text_line_start") == location.get("text_line_start"):
            return paragraph["paragraph_id"]
        if "slide" in location and p_location.get("slide") == location.get("slide") and p_location.get("text_index") == location.get("text_index"):
            return paragraph["paragraph_id"]
    return None


def _parent_table(record: dict[str, Any], tables: list[dict[str, Any]]) -> str | None:
    location = record.get("location") or {}
    did = record.get("document_id")
    for table in tables:
        if table.get("document_id") != did:
            continue
        source = table.get("source_location") or {}
        if location.get("table") == source.get("table"):
            return table["table_id"]
        if location.get("sheet_name") == source.get("sheet_name") and location.get("row_start", 0) >= source.get("row_start", 0) and location.get("row_end", 0) <= source.get("row_end", 10**9):
            return table["table_id"]
        if location.get("slide") == source.get("slide"):
            return table["table_id"]
    return None
