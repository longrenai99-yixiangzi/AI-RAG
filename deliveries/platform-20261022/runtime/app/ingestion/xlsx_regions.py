"""Read version-bound spreadsheet regions into the existing V2 table contract."""
import json
from collections import Counter
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from app.document_intelligence.v2 import _cell_with_merge, _is_summary_row, _table_row, clean_text, document_id, json_value, sha256, stable_id


def _bounds(value):
    bounds = range_boundaries(value)
    left, top, right, bottom = bounds
    if not all(isinstance(item, int) and item > 0 for item in bounds) or right < left or bottom < top or right>16384 or bottom>1048576 or (right-left+1)*(bottom-top+1) > 500_000:
        raise ValueError(f"Invalid or oversized region: {value}")
    return bounds


def _cell(sheet, cached, row, column):
    native = _cell_with_merge(sheet, cached, row, column)
    value = native['display']
    status = 'NATIVE_VALUE'
    if str(native['formula'] or '').startswith('=') and native['cached'] is None:
        status = 'FORMULA_CACHE_MISSING'
    elif str(value) in {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#SPILL!', '#CALC!'}:
        status = 'CELL_ERROR'
    origin = f'{get_column_letter(column)}{row}'
    if native['inherited_from_merged']:
        left, top, _, _ = range_boundaries(native['inherited_from_merged'])
        origin = f'{get_column_letter(left)}{top}'
    number_format=sheet[origin].number_format
    return {'sheet_name':sheet.title,'column': get_column_letter(column), 'row': row, 'value': json_value(value), 'formula': json_value(native['formula']), 'cached_value': json_value(native['cached']), 'inherited_from_merged': native['inherited_from_merged'], 'cell_reference': f"'{sheet.title}'!{get_column_letter(column)}{row}", 'value_origin': f"'{sheet.title}'!{origin}", 'origin_coordinate':origin,'value_status': status,'number_format':number_format,'display_required':isinstance(value,(int,float)) and not isinstance(value,bool) and '%' in number_format}


def extract_xlsx_regions(path: Path, expected_hash: str, definitions: list[dict]) -> dict:
    path = Path(path)
    if len(expected_hash)!=64 or sha256(path) != expected_hash:
        raise ValueError('Source version changed; region coordinates require review')
    formulas = load_workbook(path, data_only=False)
    cached = None
    did = document_id(path)
    tables, rows, contexts, coverage = [], [], [], []
    occupied = {}
    assigned = {}
    try:
        cached = load_workbook(path, data_only=True)
        for spec in definitions:
            sheet = formulas[spec['sheet_name']]
            values = cached[sheet.title]
            bounds = _bounds(spec['range'])
            left, top, right, bottom = bounds
            for old in occupied.setdefault(sheet.title, []):
                if max(left, old[0]) <= min(right, old[2]) and max(top, old[1]) <= min(bottom, old[3]):
                    raise ValueError('Overlapping data regions; preserve context links separately')
            occupied[sheet.title].append(bounds)
            used = assigned.setdefault(sheet.title, set())
            used.update((r,c) for r in range(top,bottom+1) for c in range(left,right+1))
            orientation = spec.get('orientation', 'rows')
            if orientation not in {'rows', 'columns'}:
                raise ValueError('Unknown region orientation')
            section = stable_id('section', did, 'sheet', sheet.title)
            tid = stable_id('region-table', did, expected_hash, json.dumps(spec,sort_keys=True,ensure_ascii=False))
            context_id = stable_id('region-context', tid)
            context_cells = []
            for area in spec.get('context_ranges', []):
                a,b,c,d = _bounds(area)
                context_cells.extend(_cell(sheet,values,r,col) for r in range(b,d+1) for col in range(a,c+1) if sheet.cell(r,col).value not in (None,''))
                used.update((r,col) for r in range(b,d+1) for col in range(a,c+1))
            contexts.append({'context_id': context_id, 'source_path': str(path), 'source_version': expected_hash, 'sheet_name': sheet.title, 'region': spec['range'], 'cells': context_cells, 'review_status': 'PENDING_CONTENT_REVIEW'})
            title = clean_text(sheet[spec['caption_cell']].value) if spec.get('caption_cell') else f"{sheet.title} · {spec['range']}"
            header_refs, groups = [], []
            if orientation == 'rows':
                header_rows = spec['header_rows']
                if not header_rows or sorted(set(header_rows)) != list(range(min(header_rows),max(header_rows)+1)) or min(header_rows)<top or max(header_rows)>bottom:
                    raise ValueError('Header rows must be contiguous and inside the region')
                if any(m.min_row<=max(header_rows) and m.max_row>=min(header_rows) and m.min_col<=right and m.max_col>=left and sheet.cell(m.min_row,m.min_col).value not in (None,'') and not (left<=m.min_col<=m.max_col<=right) for m in sheet.merged_cells.ranges):
                    raise ValueError('Region cuts a merged header')
                labels = []
                for col in range(left,right+1):
                    parts = list(dict.fromkeys(clean_text(_cell_with_merge(sheet,values,r,col)['display']) for r in header_rows if clean_text(_cell_with_merge(sheet,values,r,col)['display'])))
                    labels.append(' > '.join(parts) or f'列{col}')
                    header_refs.append([_cell(sheet,values,r,col)['value_origin'] for r in header_rows])
                for col in range(left,right+1):
                    shared = next((m for m in sheet.merged_cells.ranges if m.min_row<=min(header_rows) and m.max_row>=max(header_rows) and m.min_col<=col<=m.max_col and left<=m.min_col<=m.max_col<=right and len(set(labels[m.min_col-left:m.max_col-left+1]))==1),None)
                    group = list(range(shared.min_col,shared.max_col+1)) if shared else [col]
                    if not groups or groups[-1]!=group:
                        groups.append(group)
                labels = [labels[group[0]-left] for group in groups]
                header_refs = [list(dict.fromkeys(ref for col in group for ref in header_refs[col-left])) for group in groups]
                records = [(r, [_group_cells([_cell(sheet,values,r,c) for c in group], label) for group,label in zip(groups,labels)]) for r in range(max(header_rows)+1,bottom+1)]
            else:
                header_rows = []
                labels = [clean_text(sheet.cell(r,left).value) or f'列{r}' for r in range(top,bottom+1)]
                header_refs = [[_cell(sheet,values,r,left)['cell_reference']] for r in range(top,bottom+1)]
                records = [(c, [dict(_cell(sheet,values,r,c),header=label) for r,label in zip(range(top,bottom+1),labels)]) for c in range(left+1,right+1)]
            table = {'table_id':tid,'document_id':did,'section_id':section,'sheet_name':sheet.title,'sheet_state':sheet.sheet_state,'table_title':title,'header':labels,'header_rows':header_rows,'header_cell_references':header_refs,'context_id':context_id,'orientation':orientation,'source_location':{'sheet_name':sheet.title,'range':spec['range'],'row_start':top,'row_end':bottom,'column_start':left,'column_end':right},'row_count':0,'column_count':len(labels),'semantic':{'title':title,'header':labels},'structured':{'rows':[]},'region_role':spec.get('role','TABLE_CANDIDATE')}
            tables.append(table)
            table['field_columns']=[get_column_letter(group[0]) if len(group)==1 else f'{get_column_letter(group[0])}:{get_column_letter(group[-1])}' for group in groups] if orientation=='rows' else []
            table['field_groups']=spec.get('field_groups') or []
            for ordinal,cells in records:
                if not any(cell['value'] not in (None,'') for cell in cells):
                    continue
                location = dict(table['source_location'])
                if orientation=='rows':
                    location.update(row_start=ordinal,row_end=ordinal)
                else:
                    location.update(column_start=ordinal,column_end=ordinal,record_column=get_column_letter(ordinal))
                row = _table_row(tid,section,ordinal if orientation=='rows' else top,[cell['value'] for cell in cells],labels,location)
                row.update(table_row_id=stable_id('region-record',tid,ordinal),record_number=ordinal,cells=cells,is_header=False,is_summary=_is_summary_row(row['values'],cells),context_evidence_ids=[context_id])
                row_context = [_cell(sheet,values,ordinal,sheet[f'{col}1'].column) for col in spec.get('row_context_columns',[])] if orientation=='rows' else []
                if orientation=='rows':
                    for rule in spec.get('row_context_rules',[]):
                        if rule['rows'][0]<=ordinal<=rule['rows'][1]:
                            a,b,c,d=_bounds(rule['range'])
                            row_context.extend(_cell(sheet,values,r,col) for r in range(b,d+1) for col in range(a,c+1) if sheet.cell(r,col).value not in (None,''))
                    for merged in sheet.merged_cells.ranges:
                        if merged.min_row<=ordinal<=merged.max_row and merged.max_col>merged.min_col and sum(any(merged.min_col<=col<=merged.max_col for col in group) for group in groups)>1:
                            row_context.append(_cell(sheet,values,merged.min_row,merged.min_col))
                            for group,cell in zip(groups,cells):
                                if any(merged.min_col<=col<=merged.max_col for col in group):
                                    cell['value_status']='BODY_MERGE_CROSSES_FIELDS'
                row['row_context'] = [cell for cell in row_context if cell['value'] not in (None,'')]
                for cell in row['row_context']:
                    used.add((cell['row'],sheet[f"{cell['column']}1"].column))
                rows.append(row)
                table['structured']['rows'].append(ordinal)
                table['row_count']+=1
        for name,used in assigned.items():
            populated={(cell.row,cell.column) for cell in formulas[name]._cells.values() if cell.value not in (None,'')}
            coverage.append({'sheet_name':name,'populated_cells':len(populated),'assigned_cells':len(populated&used),'unassigned_cells':[f'{get_column_letter(c)}{r}' for r,c in sorted(populated-used)]})
    finally:
        formulas.close()
        if cached is not None:
            cached.close()
    if sha256(path)!=expected_hash:
        raise ValueError('Source changed during extraction')
    return _split_field_groups({'tables':tables,'table_rows':rows,'context_records':contexts,'coverage':coverage,'source_hash_unchanged':True})


def _split_field_groups(parsed):
    tables,rows=[],[]
    for table in parsed['tables']:
        source_rows=[row for row in parsed['table_rows'] if row['table_id']==table['table_id']]
        groups=table.get('field_groups') or []
        if not groups:
            tables.append(table)
            rows.extend(source_rows)
            continue
        columns=table['field_columns']
        if table['orientation']!='rows' or any(not group or len(group)!=len(set(group)) or not set(group)<=set(columns) for group in groups) or set().union(*map(set,groups))!=set(columns):
            raise ValueError('Field groups must cover the original row fields without cutting merged fields')
        repeated=[header for header,count in Counter(table['header']).items() if count>1]
        for group_index,group in enumerate(groups):
            indices=[columns.index(column) for column in group]
            projected=deepcopy(table)
            projected.update(table_id=stable_id('region-field-group',table['table_id'],','.join(group)),source_table_id=table['table_id'],original_header=table['header'],header=[table['header'][i] for i in indices],header_cell_references=[table['header_cell_references'][i] for i in indices],field_columns=group,field_groups=[],field_group_index=group_index+1,field_group_count=len(groups),requires_column_scope=bool(repeated),table_title=table['table_title']+' · 原列 '+','.join(group))
            projected['column_count']=len(group)
            projected['semantic'].update(title=projected['table_title'],header=projected['header'])
            projected['source_location']['field_columns']=group
            tables.append(projected)
            for original in source_rows:
                row=deepcopy(original)
                row.update(table_id=projected['table_id'],table_row_id=stable_id('field-group-row',original['table_row_id'],projected['table_id']),source_table_row_id=original['table_row_id'],values=[original['values'][i] for i in indices],cells=[deepcopy(original['cells'][i]) for i in indices])
                row['source_location']['field_columns']=group
                rows.append(row)
    return dict(parsed,tables=tables,table_rows=rows)


def _group_cells(cells, label):
    if len(cells)==1:
        return dict(cells[0],header=label)
    members=[cell for cell in cells if cell['value'] not in (None,'')]
    origins={cell['value_origin']:cell for cell in members}
    return {'column':f"{cells[0]['column']}:{cells[-1]['column']}",'header':label,'value':' / '.join(str(cell['value']) for cell in origins.values()) if members else None,'formula':None,'cached_value':None,'native_cells':cells,'cell_reference':f"{cells[0]['cell_reference']}:{cells[-1]['column']}{cells[-1]['row']}",'value_status':next((cell['value_status'] for cell in members if cell['value_status']!='NATIVE_VALUE'),'NATIVE_VALUE'),'display_required':any(cell['display_required'] for cell in members)}


def apply_native_display(cell,lookup):
    if cell.get('native_cells'):
        for member in cell['native_cells']:
            apply_native_display(member,lookup)
        members={member['value_origin']:member for member in cell['native_cells'] if member['value'] not in (None,'')}
        if cell.get('display_required') and all(not member.get('display_required') or member.get('display_text') for member in members.values()):
            cell['display_text']=' / '.join(str(member.get('display_text',member['value'])) for member in members.values())
    elif cell.get('display_required') and cell['value_status']=='NATIVE_VALUE':
        item=lookup.get(f"/{cell['sheet_name']}/{cell['origin_coordinate']}",{})
        properties=item.get('format',{})
        basis=properties.get('cachedValue',properties.get('value'))
        try:
            computed=properties.get('computedValue')
            compatible=basis is not None and (computed is None or abs(Decimal(str(computed))-Decimal(str(basis)))<=Decimal('1e-12'))
            if compatible and Decimal(str(cell['value']))==Decimal(str(basis)):
                cell.update(display_text=item['text'],display_source='OfficeCLI_NATIVE_FORMAT',display_source_value=str(basis))
        except InvalidOperation:
            pass
