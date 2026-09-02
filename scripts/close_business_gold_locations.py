from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "evaluation" / "business_gold_v2"
BA_DIR = PROJECT_ROOT / "evaluation" / "v1_business_acceptance"

BA007_REQUESTED_PATH = Path(
    r"D:\工作\二公司技术部\2026\示范，工程\关于发布2026年公司设计示范、深化设计示范、BIM示范工程（第一批）计划的通知.docx"
)
BA007_RESOLVED_PATH = Path(
    r"D:\工作\二公司技术部\2026\示范工程\关于发布2026年公司设计示范、深化设计示范、BIM示范工程（第一批）计划的通知.docx"
)

OWNER_SOURCES: dict[str, dict[str, Any]] = {
    "BA-001": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\工作\二公司技术部\2026\各类文件\设计管理\《项目设计管理手册》.pdf",
        "source_mode": "FROZEN_SHADOW_ARTIFACT",
        "governance": "PENDING_APPROVAL",
        "artifact": "evaluation/policy_facet_grounding/BA-001.json",
    },
    "BA-002": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\工作\二公司技术部\2026\各类文件\设计管理\《项目设计管理手册》.pdf",
        "source_mode": "FROZEN_SHADOW_ARTIFACT",
        "governance": "PENDING_APPROVAL",
        "artifact": "evaluation/ba002_formula_semantic_alignment/BA-002.json",
    },
    "BA-003": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\工作\二公司技术部\2026\知识库\《产品线（医疗、学校、厂房）方案比选案例库》\厂房\中建三局二公司产品线设计方案比选典型案例汇编（厂房）（正文）.docx",
        "source_mode": "TARGETED_READ_ONLY",
        "governance": "OUT_OF_SCOPE",
        "artifact": "data/shadow/source_discovery/source_records.jsonl",
    },
    "BA-004": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\设计管理\raw\设计支持\方案比选\方案比选库\设计方案比选提示清单7.23.xlsx",
        "source_mode": "TARGETED_READ_ONLY",
        "governance": "APPROVED",
        "artifact": "evaluation/ba004_answer_structure_hardening/run_01.json",
    },
    "BA-005": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\工作\二公司技术部\2026\知识库\全专业施工图审核要点提示汇编\全专业施工图审核要点提示汇编（2026年）.xlsx",
        "source_mode": "TARGETED_READ_ONLY",
        "governance": "PENDING_APPROVAL",
        "artifact": "evaluation/claim_preflight_positive_path/BA-005.json",
    },
    "BA-006": {
        "source_role": "SUPPORTING_SOURCE",
        "source_path": r"D:\工作\二公司技术部\2026\各类文件\设计管理\关于印发中建三局2026年设计与技术工作计划的通知.pdf",
        "source_mode": "FROZEN_SHADOW_ARTIFACT",
        "governance": "PENDING_APPROVAL",
        "artifact": "evaluation/p0_integrated_shadow_regression/BA-006.json",
        "owner_confirmed_source_label": "2026年设计与技术系统责任状正文/目标数量资料",
        "source_fidelity_status": "UNRESOLVED",
        "source_fidelity_note": "当前实际定位到的是工作计划 PDF，只能作为 Supporting Source；责任状正文尚未定位。",
    },
    "BA-007": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "requested_path": str(BA007_REQUESTED_PATH),
        "source_path": str(BA007_RESOLVED_PATH),
        "source_mode": "TARGETED_READ_ONLY",
        "governance": "PENDING_APPROVAL",
        "artifact": "evaluation/policy_local_grounding_window/BA-007.json",
        "path_note": "任务给出的“示范，工程”不存在；在指定的2026目录下定向解析到“示范工程”中的同名 DOCX。",
    },
    "BA-008": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\设计管理\raw\工作总结\2025年饶淇述职.md",
        "source_mode": "TARGETED_READ_ONLY",
        "governance": "APPROVED",
        "artifact": "evaluation/direct_fact_scope_guard_hardening/BA-008.json",
    },
    "BA-009": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\工作\二公司技术部\2026\EPC项目双周推进会\4月\EPC项目设计管理工作监督任务表（2026年4月第一周）.xlsx",
        "source_mode": "TARGETED_READ_ONLY",
        "governance": "PENDING_APPROVAL",
        "artifact": "evaluation/claim_preflight_positive_path/BA-009.json",
    },
    "BA-010": {
        "source_role": "PRIMARY_GOLD_CANDIDATE",
        "source_path": r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷\设计管理策划书-星谷科创中心项目.docx",
        "source_mode": "TARGETED_READ_ONLY",
        "governance": "PENDING_APPROVAL",
        "artifact": "evaluation/claim_preflight_positive_path/BA-010.json",
        "supporting_structured_source": r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷\星谷科创项目设计管理策划+设计示范项目打造方案\方案比选与价值创造清单方案比选及价值创造.xlsx",
        "supporting_structured_source_note": "沿用既有冻结结构化结果，仅作为 DOCX 表格关系的辅助对账，不直接继承旧统计数字。",
    },
}


def main() -> int:
    parser = argparse.ArgumentParser(description="Close owner-confirmed Business Gold locations from ten directed sources only.")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()

    questions = _load_questions()
    source_records: list[dict[str, Any]] = []
    location_records: list[dict[str, Any]] = []
    claim_records: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []
    for question_id in sorted(OWNER_SOURCES):
        source = dict(OWNER_SOURCES[question_id])
        source.update(_file_fingerprint(Path(source["source_path"])))
        source["question_id"] = question_id
        source["source_owner_confirmation"] = source.get("source_owner_confirmation", "CONFIRMED")
        source["source_governance"] = source["governance"]
        source_records.append(source)
        location, claims = _close_one(question_id, questions[question_id], source)
        location_records.append(location)
        claim_records.append({"question_id": question_id, "claims": claims})
        if location["gold_location_status"] != "CONFIRMED":
            unresolved.append(
                {
                    "question_id": question_id,
                    "gold_location_status": location["gold_location_status"],
                    "reason": location.get("unresolved_reason"),
                    "source": source,
                }
            )

    args.output_dir.mkdir(parents=True, exist_ok=True)
    _write_json(args.output_dir / "owner_confirmed_source_manifest.json", {"records": source_records})
    _write_json(args.output_dir / "gold_location_manifest.json", {"records": location_records})
    _write_json(args.output_dir / "gold_claim_manifest.json", {"records": claim_records})
    _write_json(args.output_dir / "unresolved_gold_items.json", {"records": unresolved})
    _write_json(
        args.output_dir / "gold_review_queue.json",
        {
            "schema_version": "gold_review_queue.v2.location_closure",
            "status": "PENDING_OWNER_FINAL_GOLD_CONFIRMATION",
            "items": [_review_item(location, claims) for location, claims in zip(location_records, claim_records, strict=True)],
        },
    )
    (PROJECT_ROOT / "docs" / "BUSINESS_GOLD_LOCATION_CLOSURE_REPORT.md").write_text(
        _render_report(source_records, location_records, claim_records, unresolved), encoding="utf-8"
    )
    (PROJECT_ROOT / "docs" / "BUSINESS_GOLD_OWNER_FINAL_REVIEW_SHEET.md").write_text(
        _render_review_sheet(location_records, claim_records), encoding="utf-8"
    )
    _update_v2_report(source_records, location_records, unresolved)
    print(
        json.dumps(
            {
                "sources_confirmed": sum(item["source_owner_confirmation"] == "CONFIRMED" for item in source_records),
                "location_confirmed": sum(item["gold_location_status"] == "CONFIRMED" for item in location_records),
                "location_partial": sum(item["gold_location_status"] == "PARTIAL" for item in location_records),
                "location_unresolved": sum(item["gold_location_status"] == "UNRESOLVED" for item in location_records),
                "location_pending": len(unresolved),
                "owner_final_confirmation": 0,
                "output_dir": str(args.output_dir.resolve()),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


def _load_questions() -> dict[str, str]:
    return {
        path.stem: str((_read_json(path)).get("question") or "")
        for path in sorted(BA_DIR.glob("BA-*.json"))
    }


def _close_one(question_id: str, question: str, source: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    if question_id == "BA-001":
        location = {
            "source_path": source["source_path"],
            "location_type": "page_and_section",
            "locations": [{"page": 17, "section": "9 设计任务书管理 / 9.1 设计任务书编制"}],
            "source_excerpt": "设计任务书包含项目概况、工作范围、工作要求、设计技术要点，并明确设计依据及管控要求。",
            "expected_scope": "通用项目设计任务书管理",
            "expected_authority": "L2 管理手册",
            "expected_answer_type": "TEMPLATE_QUERY",
            "expected_fact_mode": "NONE",
            "gold_location_status": "CONFIRMED",
        }
        claims = [
            _claim("C1", "设计任务书包含项目概况、工作范围、工作要求、设计技术要点。", "S3", location["locations"]),
            _claim("C2", "设计任务书应明确设计依据及管控要求。", "S3", location["locations"]),
        ]
    elif question_id == "BA-002":
        location = {
            "source_path": source["source_path"],
            "location_type": "page_and_section",
            "locations": [{"page": 30, "section": "17.3 设计创效管理 / 17.3.2 设计创效计算"}],
            "source_excerpt": "项目设计创效经济效益额=实施后实际取得的经济效益额（不包括工期效益）—实施前预期取得的经济效益额；设计创效率=总创效金额/自施产值×100%。",
            "expected_scope": "设计创效正式计算口径",
            "expected_authority": "L2 管理手册",
            "expected_answer_type": "METHOD_QUERY",
            "expected_fact_mode": "DIRECT_FACT",
            "gold_location_status": "CONFIRMED",
            "semantic_boundary": "设计效益增量与上述正式指标的完全同义关系未确认",
        }
        claims = [
            _claim("C1", "项目设计创效经济效益额按实施后实际经济效益额减去实施前预期经济效益额确定，且不包括工期效益。", "S3", location["locations"]),
            _claim("C2", "设计创效率=总创效金额/自施产值×100%。", "S3", location["locations"]),
            _claim("C3", "“设计效益增量”与正式指标的术语映射未被证据确认。", "S3", location["locations"]),
        ]
    elif question_id == "BA-003":
        extracted = _docx_keyword_extract(Path(source["source_path"]), ["专业", "厂房", "方案比选", "案例"])
        claims = _claims_from_profession_rows(extracted, source["source_path"])
        location = {
            "source_path": source["source_path"],
            "location_type": "docx_heading_and_table_or_paragraph",
            "locations": extracted["locations"],
            "source_excerpt": extracted["excerpt"],
            "expected_scope": "厂房产品线方案比选案例",
            "expected_authority": "厂房产品线案例正文；Root-002 当前治理范围外",
            "expected_answer_type": "MULTI_FACT",
            "expected_fact_mode": "DIRECT_FACT",
            "gold_location_status": "CONFIRMED" if extracted["locations"] else "UNRESOLVED",
            "unresolved_reason": None if extracted["locations"] else "厂房正文未找到可定位的专业目录或案例专业字段",
            "truth_answerability": "ANSWERABLE",
            "runtime_expected_behavior": "SOURCE_SCOPE_MISSING",
        }
    elif question_id == "BA-004":
        row = _xlsx_rows(Path(source["source_path"]), ["自动喷淋", "管材"])
        matched = [item for item in row if "自动喷淋" in item.get("text", "") and "管材" in item.get("text", "")][:1]
        claims = [
            _claim("C1", "自动喷淋系统采用传统镀锌钢管。", "S2", matched),
            _claim("C2", "自动喷淋系统采用新型材料加强氯化聚氯乙烯(PVC-C)管材。", "S2", matched),
        ]
        location = {
            "source_path": source["source_path"],
            "location_type": "xlsx_sheet_row_columns",
            "locations": matched,
            "source_excerpt": _excerpt_from_rows(matched),
            "expected_scope": "自动喷淋系统管材变更方案比选",
            "expected_authority": "L3 方案比选清单",
            "expected_answer_type": "OPTION_QUERY",
            "expected_fact_mode": "DIRECT_FACT",
            "gold_location_status": "CONFIRMED" if matched else "UNRESOLVED",
            "unresolved_reason": None if matched else "未找到自动喷淋/管材匹配行",
        }
    elif question_id == "BA-005":
        rows = _xlsx_section_rows(Path(source["source_path"]), "特殊环境条件下集电线路电气设计")
        claims = _claims_from_xlsx_rows(rows, prefix="BA005-C")
        location = {
            "source_path": source["source_path"],
            "location_type": "xlsx_sheet_row_columns",
            "locations": rows,
            "source_excerpt": _excerpt_from_rows(rows[:8]),
            "expected_scope": "特殊环境条件下集电线路电气设计图审",
            "expected_authority": "2026年全专业施工图审核要点提示汇编",
            "expected_answer_type": "DISCIPLINE_QUERY",
            "expected_fact_mode": "NONE",
            "gold_location_status": "CONFIRMED" if rows else "UNRESOLVED",
            "unresolved_reason": None if rows else "指定 XLSX 未找到特殊环境/集电线路/电气图审相关行",
        }
    elif question_id == "BA-006":
        location = {
            "source_path": source["source_path"],
            "location_type": "frozen_pdf_page",
            "locations": [{"page": 5, "field": "电子图形文件数据中心上传数量", "organization": "局层面", "year": "2026"}],
            "source_excerpt": "12月份上传不少于500个项目电子图形文件数据，沉淀设计数据资产。",
            "expected_scope": "局级设计与技术系统，2026年度",
            "expected_authority": "责任状正文待定位；当前工作计划仅为 Supporting Source",
            "expected_answer_type": "DIRECT_FACT",
            "expected_fact_mode": "DIRECT_FACT",
            "gold_location_status": "PARTIAL",
            "truth_answerability": "EVIDENCE_INSUFFICIENT",
            "runtime_expected_behavior": "PARTIAL_EVIDENCE/NO_EVIDENCE，直到责任状正文被定位",
            "source_fidelity_status": "UNRESOLVED",
            "owner_confirmed_source_label": source.get("owner_confirmed_source_label"),
            "source_semantic_note": "工作计划中的500项目要求与负责人确认的责任状来源不能自动等价；本条仅登记为 Supporting Source，不生成 Gold Claim。",
            "supporting_source_path": source["source_path"],
            "supporting_source_role": "SUPPORTING_SOURCE",
            "unresolved_reason": "尚未定位到业务负责人所指的2026年设计与技术系统责任状正文；当前工作计划仅能证明存在相关500项目上传要求。",
        }
        claims = []
    elif question_id == "BA-007":
        extracted = _docx_policy_extract(Path(source["source_path"]))
        claims = _claims_from_policy_extract(extracted)
        location = {
            "source_path": source["source_path"],
            "location_type": "docx_paragraph_and_table",
            "locations": extracted["locations"],
            "source_excerpt": extracted["excerpt"],
            "expected_scope": "2026年公司设计示范工程实施要求",
            "expected_authority": "Owner Confirmed DOCX；Root-002 governance=PENDING_APPROVAL",
            "expected_answer_type": "POLICY_QUERY",
            "expected_fact_mode": "DIRECT_FACT",
            "gold_location_status": "CONFIRMED" if extracted["locations"] else "UNRESOLVED",
            "unresolved_reason": None if extracted["locations"] else "指定示范工程通知中未找到设计示范工程实施要求定位",
            "related_context_locations": extracted.get("related_context", []),
            "gold_claim_scope_note": "仅“一、设计示范工程实施要求”进入 Gold Claims；二、深化设计示范工程、三、BIM示范工程仅保留为 Related Context。",
            "supporting_sources": [
                {"source_role": "SUPPORTING_SOURCE", "file_name": "关于印发中建三局2026年设计与技术工作计划的通知.pdf", "governance": "PENDING_APPROVAL"},
                {"source_role": "SUPPORTING_SOURCE", "file_name": "关于印发中建三局第二建设公司2026年设计与技术（科技）工作计划的通知.pdf", "governance": "PENDING_APPROVAL"},
            ],
        }
    elif question_id == "BA-008":
        lines = _markdown_lines(Path(source["source_path"]), ["4.45", "4．45", "创效金额"])
        claims = [_claim("C1", "2025年公司设计创效金额约为4.45亿元。", "S1", lines)]
        location = {
            "source_path": source["source_path"],
            "location_type": "markdown_line",
            "locations": lines,
            "source_excerpt": _excerpt_from_rows(lines),
            "expected_scope": "公司级2025年全年设计创效金额",
            "expected_authority": "公司级年度述职/总结来源",
            "expected_answer_type": "DIRECT_FACT",
            "expected_fact_mode": "DIRECT_FACT",
            "gold_location_status": "CONFIRMED" if lines else "UNRESOLVED",
            "unresolved_reason": None if lines else "未找到4.45亿元或创效金额原文行",
        }
    elif question_id == "BA-009":
        rows = [
            row
            for row in _xlsx_rows(Path(source["source_path"]), ["土木公司"])
            if row["sheet_name"] == "Sheet1 (2)"
        ]
        claims = _claims_from_xlsx_rows(rows, prefix="BA009-C")
        location = {
            "source_path": source["source_path"],
            "location_type": "xlsx_sheet_row_columns",
            "locations": rows,
            "source_excerpt": _excerpt_from_rows(rows[:8]),
            "expected_scope": "2026年4月EPC项目双周推进会，土木公司",
            "expected_authority": "EPC项目设计管理工作监督任务表正文",
            "expected_answer_type": "DIRECT_FACT",
            "expected_fact_mode": "DIRECT_FACT",
            "gold_location_status": "CONFIRMED" if rows else "UNRESOLVED",
            "unresolved_reason": None if rows else "指定监督任务表未找到土木公司对应督办行",
        }
    elif question_id == "BA-010":
        extracted = _docx_table_extract(Path(source["source_path"]), ["设计价值创造", "专业类别", "利润"])
        claims, fact = _claims_from_ba010_docx(extracted)
        location = {
            "source_path": source["source_path"],
            "location_type": "docx_heading_and_table",
            "locations": extracted["locations"],
            "source_excerpt": extracted["excerpt"],
            "expected_scope": "星谷科创中心项目设计价值创造策划",
            "expected_authority": "Owner Confirmed DOCX；Root-002 governance=PENDING_APPROVAL",
            "expected_answer_type": "AGGREGATION_QUERY",
            "expected_fact_mode": "DERIVED_FACT",
            "gold_location_status": "CONFIRMED" if fact.get("resolved") else "PARTIAL",
            "unresolved_reason": None if fact.get("resolved") else fact.get("reason"),
            "truth_answerability": "ANSWERABLE" if fact.get("resolved") else "EVIDENCE_INSUFFICIENT",
            "runtime_expected_behavior": "FACT_ANSWER_PATH；Root-002 governance=PENDING_APPROVAL" if fact.get("resolved") else "EVIDENCE_INSUFFICIENT；不得继承旧XLSX统计结果",
            "fact_reconciliation": fact,
            "supporting_sources": [
                {
                    "source_role": "STRUCTURED_DERIVED_SOURCE_RELATION_UNCONFIRMED",
                    "source_path": source.get("supporting_structured_source"),
                    "relation_status": "NOT_PROVEN",
                    "note": "仅作为既有结构化 Artifact 辅助对账，不作为 DOCX 事实来源，不继承旧统计数字。",
                }
            ],
        }
    else:
        raise KeyError(question_id)
    location.update(
        {
            "question_id": question_id,
            "question": question,
            "owner_confirmed_source": source["source_path"],
            "owner_confirmed_source_label": source.get("owner_confirmed_source_label"),
            "source_owner_confirmation": "CONFIRMED",
            "source_role": source["source_role"],
            "source_governance": source["governance"],
            "source_fidelity_status": location.get("source_fidelity_status", source.get("source_fidelity_status", "CONFIRMED")),
            "source_fidelity_note": location.get("source_fidelity_note", source.get("source_fidelity_note")),
            "owner_confirmation": "UNCONFIRMED",
            "gold_status": "GOLD_LOCATION_PENDING",
            "expected_location": location["locations"],
            "expected_scope": location.get("expected_scope"),
            "expected_authority": location.get("expected_authority"),
            "expected_answer_type": location.get("expected_answer_type"),
            "expected_fact_mode": location.get("expected_fact_mode"),
            "candidate_expected_subquestions": _subquestions(question_id),
            "expected_subquestions": _subquestions(question_id),
            "candidate_expected_claims": [claim["text"] for claim in claims],
            "expected_claims": [claim["text"] for claim in claims],
            "must_include_evidence": [claim["evidence_id"] for claim in claims],
            "must_not_use_evidence": _must_not_use(question_id),
            "truth_answerability": location.get("truth_answerability", "ANSWERABLE"),
            "runtime_expected_behavior": location.get("runtime_expected_behavior", "ANSWERABLE_IN_SHADOW_ONLY"),
        }
    )
    return location, claims


def _file_fingerprint(path: Path) -> dict[str, Any]:
    result = {"exists": path.is_file(), "size": None, "mtime": None, "sha256": None}
    if not path.is_file():
        return result
    stat = path.stat()
    result["size"] = stat.st_size
    result["mtime"] = stat.st_mtime
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1_048_576), b""):
            digest.update(block)
    result["sha256"] = digest.hexdigest()
    return result


def _xlsx_rows(path: Path, keywords: list[str]) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    workbook = load_workbook(path, read_only=True, data_only=False)
    result: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            header: list[str] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [_cell_text(value) for value in row]
                text = " | ".join(f"列{index + 1}:{value}" for index, value in enumerate(values) if value)
                if not text:
                    continue
                if (
                    not header
                    and row_number <= 20
                    and len([value for value in values if value]) >= 2
                    and any(term in value for value in values for term in ("类别", "审查内容", "专业", "项目", "监督任务", "单位", "工作任务", "完成时限", "责任人"))
                ):
                    header = values
                if any(keyword.casefold() in text.casefold() for keyword in keywords):
                    result.append(
                        {
                            "sheet_name": sheet.title,
                            "row_start": row_number,
                            "row_end": row_number,
                            "header": header,
                            "values": values,
                            "text": text,
                        }
                    )
    finally:
        workbook.close()
    return result


def _xlsx_section_rows(path: Path, start_keyword: str) -> list[dict[str, Any]]:
    """Read one titled section only; never expand to a workbook-wide candidate set."""
    rows = _xlsx_rows(path, ["新能源-线路电气"])
    start_index = next(
        (index for index, row in enumerate(rows) if start_keyword in row.get("text", "")),
        None,
    )
    if start_index is None:
        return []
    selected: list[dict[str, Any]] = []
    for row in rows[start_index:]:
        if selected and row["text"].startswith("列1:"):
            break
        selected.append(row)
    return selected


def _markdown_lines(path: Path, keywords: list[str]) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    return [
        {"line_start": index, "line_end": index, "text": line.strip()}
        for index, line in enumerate(lines, start=1)
        if line.strip() and any(keyword.casefold() in line.casefold() for keyword in keywords)
    ]


def _docx_keyword_extract(path: Path, keywords: list[str]) -> dict[str, Any]:
    if not path.is_file():
        return {"locations": [], "excerpt": "", "paragraphs": [], "tables": []}
    document = WordDocument(path)
    locations: list[dict[str, Any]] = []
    paragraphs: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = " ".join(paragraph.text.split())
        if text and any(keyword.casefold() in text.casefold() for keyword in keywords):
            item = {"paragraph": index, "style": paragraph.style.name, "text": text}
            paragraphs.append(item)
            locations.append(item)
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            text = " | ".join(" ".join(cell.text.split()) for cell in row.cells)
            if text and any(keyword.casefold() in text.casefold() for keyword in keywords):
                item = {"table": table_index, "row": row_index, "text": text}
                tables.append(item)
                locations.append(item)
    return {"locations": locations, "paragraphs": paragraphs, "tables": tables, "excerpt": "\n".join(_location_text(item) for item in locations[:12])}


def _docx_policy_extract(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {"locations": [], "related_context": [], "paragraphs": [], "tables": [], "excerpt": ""}
    document = WordDocument(path)
    primary_locations: list[dict[str, Any]] = []
    related_context: list[dict[str, Any]] = []
    current_section = ""
    primary_section = "一、设计示范工程实施要求"
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if text.startswith(("一、", "二、", "三、")) and "示范工程实施要求" in text:
            current_section = text
            item = {"paragraph": index, "style": paragraph.style.name, "section": current_section, "text": text, "kind": "heading"}
            (primary_locations if current_section == primary_section else related_context).append(item)
            continue
        if current_section:
            item = {"paragraph": index, "style": paragraph.style.name, "section": current_section, "text": text, "kind": "requirement"}
            (primary_locations if current_section == primary_section else related_context).append(item)
    return {
        "locations": primary_locations,
        "related_context": related_context,
        "paragraphs": primary_locations,
        "tables": [],
        "excerpt": "\n".join(_location_text(item) for item in primary_locations),
    }


def _docx_table_extract(path: Path, keywords: list[str]) -> dict[str, Any]:
    if not path.is_file():
        return {"locations": [], "excerpt": "", "tables": []}
    document = WordDocument(path)
    locations: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        text = " ".join(paragraph.text.split())
        if text and "设计价值创造" in text:
            item = {"paragraph": paragraph_index, "style": paragraph.style.name, "text": text}
            locations.append(item)
    for table_index, table in enumerate(document.tables, start=1):
        rows: list[dict[str, Any]] = []
        table_text = "\n".join(" | ".join(" ".join(cell.text.split()) for cell in row.cells) for row in table.rows)
        if not ("设计价值创造" in table_text and "专业类别" in table_text):
            continue
        for row_index, row in enumerate(table.rows, start=1):
            values = [" ".join(cell.text.split()) for cell in row.cells]
            text = " | ".join(values)
            if text:
                item = {"table": table_index, "row": row_index, "values": values, "text": text}
                rows.append(item)
                if any(keyword in text for keyword in keywords):
                    locations.append(item)
        tables.append({"table": table_index, "rows": rows})
    return {"locations": locations, "tables": tables, "excerpt": "\n".join(_location_text(item) for item in locations[:12])}


def _claims_from_profession_rows(extracted: dict[str, Any], source_path: str) -> list[dict[str, Any]]:
    claims: list[dict[str, Any]] = []
    professions: list[str] = []
    for item in extracted.get("paragraphs", []):
        match = re.search(r"【专业类别】\s*([^】]+)", str(item.get("text") or ""))
        if match:
            professions.append(match.group(1).strip())
    for item in extracted.get("tables", []):
        values = item.get("values") or []
        for value in values:
            candidate = value.strip()
            if candidate and len(candidate) <= 12 and any(token in candidate for token in ("建筑", "结构", "电气", "暖通", "给排水", "机电", "总图", "幕墙", "景观", "室内")):
                professions.append(candidate)
    professions = list(dict.fromkeys(professions))
    if professions:
        claims.append(_claim("C1", "厂房产品线方案比选案例覆盖专业：" + "、".join(professions) + "。", "SOURCE_DOCX", extracted.get("locations", [])))
    return claims


def _claims_from_policy_extract(extracted: dict[str, Any]) -> list[dict[str, Any]]:
    claims: list[dict[str, Any]] = []
    relevant = [
        item
        for item in extracted.get("locations", [])
        if item.get("kind") == "requirement"
        and item.get("section") == "一、设计示范工程实施要求"
    ]
    for index, item in enumerate(relevant, start=1):
        text = str(item.get("text") or "")
        if text:
            claims.append(_claim(f"C{index}", text, "SOURCE_DOCX", [item]))
    return claims


def _claims_from_xlsx_rows(rows: list[dict[str, Any]], prefix: str) -> list[dict[str, Any]]:
    return [_claim(f"{prefix}-{index}", row["text"], "SOURCE_XLSX", [row]) for index, row in enumerate(rows, start=1)]


def _claims_from_ba010_docx(extracted: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    candidate_tables = [table for table in extracted.get("tables", []) if _table_has_header(table, "专业类别")]
    if not candidate_tables:
        return [], {"resolved": False, "reason": "DOCX中未找到包含“专业类别”的设计价值创造表格。"}
    table = candidate_tables[0]
    rows = table.get("rows", [])
    header_index = next((index for index, row in enumerate(rows) if "专业类别" in " | ".join(row.get("values") or [])), None)
    if header_index is None:
        return [], {"resolved": False, "reason": "DOCX设计价值创造表格未找到可识别表头。"}
    headers = rows[header_index].get("values") or []
    professional_index = _header_index(headers, "专业类别")
    profit_index = _header_index(headers, "利润")
    profit_field_present = profit_index is not None
    detail_rows = []
    for row in rows[header_index + 1 :]:
        values = row.get("values") or []
        if professional_index is None or professional_index >= len(values):
            continue
        professional = values[professional_index].strip()
        if not professional or professional in {"专业类别", "合计", "总计"}:
            continue
        if not any(value.strip() for value in values):
            continue
        profit = _parse_number(values[profit_index]) if profit_index is not None and profit_index < len(values) else None
        detail_rows.append({**row, "professional": professional, "profit": profit})
    if not detail_rows:
        return [], {"resolved": False, "reason": "DOCX设计价值创造表格未形成可统计明细行。"}
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        groups[row["professional"]].append(row)
    claims: list[dict[str, Any]] = []
    for index, (professional, group_rows) in enumerate(groups.items(), start=1):
        positive = [row for row in group_rows if row["profit"] is not None and row["profit"] > 0]
        locations = [{"table": row["table"], "row": row["row"]} for row in group_rows]
        positive_locations = [{"table": row["table"], "row": row["row"]} for row in positive]
        claims.append(_claim(f"C{index}A", f"{professional}共有{len(group_rows)}条有效明细。", "SOURCE_DOCX", locations))
        if profit_field_present:
            claims.append(_claim(f"C{index}B", f"{professional}中利润>0的明细有{len(positive)}条。", "SOURCE_DOCX", positive_locations))
    total = len(detail_rows)
    positive_total = sum(1 for row in detail_rows if row["profit"] is not None and row["profit"] > 0)
    unresolved_profit = sum(1 for row in detail_rows if row["profit"] is None)
    claims.append(_claim("TOTAL", f"设计价值创造清单共有{total}条有效明细。", "SOURCE_DOCX", [{"table": table["table"]}]))
    if profit_field_present:
        claims.append(_claim("BENEFIT_TOTAL", f"按利润>0口径，明确增加效益的明细有{positive_total}条。", "SOURCE_DOCX", [{"table": table["table"], "row": row["row"]} for row in detail_rows if row["profit"] is not None and row["profit"] > 0]))
        claims.append(_claim("UNRESOLVED", f"有{unresolved_profit}条明细的利润字段为空或未能解析，不能直接认定为无效益。", "SOURCE_DOCX", [{"table": table["table"], "row": row["row"]} for row in detail_rows if row["profit"] is None]))
    return claims, {
        "resolved": profit_field_present,
        "table": table["table"],
        "header": headers,
        "total_rows": total,
        "increase_effect_count": positive_total if profit_field_present else None,
        "undetermined_profit_count": unresolved_profit if profit_field_present else None,
        "profit_field_present": profit_field_present,
        "group_counts": {key: {"total": len(value), "benefit": sum(1 for row in value if row["profit"] is not None and row["profit"] > 0)} for key, value in groups.items()},
        "computed_from_owner_source": True,
        "business_rule": "利润 > 0",
        "reason": None if profit_field_present else "Owner Confirmed DOCX 的设计价值创造清单表未提供利润字段，无法从该 DOCX 确定性计算增加效益条数；旧 XLSX 统计不可直接继承。",
    }


def _table_has_header(table: dict[str, Any], header: str) -> bool:
    return any(header in " | ".join(row.get("values") or []) for row in table.get("rows", []))


def _header_index(headers: list[str], term: str) -> int | None:
    for index, value in enumerate(headers):
        if term in value:
            return index
    return None


def _parse_number(value: str) -> float | None:
    text = str(value or "").strip().replace(",", "")
    if not text or text in {"-", "—", "/"}:
        return None
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def _claim(claim_id: str, text: str, evidence_id: str, locations: list[dict[str, Any]]) -> dict[str, Any]:
    return {"claim_id": claim_id, "text": text, "evidence_id": evidence_id, "source_locations": locations}


def _subquestions(question_id: str) -> list[str]:
    return {
        "BA-001": ["设计任务书包含哪些内容"],
        "BA-002": ["设计创效经济效益额公式", "设计创效率公式", "术语边界"],
        "BA-003": ["厂房产品线案例覆盖哪些专业"],
        "BA-004": ["喷淋管材可选方案", "适用条件"],
        "BA-005": ["特殊环境集电线路电气图审要点"],
        "BA-006": ["2026年局级DOP电子图形文件上传数量"],
        "BA-007": ["设计示范工程实施要求"],
        "BA-008": ["2025年公司设计创效金额"],
        "BA-009": ["土木公司对应督办事项"],
        "BA-010": ["专业列表", "各专业条数", "利润>0条数", "总条数"],
    }[question_id]


def _must_not_use(question_id: str) -> list[str]:
    return {
        "BA-001": ["Wiki导航页单独作为最终事实", "单一项目复盘替代通用要求"],
        "BA-002": ["项目案例替代正式公式", "自行把设计效益增量定义为同一指标"],
        "BA-003": ["Root-001登记页替代厂房正文", "医疗/学校案例替代厂房案例"],
        "BA-004": ["其他喷淋或消防方案行替代第89行"],
        "BA-005": ["项目经验资料替代正式图审依据"],
        "BA-006": ["项目级或二公司级数量替代局级目标"],
        "BA-007": ["旧工作计划作为Primary Gold", "深化设计示范工程要求或BIM示范工程要求作为设计示范工程要求的答案", "科技/BIM示范要求与设计示范要求混为一谈"],
        "BA-008": ["项目级创效金额替代公司年度金额", "删除“约”"],
        "BA-009": ["通用计划或其他项目督办替代土木公司行"],
        "BA-010": ["旧XLSX统计结果未经DOCX对账直接继承", "第88行汇总公式作为明细", "其他项目价值创造清单"],
    }[question_id]


def _review_item(location: dict[str, Any], claim_record: dict[str, Any]) -> dict[str, Any]:
    return {
        "question_id": location["question_id"],
        "question": location["question"],
        "source_owner_confirmation": location["source_owner_confirmation"],
        "gold_location_status": location["gold_location_status"],
        "candidate_correct_file": location["owner_confirmed_source"],
        "expected_location": location["locations"],
        "candidate_claims": claim_record["claims"],
        "owner_confirmation": location["owner_confirmation"],
    }


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip().replace("\n", " ")


def _location_text(item: dict[str, Any]) -> str:
    if "paragraph" in item:
        return f"paragraph {item['paragraph']} ({item.get('style')}): {item.get('text')}"
    return f"table {item.get('table')} row {item.get('row')}: {item.get('text')}"


def _excerpt_from_rows(rows: list[dict[str, Any]]) -> str:
    return "\n".join(str(row.get("text") or row.get("source_excerpt") or "") for row in rows[:12])


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, value: dict[str, Any]) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _update_v2_report(source_records: list[dict[str, Any]], locations: list[dict[str, Any]], unresolved: list[dict[str, Any]]) -> None:
    path = PROJECT_ROOT / "docs" / "BUSINESS_GOLD_V2_REPORT.md"
    previous = path.read_text(encoding="utf-8") if path.exists() else "# Business Gold V2 Report\n"
    previous = re.split(r"\n## TASK-020A\.2(?:\.1)?[^\n]*", previous, maxsplit=1)[0]
    block = [
        "",
        "## TASK-020A.2.1 忠实性与范围修正",
        "",
        f"- Owner Source 记录：{len(source_records)}/10；其中物理来源忠实性已确认 {sum(item.get('source_fidelity_status') != 'UNRESOLVED' for item in source_records)}/10。",
        f"- Gold Location：CONFIRMED {sum(item['gold_location_status'] == 'CONFIRMED' for item in locations)}、PARTIAL {sum(item['gold_location_status'] == 'PARTIAL' for item in locations)}、UNRESOLVED {sum(item['gold_location_status'] == 'UNRESOLVED' for item in locations)}；合计 {len(locations)}/10。",
        "- BA-006 的工作计划仅作为 Supporting Source；责任状正文尚未定位。",
        "- BA-007 Gold Claim 已收窄为“一、设计示范工程实施要求”四项，其他示范类型仅为 Related Context。",
        "- BA-010 继续保持 PARTIAL，不继承旧 XLSX 的统计数字。",
        "- `owner_confirmation` 仍全部为 `UNCONFIRMED`；Source Confirmation 不等于最终 Gold Confirmation。",
        "- `Root-002` 文件治理状态继续保留 `PENDING_APPROVAL`。",
        "- 不生成 final_failure_atlas，不进入 TASK-020B。",
        "",
    ]
    path.write_text(previous.rstrip() + "\n" + "\n".join(block), encoding="utf-8")


def _render_report(sources: list[dict[str, Any]], locations: list[dict[str, Any]], claims: list[dict[str, Any]], unresolved: list[dict[str, Any]]) -> str:
    claims_by_id = {item["question_id"]: item["claims"] for item in claims}
    lines = [
        "# Business Gold Location Closure Report",
        "",
        "> TASK-020A.2.1：修正 Owner Source Fidelity、Gold Claim Scope、BA-010 未决语义和汇总计数。不运行 Fresh Retrieval，不调用 LLM，不扫描新 Root，不修改 Retriever、Qdrant 或 8000。",
        "",
        "## 1. 状态定义",
        "",
        "- `source_owner_confirmation=CONFIRMED`：保留业务负责人此前确认的来源候选记录，不代表当前物理文件的 Fidelity 或最终 Gold 已确认。",
        "- `source_fidelity_status=UNRESOLVED`：当前定位文件不能证明就是负责人所指的 Owner Source，只能作为 Supporting Source。",
        "- `gold_location_status=CONFIRMED`：已找到可复核的最小位置和候选 Claim。",
        "- `owner_confirmation=UNCONFIRMED`：最终 Gold/答案仍待业务负责人审核。",
        "- `Root-002` 继续是 `PENDING_APPROVAL`，不写成 `APPROVED_ROOT`。",
        "",
        "## 2. 汇总",
        "",
        "| 指标 | 数量 |",
        "|---|---:|",
        f"| Owner Source Records | {sum(item['source_owner_confirmation'] == 'CONFIRMED' for item in sources)} / 10 |",
        f"| Physical Source Fidelity Confirmed | {sum(item.get('source_fidelity_status') != 'UNRESOLVED' for item in sources)} / 10 |",
        f"| Gold Location CONFIRMED | {sum(item['gold_location_status'] == 'CONFIRMED' for item in locations)} / 10 |",
        f"| Gold Location PARTIAL | {sum(item['gold_location_status'] == 'PARTIAL' for item in locations)} |",
        f"| Gold Location UNRESOLVED | {sum(item['gold_location_status'] == 'UNRESOLVED' for item in locations)} |",
        f"| Gold Location 待处理（PARTIAL+UNRESOLVED） | {len(unresolved)} |",
        "| Owner Final Gold Confirmed | 0 |",
        "",
        "## 3. 逐题闭环",
        "",
        "| BA | 来源 | Location 状态 | Truth Answerability | Runtime Expected | Governance |",
        "|---|---|---|---|---|---|",
    ]
    for item in locations:
        lines.append(
            f"| {item['question_id']} | `{Path(item['owner_confirmed_source']).name}` | `{item['gold_location_status']}` | `{item['truth_answerability']}` | `{item['runtime_expected_behavior']}` | `{item['source_governance']}` |"
        )
    lines += ["", "## 4. 关键对账", ""]
    for question_id in sorted(locations_by_id(locations)):
        item = locations_by_id(locations)[question_id]
        lines += [f"### {question_id}：{item['question']}", "", f"- Owner Source Label：{item.get('owner_confirmed_source_label') or '未单独提供'}", f"- 当前定位 Source：`{item['owner_confirmed_source']}`", f"- Source Role：`{item.get('source_role')}`", f"- Source Fidelity：`{item.get('source_fidelity_status')}`", f"- 最小位置：`{json.dumps(item['locations'], ensure_ascii=False)}`", f"- Scope：{item.get('expected_scope')}", f"- Authority：{item.get('expected_authority')}", f"- Gold Location：`{item['gold_location_status']}`", "- 候选 Claims："]
        lines.extend(f"  - {claim['claim_id']}：{claim['text']}；evidence=`{claim['evidence_id']}`；location=`{claim['source_locations']}`" for claim in claims_by_id[question_id])
        if item.get("gold_claim_scope_note"):
            lines.append(f"- Gold Claim Scope：{item['gold_claim_scope_note']}")
        if item.get("related_context_locations"):
            lines.append(f"- Related Context（不进入 Gold Claim）：`{json.dumps(item['related_context_locations'], ensure_ascii=False)}`")
        if item.get("supporting_source_path"):
            lines.append(f"- Supporting Source：`{item['supporting_source_path']}`")
        if item.get("source_fidelity_note"):
            lines.append(f"- Source Fidelity Note：{item['source_fidelity_note']}")
        if item.get("fact_reconciliation"):
            lines.append(f"- Fact 对账：`{json.dumps(item['fact_reconciliation'], ensure_ascii=False)}`")
        if item.get("unresolved_reason"):
            lines.append(f"- 未解析原因：{item['unresolved_reason']}")
        lines.append("")
    lines += [
        "## 5. 特别结论",
        "",
        "- BA-005 已切换为 Owner Confirmed XLSX Primary Source，不再保留旧的普通 `AUTHORITY_INSUFFICIENT` Candidate 作为最终状态。",
        "- BA-006 当前仅定位到工作计划 PDF；该文件只作为 Supporting Source，责任状正文尚未定位，不能生成500项目的 Gold Claim。",
        "- BA-007 已切换为 Owner Confirmed DOCX Primary Source；旧的两个 2026 工作计划降级为 Supporting Source。",
        "- BA-009 已切换为指定的 2026 年 4 月监督任务表 Primary Source。",
        "- BA-010 已切换为 Owner Confirmed 星谷项目策划书 DOCX Primary Source；旧 XLSX 只作为结构化辅助对账，统计数字不得直接继承。",
        "- BA-003 即使 Truth Source 已确认，因正文位于当前批准范围外，Runtime 仍应保持 `SOURCE_SCOPE_MISSING`。",
        "",
        "## 6. 输出",
        "",
        "- `evaluation/business_gold_v2/owner_confirmed_source_manifest.json`",
        "- `evaluation/business_gold_v2/gold_location_manifest.json`",
        "- `evaluation/business_gold_v2/gold_claim_manifest.json`",
        "- `evaluation/business_gold_v2/unresolved_gold_items.json`",
        "- `docs/BUSINESS_GOLD_OWNER_FINAL_REVIEW_SHEET.md`",
        "",
        "## 7. 停止点",
        "",
        "本任务不运行 Fresh Retrieval，Fresh Run 留给 TASK-020A.3；本任务不进入 TASK-020B。",
        "",
    ]
    return "\n".join(lines)


def locations_by_id(locations: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {item["question_id"]: item for item in locations}


def _render_review_sheet(locations: list[dict[str, Any]], claims: list[dict[str, Any]]) -> str:
    claims_by_id = {item["question_id"]: item["claims"] for item in claims}
    lines = ["# Business Gold Owner Final Review Sheet", "", "> 来源确认与 Gold 位置/Claim 确认是两件事。以下为系统定向解析结果，最终 Gold 仍需人工确认。", ""]
    for item in locations:
        lines += [
            f"## {item['question_id']}",
            "",
            f"问题：{item['question']}",
            "",
            f"当前定位文件：`{item['owner_confirmed_source']}`",
            f"Owner Source Label：{item.get('owner_confirmed_source_label') or '未单独提供'}",
            f"Source Owner Confirmation：`{item['source_owner_confirmation']}`",
            f"Source Role：`{item.get('source_role')}`",
            f"Source Fidelity：`{item.get('source_fidelity_status')}`",
            f"Source Governance：`{item['source_governance']}`",
            f"Gold Location Status：`{item['gold_location_status']}`",
            "",
            "最小引用位置：",
            f"`{json.dumps(item['locations'], ensure_ascii=False)}`",
            "",
            f"适用 Scope：{item.get('expected_scope')}",
            f"Authority：{item.get('expected_authority')}",
            f"Truth Answerability：`{item['truth_answerability']}`",
            f"Runtime Expected Behavior：`{item['runtime_expected_behavior']}`",
            "",
        ]
        if item.get("gold_claim_scope_note"):
            lines.append(f"Gold Claim Scope：{item['gold_claim_scope_note']}")
            lines.append("")
        if item.get("related_context_locations"):
            lines.append(f"Related Context（不进入 Gold Claim）：`{json.dumps(item['related_context_locations'], ensure_ascii=False)}`")
            lines.append("")
        if item.get("supporting_source_path"):
            lines.append(f"Supporting Source：`{item['supporting_source_path']}`")
            lines.append("")
        if item.get("source_fidelity_note"):
            lines.append(f"Source Fidelity Note：{item['source_fidelity_note']}")
            lines.append("")
        lines += [
            "候选 Claims：",
        ]
        lines.extend(f"- `{claim['claim_id']}` {claim['text']}；evidence=`{claim['evidence_id']}`；位置=`{claim['source_locations']}`" for claim in claims_by_id[item["question_id"]])
        lines += [
            "",
            f"绝不能使用：{'; '.join(item['must_not_use_evidence'])}",
            "",
            "业务负责人选择：",
            "",
            "[ ] 确认",
            "",
            "[ ] 修改后确认",
            "",
            "[ ] 不确认",
            "",
            "业务负责人备注：",
            "",
            "________________",
            "",
        ]
    return "\n".join(lines)


if __name__ == "__main__":
    raise SystemExit(main())
