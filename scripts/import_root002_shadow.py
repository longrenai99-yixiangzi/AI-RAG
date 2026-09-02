from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import time
from collections import Counter
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import torch
from qdrant_client import QdrantClient, models

from app.ingestion.pipeline import PipelineDocument, PipelineResult, run_document_pipeline
from app.retrieval.dense_provider import BGEM3DenseProvider

ROOT_ID = "Root-002"
ROOT = Path(r"D:\工作\二公司技术部")
SUPPORTED = {".pdf", ".docx", ".xlsx", ".pptx", ".md", ".markdown"}
FORBIDDEN_PARTS = ("微信", "临时", "temp", "缓存", "cache", "个人", "草稿")
COLLECTION = "root002_shadow_bge_m3"

APPROVED_DIRS = (
    ("P0", "设计管理", Path(r"D:\工作\二公司技术部\2026\各类文件\设计管理")),
    ("P0", "设计服务台账", Path(r"D:\工作\二公司技术部\2026\设计服务台账")),
    ("P0", "新洲星谷", Path(r"D:\工作\二公司技术部\2026\概算及策划评审\新洲星谷")),
    ("P0", "应城智汇港", Path(r"D:\工作\二公司技术部\2026\概算及策划评审\应城智汇港")),
    ("P1", "设计复盘", Path(r"D:\工作\二公司技术部\2026\设计复盘")),
)


def jsonl_read(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def jsonl_write(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def norm(path: Path | str) -> str:
    return os.path.normcase(str(Path(path).resolve(strict=False)))


def inside(path: Path, root: Path) -> bool:
    try:
        path.resolve(strict=False).relative_to(root.resolve(strict=False))
        return True
    except ValueError:
        return False


def forbidden(path: Path) -> str | None:
    for part in path.parts:
        lowered = part.casefold()
        for token in FORBIDDEN_PARTS:
            if token.casefold() in lowered:
                return token
    return None


def scope_for(path: Path) -> tuple[str, str, Path] | None:
    matches = [item for item in APPROVED_DIRS if inside(path, item[2])]
    return max(matches, key=lambda item: len(str(item[2]))) if matches else None


def discover_files() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    files: dict[str, dict[str, Any]] = {}
    skipped: list[dict[str, Any]] = []
    for priority, label, directory in APPROVED_DIRS:
        if not directory.is_dir():
            skipped.append({"path": str(directory), "reason": "approved_directory_missing"})
            continue
        for path in directory.rglob("*"):
            if not path.is_file():
                continue
            reason = forbidden(path)
            if reason:
                skipped.append({"path": str(path), "reason": f"forbidden:{reason}"})
                continue
            if path.suffix.casefold() not in SUPPORTED:
                skipped.append({"path": str(path), "reason": "unsupported_extension"})
                continue
            stat = path.stat()
            key = norm(path)
            files[key] = {
                "path": str(path.resolve()),
                "filename": path.name,
                "size": stat.st_size,
                "mtime": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(),
                "mtime_ns": stat.st_mtime_ns,
                "sha256": sha256_file(path),
                "file_type": path.suffix.casefold(),
                "knowledge_root_id": ROOT_ID,
                "approval_priority": priority,
                "approval_scope": label,
                "approval_directory": str(directory),
            }
    return sorted(files.values(), key=lambda item: item["path"].casefold()), skipped


def build_source_records(
    snapshots: list[dict[str, Any]], discovery_dir: Path
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    discovered = jsonl_read(discovery_dir / "source_records.jsonl")
    discovered_by_path: dict[str, list[dict[str, Any]]] = {}
    for item in discovered:
        if item.get("resolved_path"):
            discovered_by_path.setdefault(norm(item["resolved_path"]), []).append(item)

    records: list[dict[str, Any]] = []
    links: list[dict[str, Any]] = []
    for snapshot in snapshots:
        path = Path(snapshot["path"])
        source_id = "src-root002-" + hashlib.sha256(
            f"{ROOT_ID}|{norm(path)}".encode("utf-8")
        ).hexdigest()[:24]
        matches = discovered_by_path.get(norm(path), [])
        records.append(
            {
                "source_id": source_id,
                "knowledge_root_id": ROOT_ID,
                "parent_document": (matches[0].get("parent_document") if matches else f"approval:{snapshot['approval_scope']}"),
                "raw_link": (matches[0].get("raw_link") if matches else str(path)),
                "resolved_path": str(path),
                "file_type": snapshot["file_type"],
                "exists": True,
                "inside_allowed_root": inside(path, ROOT),
                "size": snapshot["size"],
                "mtime": snapshot["mtime"],
                "mtime_ns": snapshot["mtime_ns"],
                "sha256": snapshot["sha256"],
                "source_status": "BODY_AVAILABLE",
                "resolution_status": "APPROVED_BODY_AVAILABLE",
                "approval_priority": snapshot["approval_priority"],
                "approval_scope": snapshot["approval_scope"],
                "source_discovery_ids": [item.get("source_id") for item in matches],
            }
        )
        if matches:
            for item in matches:
                links.append(
                    {
                        "source_id": source_id,
                        "from_document": item.get("parent_document"),
                        "target": item.get("raw_link"),
                        "link_type": item.get("link_type"),
                        "line_location": item.get("line_location"),
                        "resolution_status": "APPROVED_BODY_AVAILABLE",
                    }
                )
        else:
            links.append(
                {
                    "source_id": source_id,
                    "from_document": f"approval:{snapshot['approval_scope']}",
                    "target": str(path),
                    "link_type": "approved_path",
                    "line_location": None,
                    "resolution_status": "APPROVED_BODY_AVAILABLE",
                }
            )
    return records, links


def run_pipeline_resilient(files: list[Path]) -> tuple[PipelineResult, list[dict[str, Any]]]:
    documents: list[PipelineDocument] = []
    errors: list[str] = []
    for path in files:
        try:
            result = run_document_pipeline(ROOT, files=[path])
            documents.extend(result.documents)
            errors.extend(result.scan_errors)
        except Exception as error:  # One bad external file cannot stop the Shadow batch.
            errors.append(f"{path}: {type(error).__name__}: {error}")
    return PipelineResult(ROOT, documents, errors, len(files)), []


def xlsx_details(document: PipelineDocument) -> list[dict[str, Any]]:
    if document.file_type != ".xlsx":
        return []
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(document.path, read_only=True, data_only=True)
        sheet_names = list(workbook.sheetnames)
        blocks_by_sheet = {block.heading_path: block for block in document.source_blocks}
        result = []
        for name in sheet_names:
            block = blocks_by_sheet.get(name)
            location = block.location if block else {}
            header = ""
            if block:
                header_line = next((line for line in block.text.splitlines() if line.startswith("表头")), "")
                header = header_line.split("：", 1)[1] if "：" in header_line else ""
            result.append(
                {
                    "workbook": document.path.name,
                    "sheet_name": name,
                    "valid_sheet": bool(block),
                    "headers": header.split(" | ") if header else [],
                    "row_start": location.get("row_start"),
                    "row_end": location.get("row_end"),
                    "column_count": location.get("column_count"),
                    "source_location": location,
                }
            )
        workbook.close()
        return result
    except Exception as error:
        return [{"workbook": document.path.name, "error": f"{type(error).__name__}: {error}"}]


def staging_record(document: PipelineDocument, snapshot: dict[str, Any], source_id: str) -> dict[str, Any]:
    metadata = dict(document.metadata)
    metadata.update(
        {
            "knowledge_root_id": ROOT_ID,
            "source_record_id": source_id,
            "approval_priority": snapshot["approval_priority"],
            "approval_scope": snapshot["approval_scope"],
        }
    )
    chunks = [asdict(chunk) for chunk in document.chunks]
    blocks = [asdict(block) for block in document.source_blocks]
    chunk_metadata = {}
    for chunk_id, value in document.chunk_metadata.items():
        item = dict(value)
        item.update(
            {
                "knowledge_root_id": ROOT_ID,
                "source_record_id": source_id,
                "approval_priority": snapshot["approval_priority"],
                "approval_scope": snapshot["approval_scope"],
            }
        )
        chunk_metadata[chunk_id] = item
    return {
        "source_record_id": source_id,
        "knowledge_root_id": ROOT_ID,
        "approval_priority": snapshot["approval_priority"],
        "approval_scope": snapshot["approval_scope"],
        "path": str(document.path),
        "file_name": document.path.name,
        "file_type": document.file_type,
        "status": document.status,
        "error": document.error,
        "source_blocks": blocks,
        "chunks": chunks,
        "metadata": metadata,
        "chunk_metadata": chunk_metadata,
        "quality": document.quality,
        "xlsx": xlsx_details(document),
    }


def point_id(chunk_id: str) -> str:
    import uuid

    try:
        uuid.UUID(chunk_id)
        return chunk_id
    except ValueError:
        return str(uuid.uuid5(uuid.NAMESPACE_URL, chunk_id))


def payload(chunk: Any, metadata: dict[str, Any], source_id: str, snapshot: dict[str, Any]) -> dict[str, Any]:
    return {
        "chunk_id": chunk.chunk_id,
        "document_id": chunk.document_id,
        "source_record_id": source_id,
        "knowledge_root_id": ROOT_ID,
        "approval_priority": snapshot["approval_priority"],
        "approval_scope": snapshot["approval_scope"],
        "source_path": chunk.source_path,
        "file_name": chunk.file_name,
        "file_type": snapshot["file_type"],
        "heading_path": chunk.heading_path,
        "location": chunk.location,
        "text": chunk.text,
        "metadata": metadata,
    }


def build_shadow(
    chunks: list[tuple[Any, dict[str, Any], str, dict[str, Any]]],
    shadow_dir: Path,
    model_path: Path,
    batch_size: int,
) -> dict[str, Any]:
    started = time.perf_counter()
    qdrant_dir = shadow_dir / "qdrant"
    qdrant_dir.mkdir(parents=True, exist_ok=True)
    result: dict[str, Any] = {
        "cuda_available": bool(torch.cuda.is_available()),
        "torch": torch.__version__,
        "torch_cuda": torch.version.cuda,
        "gpu_name": torch.cuda.get_device_name(0) if torch.cuda.is_available() else None,
        "model_path": str(model_path),
        "embedding_count": 0,
        "embedding_failures": [],
        "elapsed_seconds": None,
        "gpu_peak_memory_allocated_bytes": 0,
        "gpu_peak_memory_reserved_bytes": 0,
    }
    if not torch.cuda.is_available():
        result["error"] = "torch.cuda.is_available() is False; embedding not run"
        result["elapsed_seconds"] = round(time.perf_counter() - started, 3)
        return result

    client = QdrantClient(path=str(qdrant_dir))
    if client.collection_exists(COLLECTION):
        client.delete_collection(COLLECTION)
    client.create_collection(
        collection_name=COLLECTION,
        vectors_config=models.VectorParams(size=BGEM3DenseProvider.VECTOR_SIZE, distance=models.Distance.COSINE),
    )
    provider = BGEM3DenseProvider(model_path, use_fp16=True, batch_size=batch_size)
    try:
        provider.load()
        torch.cuda.synchronize()
        torch.cuda.reset_peak_memory_stats()
        for start in range(0, len(chunks), batch_size):
            batch = chunks[start : start + batch_size]
            try:
                vectors = provider.embed_documents([item[0].text for item in batch])
                points = [
                    models.PointStruct(
                        id=point_id(item[0].chunk_id),
                        vector=vector,
                        payload=payload(item[0], item[1], item[2], item[3]),
                    )
                    for item, vector in zip(batch, vectors, strict=True)
                ]
                client.upsert(collection_name=COLLECTION, points=points, wait=True)
                result["embedding_count"] += len(points)
            except Exception as error:
                for item in batch:
                    result["embedding_failures"].append(
                        {
                            "chunk_id": item[0].chunk_id,
                            "source_path": item[0].source_path,
                            "error": f"{type(error).__name__}: {error}",
                        }
                    )
                torch.cuda.empty_cache()
        torch.cuda.synchronize()
        result["gpu_peak_memory_allocated_bytes"] = int(torch.cuda.max_memory_allocated())
        result["gpu_peak_memory_reserved_bytes"] = int(torch.cuda.max_memory_reserved())
    except Exception as error:
        result["error"] = f"{type(error).__name__}: {error}"
    finally:
        provider.close()
        client.close()
    result["elapsed_seconds"] = round(time.perf_counter() - started, 3)
    return result


def qdrant_validation(shadow_dir: Path, chunks: list[tuple[Any, dict[str, Any], str, dict[str, Any]]]) -> dict[str, Any]:
    client = QdrantClient(path=str(shadow_dir / "qdrant"))
    try:
        if not client.collection_exists(COLLECTION):
            return {"collection_exists": False, "points": 0, "payload_valid": 0, "citation_location_complete": 0}
        count = client.count(collection_name=COLLECTION, exact=True).count
        expected = {item[0].chunk_id for item in chunks}
        seen: set[str] = set()
        payload_valid = 0
        citation_complete = 0
        offset = None
        while True:
            points, offset = client.scroll(
                collection_name=COLLECTION,
                limit=256,
                offset=offset,
                with_payload=True,
                with_vectors=False,
            )
            for point in points:
                data = point.payload or {}
                chunk_id = str(data.get("chunk_id") or "")
                seen.add(chunk_id)
                if data.get("knowledge_root_id") == ROOT_ID and data.get("source_path") and data.get("location"):
                    payload_valid += 1
                if data.get("location"):
                    citation_complete += 1
            if offset is None:
                break
        return {
            "collection_exists": True,
            "points": count,
            "payload_valid": payload_valid,
            "citation_location_complete": citation_complete,
            "expected_chunk_ids": len(expected),
            "chunk_ids_present": len(expected & seen),
            "chunk_ids_missing": len(expected - seen),
        }
    finally:
        client.close()


def focus_rows(staging: list[dict[str, Any]], source_records: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    focus = {
        "BA-007": ("设计与技术工作计划", "2026年设计与技术工作计划"),
        "BA-009": ("设计服务管理台帐2026", "设计服务管理台账"),
        "BA-010": ("新洲星谷", "星谷科创中心", "方案比选与价值创造清单方案比选及价值创造"),
    }
    rows: dict[str, list[dict[str, Any]]] = {key: [] for key in focus}
    for record in source_records:
        haystack = f"{record['resolved_path']} {record['raw_link']}"
        for key, terms in focus.items():
            if any(term.casefold() in haystack.casefold() for term in terms):
                rows[key].append({"source_record_id": record["source_id"], **record})
    for record in staging:
        haystack = f"{record['path']} {record['file_name']}"
        for key, terms in focus.items():
            if any(term.casefold() in haystack.casefold() for term in terms):
                rows[key].append(
                    {
                        "path": record["path"],
                        "file_name": record["file_name"],
                        "status": record["status"],
                        "source_blocks": len(record["source_blocks"]),
                        "chunks": len(record["chunks"]),
                        "xlsx": record.get("xlsx", []),
                    }
                )
    return rows


def render_report(
    *,
    args: argparse.Namespace,
    snapshots: list[dict[str, Any]],
    skipped: list[dict[str, Any]],
    source_records: list[dict[str, Any]],
    pipeline: PipelineResult,
    staging: list[dict[str, Any]],
    embedding: dict[str, Any],
    qdrant: dict[str, Any],
    focus: dict[str, list[dict[str, Any]]],
) -> str:
    ext_counts = Counter(item["file_type"] for item in snapshots)
    status_counts = pipeline.status_counts()
    metadata_complete = pipeline.metadata_complete_chunk_count()
    location_complete = pipeline.location_valid_chunk_count()
    lines = [
        "# ROOT-002 Shadow Import & Pipeline Validation Report",
        "",
        "> TASK-016C-2 只对 TASK-016C-1 批准范围运行。结果写入独立 Shadow 目录，不进入正式 Retriever、8000 服务或正式 Qdrant。",
        "",
        "## 1. 执行边界",
        "",
        f"- Knowledge Root：`{ROOT_ID}` — `{ROOT}`",
        f"- Shadow 目录：`{args.shadow_dir.resolve()}`",
        f"- Shadow Collection：`{COLLECTION}`",
        "- 未修改正式 Retriever、`app/main.py`、8000 服务和正式 `data\\qdrant`。",
        "- 未自动导入个人资料、临时文件、微信文件、缓存文件和未审核草稿。",
        "",
        "## 2. Root-002 审批快照",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| 批准目录数 | {len(APPROVED_DIRS)} |",
        f"| 快照文件数 | {len(snapshots)} |",
        f"| 跳过文件/目录数 | {len(skipped)} |",
        f"| PDF | {ext_counts.get('.pdf', 0)} |",
        f"| DOCX | {ext_counts.get('.docx', 0)} |",
        f"| XLSX | {ext_counts.get('.xlsx', 0)} |",
        f"| PPTX | {ext_counts.get('.pptx', 0)} |",
        f"| Markdown | {ext_counts.get('.md', 0) + ext_counts.get('.markdown', 0)} |",
        "",
        "快照字段：`path`、`filename`、`size`、`mtime`、`mtime_ns`、`sha256`、审批优先级和批准子目录。",
        f"快照文件：`{args.shadow_dir / 'approval_snapshot.jsonl'}`",
        "",
        "### 批准目录",
        "",
        "| 优先级 | 子目录 | 路径 |",
        "|---|---|---|",
    ]
    lines.extend(f"| {priority} | {label} | `{path}` |" for priority, label, path in APPROVED_DIRS)
    lines += [
        "",
        "## 3. Source Resolver",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| Root-002 SourceRecord | {len(source_records)} |",
        "| source_status | BODY_AVAILABLE |",
        f"| Root-002 文件存在数 | {sum(item['exists'] for item in source_records)} |",
        f"| Root-002 文件类型 | {json.dumps(dict(ext_counts), ensure_ascii=False)} |",
        "",
        f"SourceRecord：`{args.shadow_dir / 'source_records.jsonl'}`",
        f"SourceLink：`{args.shadow_dir / 'source_links.jsonl'}`",
        "",
        "## 4. Document Pipeline",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| Document | {len(pipeline.documents)} |",
        f"| SourceBlock | {pipeline.source_block_count} |",
        f"| Chunk | {pipeline.chunk_count} |",
        f"| Metadata 完整 Chunk | {metadata_complete} |",
        f"| Citation location 有效 Chunk | {location_complete} |",
        f"| Pipeline 状态 | `{json.dumps(status_counts, ensure_ascii=False)}` |",
        f"| Pipeline 错误 | {len(pipeline.scan_errors)} |",
        "",
        f"Metadata staging：`{args.shadow_dir / 'pipeline_staging.jsonl'}`",
        "",
        "## 5. Excel 特殊处理",
        "",
        "每个 XLSX 记录 Workbook、Sheet 名称、有效 Sheet、表头、行列范围和 SourceBlock location；不执行复杂公式计算，也不回写源文件。",
        "",
        "| Workbook | Sheet | 有效 | 表头 | 行范围 | 列数 | 来源位置 |",
        "|---|---|---|---|---:|---:|---|",
    ]
    xlsx_rows = [sheet for record in staging for sheet in record.get("xlsx", [])]
    for sheet in xlsx_rows:
        if "error" in sheet:
            lines.append(f"| `{sheet['workbook']}` | - | 否 | - | - | - | {sheet['error']} |")
        else:
            row_range = f"{sheet.get('row_start') or '-'}-{sheet.get('row_end') or '-'}"
            headers = "、".join(sheet.get("headers") or [])[:120]
            lines.append(
                f"| `{sheet['workbook']}` | `{sheet['sheet_name']}` | {'是' if sheet['valid_sheet'] else '否'} | {headers} | {row_range} | {sheet.get('column_count') or '-'} | `{json.dumps(sheet.get('source_location') or {}, ensure_ascii=False)}` |"
            )
    if not xlsx_rows:
        lines.append("| 无 | - | - | - | - | - | - |")
    lines += [
        "",
        "## 6. Shadow Embedding 与 Qdrant",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| CUDA 可用 | `{embedding.get('cuda_available')}` |",
        f"| Torch | `{embedding.get('torch')}` |",
        f"| CUDA Runtime | `{embedding.get('torch_cuda')}` |",
        f"| GPU | `{embedding.get('gpu_name') or '-'}` |",
        f"| BGE-M3 路径 | `{embedding.get('model_path')}` |",
        f"| Embedding 数量 | {embedding.get('embedding_count', 0)} |",
        f"| Embedding 失败数量 | {len(embedding.get('embedding_failures', []))} |",
        f"| Embedding 耗时（秒） | {embedding.get('elapsed_seconds')} |",
        f"| GPU 峰值显存 allocated（bytes） | {embedding.get('gpu_peak_memory_allocated_bytes', 0)} |",
        f"| GPU 峰值显存 reserved（bytes） | {embedding.get('gpu_peak_memory_reserved_bytes', 0)} |",
        f"| Shadow Collection 存在 | `{qdrant.get('collection_exists')}` |",
        f"| Shadow Qdrant 数量 | {qdrant.get('points', 0)} |",
        f"| Payload Root-002 完整数量 | {qdrant.get('payload_valid', 0)} |",
        f"| Citation location 完整数量 | {qdrant.get('citation_location_complete', 0)} |",
        f"| Chunk ID 应有/已入 Shadow | {qdrant.get('expected_chunk_ids', 0)}/{qdrant.get('chunk_ids_present', 0)} |",
        f"| Chunk ID 缺失 | {qdrant.get('chunk_ids_missing', 0)} |",
        "",
        f"Shadow Qdrant 路径：`{args.shadow_dir / 'qdrant'}`",
        "",
        "## 7. 重点业务问题验证",
        "",
    ]
    focus_titles = {
        "BA-007": "中建三局2026年设计与技术工作计划",
        "BA-009": "2026设计服务管理台账",
        "BA-010": "星谷价值创造清单",
    }
    for key, title in focus_titles.items():
        lines += [f"### {key}：{title}", ""]
        items = focus.get(key, [])
        if not items:
            lines.append("- 未在批准范围快照中匹配到文件。")
        for item in items[:20]:
            if "source_record_id" in item and "status" not in item:
                lines.append(
                    f"- SourceRecord `{item['source_record_id']}`：`{item['resolved_path']}`；source_status=`{item['source_status']}`；sha256=`{item['sha256']}`"
                )
            else:
                lines.append(
                    f"- Pipeline：`{item['file_name']}`；status=`{item['status']}`；SourceBlock={item['source_blocks']}；Chunk={item['chunks']}"
                )
                for sheet in item.get("xlsx", []):
                    if "sheet_name" in sheet:
                        lines.append(
                            f"  - Sheet `{sheet['sheet_name']}`：有效={sheet['valid_sheet']}，行={sheet.get('row_start')}-{sheet.get('row_end')}，列数={sheet.get('column_count')}"
                        )
        lines.append("")
    lines += [
        "## 8. 异常与未导入项",
        "",
        f"- 跳过项总数：{len(skipped)}。原因分布：`{json.dumps(dict(Counter(item['reason'] for item in skipped)), ensure_ascii=False)}`。",
        f"- Pipeline 错误：{json.dumps(pipeline.scan_errors, ensure_ascii=False) if pipeline.scan_errors else '无'}。",
        f"- Embedding 错误：{json.dumps(embedding.get('embedding_failures', []), ensure_ascii=False) if embedding.get('embedding_failures') else '无'}。",
        "- `.docm`、`.wps`、`.7z` 等不在本任务 Loader 支持范围内的文件只登记为跳过项，不自动转换、不自动导入。",
        "- 未审核草稿、个人资料、临时文件、微信文件和缓存目录不进入 Shadow。",
        "",
        "## 9. 结论",
        "",
        "本次已按 Root-002 选择性批准范围完成只读快照、SourceRecord、Document Pipeline、Metadata staging、BGE-M3 Embedding 和独立 Shadow Qdrant 验证。该结果仅证明候选资料可被技术链路处理，不代表已进入正式知识库，也不代表业务内容已经审核通过。",
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Import approved Root-002 files into an isolated Shadow pipeline")
    parser.add_argument("--shadow-dir", type=Path, default=Path("data") / "shadow" / "root002_import")
    parser.add_argument("--discovery-dir", type=Path, default=Path("data") / "shadow" / "source_discovery")
    parser.add_argument("--batch-size", type=int, default=4)
    args = parser.parse_args()
    if args.batch_size < 1:
        parser.error("--batch-size must be positive")

    snapshots, skipped = discover_files()
    args.shadow_dir.mkdir(parents=True, exist_ok=True)
    jsonl_write(args.shadow_dir / "approval_snapshot.jsonl", snapshots)
    jsonl_write(args.shadow_dir / "approval_skipped.jsonl", skipped)
    source_records, source_links = build_source_records(snapshots, args.discovery_dir)
    jsonl_write(args.shadow_dir / "source_records.jsonl", source_records)
    jsonl_write(args.shadow_dir / "source_links.jsonl", source_links)

    snapshot_by_path = {norm(item["path"]): item for item in snapshots}
    source_id_by_path = {norm(item["resolved_path"]): item["source_id"] for item in source_records}
    files = [Path(item["path"]) for item in snapshots]
    pipeline, _ = run_pipeline_resilient(files)
    staging = [
        staging_record(
            document,
            snapshot_by_path[norm(document.path)],
            source_id_by_path[norm(document.path)],
        )
        for document in pipeline.documents
    ]
    jsonl_write(args.shadow_dir / "pipeline_staging.jsonl", staging)

    chunk_rows = []
    staging_by_path = {norm(item["path"]): item for item in staging}
    for document in pipeline.documents:
        if document.status not in {"parsed", "empty"}:
            continue
        snapshot = snapshot_by_path[norm(document.path)]
        source_id = source_id_by_path[norm(document.path)]
        metadata = staging_by_path[norm(document.path)]["metadata"]
        chunk_rows.extend((chunk, metadata, source_id, snapshot) for chunk in document.chunks)

    embedding = build_shadow(
        chunk_rows,
        args.shadow_dir,
        PROJECT_ROOT / "models" / "bge-m3",
        args.batch_size,
    )
    qdrant = qdrant_validation(args.shadow_dir, chunk_rows)
    focus = focus_rows(staging, source_records)
    report = render_report(
        args=args,
        snapshots=snapshots,
        skipped=skipped,
        source_records=source_records,
        pipeline=pipeline,
        staging=staging,
        embedding=embedding,
        qdrant=qdrant,
        focus=focus,
    )
    report_path = Path("docs") / "ROOT002_SHADOW_IMPORT_REPORT.md"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(report, encoding="utf-8")
    summary = {
        "snapshot_files": len(snapshots),
        "skipped": len(skipped),
        "source_records": len(source_records),
        "documents": len(pipeline.documents),
        "source_blocks": pipeline.source_block_count,
        "chunks": pipeline.chunk_count,
        "embedding_count": embedding.get("embedding_count", 0),
        "qdrant_points": qdrant.get("points", 0),
        "pipeline_status": pipeline.status_counts(),
        "report": str(report_path.resolve()),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if not embedding.get("error") and not embedding.get("embedding_failures") else 1


if __name__ == "__main__":
    raise SystemExit(main())
