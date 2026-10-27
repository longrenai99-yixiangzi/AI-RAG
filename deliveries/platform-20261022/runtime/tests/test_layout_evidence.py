from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from app.ingestion.layout_evidence import list_regions, slide_nodes, record,labeled_text_regions
from app.layout_answers import layout_proposal, render_layout
from app.retrieval.query_planner_v1 import plan_query
from app.verified_answer_engine_v2 import render, validate, _explicit_list_item_count


def node(text,x,y,*,bold=False,size=18):
    return {'text':text,'shape_id':text,'groups':[],'bbox':[x,y,x+.15,y+.05],'bold':bold,'font_size':size,'paragraphs':[{'text':text,'bold':bold,'font_size':size}]}


def sample(label='设计管控四大模块',items=None,context='给排水专业特点分析',kind='list_group'):
    items = items or ['室内供水','室外排水','专项协调','消防设计']
    body = context+'\n'+label+'\n'+'\n'.join(items)
    return record(Path('source.pptx'),'hash',{'slide':7},context+' > '+label,body,
                  {'kind':kind,'label':label,'context':context,'source_context':body,'items':[{'text':v} for v in items],'region_count':2})


def propose(question,rows):
    return layout_proposal(question,plan_query(question).to_dict(),rows)


def test_nested_group_coordinates_are_converted_to_slide_coordinates():
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);group=slide.shapes.add_group_shape()
    shape=group.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1));shape.text='模块'
    group.left=Inches(6);group.top=Inches(3);group.width=Inches(2);group.height=Inches(1)
    nodes=slide_nodes(slide,prs.slide_width,prs.slide_height)
    assert abs(nodes[0]['bbox'][0]-6/prs.slide_width*914400)<.001
    assert nodes[0]['groups']==[str(group.shape_id)]


def test_ppt_table_intro_is_kept_as_native_context():
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    table=slide.shapes.add_table(1,1,Inches(1),Inches(1),Inches(6),Inches(1)).table
    table.cell(0,0).text='从目前EPC项目的运行情况来看，存在三个方面的问题。'
    nodes=slide_nodes(slide,prs.slide_width,prs.slide_height)
    assert nodes[0]['kind']=='table' and 'EPC项目' in nodes[0]['text']


def test_peer_columns_do_not_mix_items():
    nodes=[node('四大模块',.05,.2,bold=True),node('管控要点',.55,.2,bold=True)]
    nodes += [node(s,x,y) for s,x,y in [('甲模块',.1,.4),('乙模块',.3,.4),('丙模块',.1,.7),('丁模块',.3,.7),('错误栏目甲',.6,.4),('错误栏目乙',.6,.6)]]
    regions=list_regions(nodes);left=next(r for r in regions if r['label']=='四大模块')
    assert {i['text'] for i in left['items']}=={'甲模块','乙模块','丙模块','丁模块'}


def test_labeled_problem_and_measure_columns_do_not_mix_full_paragraphs():
    nodes=[node('【存在问题】',.1,.2),node('【改进措施】',.65,.2),node('原文问题及其条件。',.1,.4),node('原文改进措施及例外。',.65,.4)]
    regions={r['label']:r for r in labeled_text_regions(nodes)}
    assert regions['存在问题']['text']=='原文问题及其条件。'
    assert regions['改进措施']['text']=='原文改进措施及例外。'


def test_stacked_measure_panel_does_not_enter_the_problem_region():
    nodes=[node('存在问题',.45,.15),node('问题一。',.45,.3),node('改进措施',.45,.55),node('完整措施。',.45,.7)]
    regions={r['label']:r for r in labeled_text_regions(nodes)}
    assert regions['存在问题']['text']=='问题一。' and regions['改进措施']['text']=='完整措施。'


def test_cross_panel_text_is_held_instead_of_assigned_by_its_center_only():
    nodes=[node('存在问题',.1,.2),node('改进措施',.65,.2)]
    text=node('横跨两栏的文字',.1,.4);text['bbox']=[.1,.4,.9,.5];nodes.append(text)
    regions=labeled_text_regions(nodes)
    assert not any('横跨' in r['text'] for r in regions)
    assert any(text['shape_id'] in r['cross_boundary_shape_ids'] for r in regions)


def test_list_answers_use_the_requested_region_and_correct_count():
    right=sample('设计管控要点',['系统规模','设备选型','接口','永临结合'])
    result=propose('给排水设计管控的四大模块有哪些？',[right,sample()])
    assert result and result['items']==['室内供水','室外排水','专项协调','消防设计']
    assert propose('给排水设计管控的五个模块有哪些？',[sample()]) is None
    assert propose('给排水专业特点分析，设计管控要点有哪些？',[right,sample()])['items']==['系统规模','设备选型','接口','永临结合']


def test_wrong_object_and_unseen_replacement_values_are_not_answer_lookups():
    source=sample('冷却塔选型方案比选',['甲型冷却塔','乙型冷却塔'],'冷却塔选型方案比选','comparison_options')
    assert propose('冷却塔方案比选有哪些方案？',[source])['items']==['甲型冷却塔','乙型冷却塔']
    assert propose('水泵选型比选有哪些方案？',[source]) is None
    assert propose('请给出某医院项目的冷却塔选型方案？',[source]) is None
    generic=sample('空调方案比选',['制冷机组','VRV'],'空调','comparison_options')
    assert propose('空调水泵选型比选有哪些方案？',[generic]) is None


def test_generic_epc_phrase_is_not_a_project_scope_but_named_epc_project_is():
    assert plan_query('从目前EPC项目的实际运行情况来看，问题体现在哪三个方面').project==[]
    assert plan_query('丰台长辛店EPC项目合同金额是多少').project


def test_unmatched_stage_compound_question_and_unbacked_items_are_rejected():
    source=sample('控制要点',['甲','乙','丙','丁','戊'],'暖通专业 初步设计阶段管控要点')
    source['layout']['region_count']=1
    assert propose('暖通施工图设计阶段哪五个方面？',[source]) is None
    assert propose('暖通初步设计有哪五个方面，并说明负责人？',[source]) is None
    source['layout']['items'][0]['text']='不在原文里的新增结论'
    assert propose('暖通初步设计管控哪五个方面？',[source]) is None


def test_layout_render_requires_retrieved_evidence_and_preserves_citation():
    source=sample();question='给排水设计管控四大模块有哪些？'
    evidence={**source,'role':'CONTEXT_ONLY'}
    bundle={'query_id':'q','question':question,'query_plan':plan_query(question).to_dict(),'bundle_status':'VERIFIED_PARTIAL','candidate_evidence':[evidence], 'verified_evidence':[],'supporting_evidence':[],'conflicting_evidence':[],'coverage_map':[],'layout_source_records':{source['evidence_id']:source}}
    answer=render_layout(bundle)
    assert answer['answer_status']=='ANSWERED' and validate(answer,{**bundle,'candidate_evidence':[{**evidence,'role':'DIRECT'}]})['valid']
    assert answer['citations'][0]['location']['slide']==7
    bundle['candidate_evidence']=[]
    assert render_layout(bundle) is None
    bundle['candidate_evidence']=[evidence];bundle['bundle_status']='CONFLICTING_EVIDENCE'
    assert render_layout(bundle) is None


def test_explicit_numbered_items_take_precedence_over_stage_heading():
    text='6、暖通专业 > 初步设计阶段管控要点：\n1. 主要机房\n2. 主要设备\n3. 风水系统\n4. 管井\n5. 冷热指标'
    assert _explicit_list_item_count(text,question='初步设计阶段哪五个方面？')==5
    assert _explicit_list_item_count('问题与困难：\n1. 甲\n2. 乙\n3. 丙',question='体现在哪三个方面？')==3
    assert _explicit_list_item_count('1. 甲\n3. 丙',question='哪三个方面？')!=3
