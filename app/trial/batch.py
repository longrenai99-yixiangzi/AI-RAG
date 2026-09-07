from __future__ import annotations

import hashlib
import hmac
import re
import secrets
import uuid
from collections import Counter
from pathlib import Path
from pathlib import PurePosixPath
from typing import Any

from docx import Document
from fastapi import APIRouter, Header, HTTPException, Query, Request
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.ingestion.loaders.pdf_loader import extract_value_creation_rows, extract_value_creation_summary, load_pdf

from . import v2


router = APIRouter(prefix="/api/v2/batch", tags=["V2 Batch Workflow"])
SUPPORTED_EXTENSIONS = {".md", ".pdf", ".docx", ".xlsx", ".pptx", ".wps"}
UPLOAD_ROOT = v2.BATCH_ROOT / "uploads"
UPLOAD_SESSIONS = v2.BATCH_ROOT / "upload_sessions.jsonl"
MAX_UPLOAD_FILE_BYTES = 100 * 1024 * 1024
MAX_UPLOAD_TOTAL_BYTES = 512 * 1024 * 1024


class BatchPathsRequest(BaseModel):
    paths: list[str] = Field(default_factory=list)
    recursive: bool = True
    trial_user: str = "reviewer-001"


class BatchGenerateRequest(BaseModel):
    paths: list[str] = Field(default_factory=list)
    trial_user: str = "reviewer-001"


class BatchRunRequest(BaseModel):
    case_ids: list[str] = Field(default_factory=list)
    trial_user: str = "reviewer-001"


class UploadStartRequest(BaseModel):
    trial_user: str = "reviewer-001"


class UploadCompleteRequest(BaseModel):
    trial_user: str = "reviewer-001"


def _ensure_batch_user(user_id: str) -> None:
    v2._ensure_user(user_id)
    if v2.USERS[user_id].get("role") not in {"BUSINESS_REVIEWER", "BATCH_OPERATOR"}:
        raise HTTPException(status_code=403, detail="BATCH_PERMISSION_DENIED")


@router.post("/sources/preview")
def preview_sources(request: BatchPathsRequest) -> dict[str, Any]:
    _ensure_batch_user(request.trial_user)
    files, errors = _discover_files(request.paths, request.recursive)
    return {
        "source_count": len(files),
        "files": [_source_probe(path) for path in files],
        "errors": errors,
        "read_only": True,
        "approval_required": True,
    }


@router.post("/sources/approve")
def approve_sources(request: BatchPathsRequest) -> dict[str, Any]:
    _ensure_batch_user(request.trial_user)
    files, errors = _discover_files(request.paths, request.recursive)
    approved = []
    for path in files:
        source_path = str(path)
        row = {
            "source_path": source_path,
            "path": source_path,
            "file_name": path.name,
            "file_type": path.suffix.lower(),
            "sha256": _sha256(path),
            "size": path.stat().st_size,
            "status": "APPROVED",
            "approval_status": "USER_APPROVED_SHADOW_READ",
            "approved_at": v2._now(),
            "created_by": request.trial_user,
            "knowledge_root_id": "Root-002",
            "read_only": True,
        }
        if not any(item.get("source_path") == source_path and item.get("status") == "APPROVED" for item in v2._read_jsonl(v2.BATCH_SOURCE_REGISTER)):
            v2._append_jsonl(v2.BATCH_SOURCE_REGISTER, row)
        approved.append(row)
    if approved and v2._engine is not None:
        v2._engine.load_approved_sources()
    return {"approved_count": len(approved), "approved": approved, "errors": errors, "formal_knowledge_publish": 0}


@router.post("/questions/generate")
def generate_questions(request: BatchGenerateRequest) -> dict[str, Any]:
    _ensure_batch_user(request.trial_user)
    approved = _approved_sources(request.paths)
    cases = v2._read_jsonl(v2.BATCH_CASES)
    existing = {row.get("fingerprint") for row in cases}
    created = []
    skipped = []
    for source in approved:
        path = Path(str(source["source_path"]))
        generated = _generate_cases(path)
        if not generated:
            skipped.append({"source_path": str(path), "reason": "当前规则无法安全推导验收事实"})
            continue
        for case in generated:
            if case["fingerprint"] in existing:
                continue
            v2._append_jsonl(v2.BATCH_CASES, case)
            existing.add(case["fingerprint"])
            created.append(case)
    return {"created_count": len(created), "created": created, "skipped": skipped, "runtime_input_injection": 0}


@router.get("/questions")
def batch_questions(source_paths: list[str] = Query(default=[]), trial_user: str = Query(default="reviewer-001")) -> dict[str, Any]:
    _ensure_batch_user(trial_user)
    return {"questions": _filter_cases(list(v2._latest_by_id(v2._read_jsonl(v2.BATCH_CASES), "case_id").values()), source_paths)}


@router.post("/regressions/run")
def run_batch_regressions(request: BatchRunRequest) -> dict[str, Any]:
    v2._ensure_user(request.trial_user)
    engine = v2._engine_instance()
    cases = list(v2._latest_by_id(v2._read_jsonl(v2.BATCH_CASES), "case_id").values())
    if request.case_ids:
        wanted = set(request.case_ids)
        cases = [case for case in cases if case.get("case_id") in wanted]
    runs = []
    for case in cases:
        expected_path = v2._normalize_source_path(case.get("expected_source_path"))
        if expected_path not in v2.LOADED_SHADOW_PATHS:
            run = _not_ready_run(case, request.trial_user)
        else:
            result = engine.answer(str(case.get("question") or ""), detect_growth=False, source_paths=[expected_path])
            run = _evaluate_case(case, result, request.trial_user)
        v2._append_jsonl(v2.BATCH_RUNS, run)
        runs.append(run)
    return {"run_count": len(runs), "passed_count": sum(row["regression_status"] == "PASSED" for row in runs), "failed_count": sum(row["regression_status"] == "FAILED" for row in runs), "not_ready_count": sum(row["regression_status"] == "NOT_READY" for row in runs), "observed_count": sum(row["regression_status"] == "OBSERVED" for row in runs), "runs": runs, "formal_knowledge_publish": 0}


@router.get("/anomalies")
def batch_anomalies(source_paths: list[str] = Query(default=[]), trial_user: str = Query(default="reviewer-001")) -> dict[str, Any]:
    _ensure_batch_user(trial_user)
    cases = _filter_cases(list(v2._latest_by_id(v2._read_jsonl(v2.BATCH_CASES), "case_id").values()), source_paths)
    latest = {case_id: row for case_id, row in v2._latest_by_id(v2._read_jsonl(v2.BATCH_RUNS), "case_id").items() if case_id in {case["case_id"] for case in cases}}
    return {"anomalies": [row for row in latest.values() if row.get("regression_status") != "PASSED"]}


@router.get("/summary")
def batch_summary(source_paths: list[str] = Query(default=[]), trial_user: str = Query(default="reviewer-001")) -> dict[str, Any]:
    _ensure_batch_user(trial_user)
    sources = v2._latest_by_id(v2._read_jsonl(v2.BATCH_SOURCE_REGISTER), "source_path")
    if source_paths:
        wanted = {_normalize_path(path) for path in source_paths}
        sources = {key: row for key, row in sources.items() if _normalize_path(row.get("source_path")) in wanted}
    cases = {key: row for key, row in v2._latest_by_id(v2._read_jsonl(v2.BATCH_CASES), "case_id").items() if not source_paths or _normalize_path(row.get("expected_source_path")) in {_normalize_path(path) for path in source_paths}}
    runs = {key: row for key, row in v2._latest_by_id(v2._read_jsonl(v2.BATCH_RUNS), "case_id").items() if key in cases}
    anomalies = [row for row in runs.values() if row.get("regression_status") != "PASSED"]
    return {"approved_source_count": sum(row.get("status") == "APPROVED" for row in sources.values()), "question_count": len(cases), "run_count": len(runs), "anomaly_count": len(anomalies), "passed_count": sum(row.get("regression_status") == "PASSED" for row in runs.values()), "formal_knowledge_publish": 0, "read_only": True}


@router.post("/uploads")
def start_upload(request: UploadStartRequest) -> dict[str, Any]:
    _ensure_batch_user(request.trial_user)
    upload_id = "UP_" + uuid.uuid4().hex
    token = secrets.token_urlsafe(32)
    session_root = (UPLOAD_ROOT / upload_id).resolve()
    session_root.mkdir(parents=True, exist_ok=False)
    v2._append_jsonl(UPLOAD_SESSIONS, {
        "event": "START",
        "upload_id": upload_id,
        "owner": request.trial_user,
        "token_hash": hashlib.sha256(token.encode("utf-8")).hexdigest(),
        "status": "OPEN",
        "file_count": 0,
        "total_bytes": 0,
        "created_at": v2._now(),
    })
    return {"upload_id": upload_id, "upload_token": token, "max_file_bytes": MAX_UPLOAD_FILE_BYTES, "max_total_bytes": MAX_UPLOAD_TOTAL_BYTES, "allowed_extensions": sorted(SUPPORTED_EXTENSIONS), "read_only_after_approval": True}


@router.put("/uploads/{upload_id}/file")
async def upload_file(upload_id: str, request: Request, relative_path: str = Query(...), x_upload_token: str = Header(default="")) -> dict[str, Any]:
    session = _upload_session(upload_id)
    _authorize_upload(session, x_upload_token)
    relative = _safe_relative_path(relative_path)
    if Path(relative).suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="UPLOAD_EXTENSION_NOT_ALLOWED")
    session_root = (UPLOAD_ROOT / upload_id).resolve()
    target = (session_root / Path(*PurePosixPath(relative).parts)).resolve()
    if not target.is_relative_to(session_root):
        raise HTTPException(status_code=400, detail="UPLOAD_PATH_TRAVERSAL")
    content_length = int(request.headers.get("content-length") or 0)
    if content_length > MAX_UPLOAD_FILE_BYTES:
        raise HTTPException(status_code=413, detail="UPLOAD_FILE_TOO_LARGE")
    current_total = sum(int(row.get("size") or 0) for row in v2._read_jsonl(UPLOAD_SESSIONS) if row.get("event") == "FILE" and row.get("upload_id") == upload_id)
    target.parent.mkdir(parents=True, exist_ok=True)
    written = 0
    try:
        with target.open("wb") as handle:
            async for chunk in request.stream():
                written += len(chunk)
                if written > MAX_UPLOAD_FILE_BYTES or current_total + written > MAX_UPLOAD_TOTAL_BYTES:
                    raise HTTPException(status_code=413, detail="UPLOAD_TOTAL_TOO_LARGE")
                handle.write(chunk)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    v2._append_jsonl(UPLOAD_SESSIONS, {"event": "FILE", "upload_id": upload_id, "owner": session["owner"], "relative_path": relative, "stored_path": str(target), "size": written, "uploaded_at": v2._now()})
    return {"uploaded": True, "upload_id": upload_id, "relative_path": relative, "size": written, "stored_path": str(target)}


@router.post("/uploads/{upload_id}/complete")
def complete_upload(upload_id: str, request: UploadCompleteRequest, x_upload_token: str = Header(default="")) -> dict[str, Any]:
    _ensure_batch_user(request.trial_user)
    session = _upload_session(upload_id)
    _authorize_upload(session, x_upload_token, request.trial_user)
    files = [row for row in v2._read_jsonl(UPLOAD_SESSIONS) if row.get("event") == "FILE" and row.get("upload_id") == upload_id]
    if not files:
        raise HTTPException(status_code=400, detail="UPLOAD_EMPTY")
    v2._append_jsonl(UPLOAD_SESSIONS, {"event": "COMPLETE", "upload_id": upload_id, "owner": session["owner"], "status": "COMPLETED", "file_count": len(files), "total_bytes": sum(int(row.get("size") or 0) for row in files), "completed_at": v2._now()})
    return {"completed": True, "upload_id": upload_id, "file_count": len(files), "files": [row["stored_path"] for row in files], "next_step": "preview_then_approve"}


def _upload_session(upload_id: str) -> dict[str, Any]:
    rows = [row for row in v2._read_jsonl(UPLOAD_SESSIONS) if row.get("upload_id") == upload_id]
    start = next((row for row in rows if row.get("event") == "START"), None)
    if start is None:
        raise HTTPException(status_code=404, detail="UPLOAD_SESSION_NOT_FOUND")
    latest = rows[-1]
    return {**start, **latest}


def _authorize_upload(session: dict[str, Any], token: str, owner: str | None = None) -> None:
    expected = str(session.get("token_hash") or "")
    actual = hashlib.sha256(token.encode("utf-8")).hexdigest() if token else ""
    if not expected or not hmac.compare_digest(expected, actual) or (owner is not None and owner != session.get("owner")):
        raise HTTPException(status_code=403, detail="UPLOAD_PERMISSION_DENIED")
    if session.get("status") == "COMPLETED":
        raise HTTPException(status_code=409, detail="UPLOAD_SESSION_COMPLETED")


def _safe_relative_path(value: str) -> str:
    normalized = str(value or "").replace("\\", "/").strip()
    path = PurePosixPath(normalized)
    if not normalized or path.is_absolute() or ":" in normalized.split("/", 1)[0] or any(part in {"", ".", ".."} for part in path.parts) or "\x00" in normalized:
        raise HTTPException(status_code=400, detail="UPLOAD_PATH_TRAVERSAL")
    return "/".join(path.parts)


def _discover_files(values: list[str], recursive: bool) -> tuple[list[Path], list[dict[str, str]]]:
    files: dict[str, Path] = {}
    errors = []
    for raw in values:
        value = str(raw or "").strip().strip('"')
        if not value:
            continue
        path = Path(value)
        if not path.exists():
            errors.append({"path": value, "reason": "路径不存在"})
            continue
        candidates = [path] if path.is_file() else (path.rglob("*") if recursive else path.glob("*"))
        for candidate in candidates:
            if candidate.is_file() and candidate.suffix.lower() in SUPPORTED_EXTENSIONS:
                files.setdefault(str(candidate.resolve()).casefold(), candidate.resolve())
    return sorted(files.values(), key=lambda item: str(item).casefold()), errors


def _source_probe(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {"source_path": str(path), "file_name": path.name, "file_type": path.suffix.lower(), "size": path.stat().st_size, "sha256": _sha256(path), "parse_probe": "deferred"}
    try:
        if path.suffix.lower() == ".docx":
            document = Document(path)
            result["parse_probe"] = "docx_readable"
            result["paragraphs"] = len(document.paragraphs)
            result["tables"] = len(document.tables)
        elif path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            result["parse_probe"] = "xlsx_readable"
            result["sheets"] = list(workbook.sheetnames)
            workbook.close()
        elif path.suffix.lower() == ".wps":
            result["parse_probe"] = "wps_com_read_only"
        elif path.suffix.lower() == ".pdf":
            pdf = load_pdf(path, "batch-probe-" + hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:24])
            result["parse_probe"] = "pdf_readable" if pdf.status in {"parsed", "empty"} else f"pdf_{pdf.status}"
            result["pages"] = len(pdf.blocks)
            result["value_creation_summary"] = any(extract_value_creation_summary(block.text) for block in pdf.blocks)
            if pdf.status == "read_error":
                result["error"] = pdf.error
        return result
    except Exception as error:
        result["parse_probe"] = "read_error"
        result["error"] = f"{type(error).__name__}: {error}"
        return result


def _approved_sources(paths: list[str]) -> list[dict[str, Any]]:
    rows = list(v2._latest_by_id(v2._read_jsonl(v2.BATCH_SOURCE_REGISTER), "source_path").values())
    wanted = {v2._normalize_source_path(path) for path in paths if str(path).strip()}
    return [row for row in rows if row.get("status") == "APPROVED" and (not wanted or v2._normalize_source_path(row.get("source_path")) in wanted)]


def _filter_cases(cases: list[dict[str, Any]], source_paths: list[str]) -> list[dict[str, Any]]:
    if not source_paths:
        return cases
    wanted = {_normalize_path(path) for path in source_paths}
    return [case for case in cases if _normalize_path(case.get("expected_source_path")) in wanted]


def _normalize_path(value: Any) -> str:
    return v2._normalize_source_path(value).casefold()


def _generate_cases(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".wps":
        return _generate_wps_cases(path)
    if path.suffix.lower() == ".docx":
        return _generate_docx_cases(path)
    if path.suffix.lower() == ".xlsx":
        return _generate_xlsx_cases(path)
    if path.suffix.lower() == ".pdf":
        return _generate_pdf_cases(path)
    return []


def _generate_pdf_cases(path: Path) -> list[dict[str, Any]]:
    document_id = "batch-probe-" + hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:24]
    result = load_pdf(path, document_id)
    summary = {}
    summary_page = "?"
    detail_rows = []
    detail_page = "?"
    for block in result.blocks:
        detail_rows.extend((row, block.location.get("page", "?")) for row in extract_value_creation_rows(block.text))
        candidate_summary = extract_value_creation_summary(block.text)
        if not summary and candidate_summary.get("合计"):
            summary = candidate_summary
            summary_page = block.location.get("page", "?")
    cases = []
    if summary:
        total = summary["合计"]
        stage_terms = [f"方案设计：{total[0]}条", f"初步设计：{total[1]}条", f"施工图设计：{total[2]}条", f"合计：{sum(total)}条"]
        stage_question = f"{path.stem}中，各专业、阶段数量统计的方案设计、初步设计、施工图设计和合计分别是多少条？"
        cases.append(_case(path, stage_question, f"第{summary_page}页", stage_terms, "PDF_VALUE_CREATION_STAGE_SUMMARY", "STRICT"))
        professional_terms = [f"{name}：{sum(values)}条" for name, values in summary.items() if name != "合计"]
        professional_question = f"{path.stem}中，各专业分别有多少条设计价值创造点？"
        cases.append(_case(path, professional_question, f"第{summary_page}页", professional_terms, "PDF_VALUE_CREATION_PROFESSION_COUNTS", "STRICT"))
        profession_count_question = f"{path.stem}里，包含了多少个专业？"
        cases.append(_case(path, profession_count_question, f"第{summary_page}页", [f"{len(summary) - 1}个专业"], "PDF_VALUE_CREATION_PROFESSION_COUNT", "STRICT"))
    if detail_rows:
        row, detail_page = detail_rows[0]
        location = f"第{detail_page}页"
        if row["value_item"]:
            cases.append(_case(path, f"{row['professional']}专业{row['stage']}的第{row['number']}条价值创造点是什么？", location, [row["value_item"]], "PDF_VALUE_CREATION_POINT_SAMPLE", "STRICT"))
        if row["applicability"]:
            cases.append(_case(path, f"{row['professional']}专业{row['stage']}的第{row['number']}条价值创造点，适用条件是什么？", location, [row["applicability"]], "PDF_VALUE_CREATION_APPLICABILITY_SAMPLE", "STRICT"))
    return cases


def _generate_wps_cases(path: Path) -> list[dict[str, Any]]:
    document_id = "batch-probe-" + hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:24]
    document = v2.V2TrialEngine._load_approved_wps_read_only(path, document_id).documents[0]
    text = "\n".join(chunk.text for chunk in document.chunks)
    names = _project_names(text)
    if not names:
        return []
    question = "2026年立项的设计示范项目有哪些"
    return [_case(path, question, "表1", names, "WPS_PROJECT_LIST", "STRICT")]


def _generate_docx_cases(path: Path) -> list[dict[str, Any]]:
    document = Document(path)
    project_name = _docx_project_name(document)
    cases = []
    risk_case_created = False
    for table_number, table in enumerate(document.tables, start=1):
        rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
        header_index = next((index for index, values in enumerate(rows) if "专业类别" in values and "价值创造策划点" in values and "价值创造分析" in values), None)
        if header_index is None:
            if any(any(marker in value for marker in ("风险内容项", "风险事项", "风险描述")) for values in rows[:4] for value in values) and any(any(marker in value for marker in ("应对措施", "化解建议", "应对策略")) for values in rows[:8] for value in values):
                subject = project_name or path.stem
                cases.append(_case(path, f"{subject}的设计风险识别清单有哪些风险和应对措施？", f"表{table_number}", [], "DOCX_RISK_OBSERVATION", "OBSERVE_ONLY"))
                risk_case_created = True
            continue
        items = [row for row in rows[header_index + 1 :] if row and row[0] and row[0] != "专业类别"]
        counts = Counter(row[0] for row in items)
        if not counts:
            continue
        subject = project_name or path.stem
        question = f"{subject}，设计策划中的设计价值创造清单，包含了哪几个专业，每个专业分别有多少条"
        terms = [f"{name}：{count}条" for name, count in counts.items()] + [f"合计：{len(items)}条"]
        cases.append(_case(path, question, f"表{table_number}", terms, "DOCX_VALUE_CREATION_COUNTS", "STRICT"))
    if not cases and not risk_case_created:
        return []
    return cases


def _generate_xlsx_cases(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    cases = []
    try:
        for worksheet in workbook.worksheets:
            rows = [[str(value).strip() if value is not None else "" for value in row] for row in worksheet.iter_rows(values_only=True)]
            for group_row_index, values in enumerate(rows):
                groups = [value for value in values if re.fullmatch(r"[\u3400-\u9fff]{1,8}专业设计参数", value)]
                if not groups:
                    continue
                group = groups[0]
                header_index = next((index for index in range(group_row_index, min(len(rows), group_row_index + 3)) if sum(bool(item) for item in rows[index]) >= len(groups) and not any("专业设计参数" in item for item in rows[index] if item)), None)
                if header_index is None:
                    continue
                start_column = values.index(group)
                fields = [value for value in rows[header_index][start_column:] if value]
                if len(fields) < 2:
                    continue
                cases.append(_case(path, f"{path.stem}中，{group}，包含哪些", f"{worksheet.title}，第{header_index + 1}行", fields, "XLSX_PROFESSIONAL_FIELDS", "STRICT"))
                break
    finally:
        workbook.close()
    return cases


def _project_names(text: str) -> list[str]:
    names = []
    for line in text.splitlines():
        cells = [value.strip() for value in line.strip().strip("|").split("|")]
        if len(cells) >= 2 and re.fullmatch(r"\d+", cells[0]) and cells[1]:
            names.append(cells[1])
    return list(dict.fromkeys(names))


def _docx_project_name(document: Document) -> str:
    if not document.tables:
        return ""
    for row in document.tables[0].rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if len(cells) >= 2 and cells[0] == "项目名称":
            return cells[1]
    return ""


def _case(path: Path, question: str, location: str, terms: list[str], generation_mode: str, verification_mode: str) -> dict[str, Any]:
    fingerprint = hashlib.sha256(f"{path}|{question}|{location}".encode("utf-8")).hexdigest()
    return {"case_id": "BQ_" + fingerprint[:16], "fingerprint": fingerprint, "question": question, "expected_source_path": str(path), "expected_source_location": location, "required_terms": terms, "verification_mode": verification_mode, "generation_mode": generation_mode, "created_at": v2._now(), "runtime_input": False, "formal_knowledge_publish": False}


def _not_ready_run(case: dict[str, Any], reviewer: str) -> dict[str, Any]:
    return {"run_id": "BR_" + uuid.uuid4().hex, "case_id": case["case_id"], "run_at": v2._now(), "reviewer": reviewer, "regression_status": "NOT_READY", "repair_status": "SOURCE_CLOSURE_REQUIRED", "reason": "来源尚未进入当前 Shadow 索引。", "source_runtime_status": "SOURCE_IDENTIFIED", "formal_knowledge_publish": False}


def _evaluate_case(case: dict[str, Any], result: dict[str, Any], reviewer: str) -> dict[str, Any]:
    answer = str(result.get("answer") or "")
    citations = result.get("citations") or []
    source_hit = any(v2._same_source_path(case["expected_source_path"], item.get("source_path")) for item in citations)
    expected_location = str(case.get("expected_source_location") or "").strip()
    location_hit = not expected_location or any(expected_location in str(item.get("display_location") or "") for item in citations)
    missing = [term for term in case.get("required_terms", []) if term not in answer]
    if case.get("verification_mode") == "OBSERVE_ONLY":
        status = "OBSERVED" if source_hit and location_hit and result.get("answer_status") != "SOURCE_SCOPE_MISSING" else "FAILED"
        repair_status = "HUMAN_CONFIRMATION_REQUIRED" if status == "OBSERVED" else "REPAIR_REQUIRED"
    else:
        passed = result.get("answer_status") == "ANSWERED" and source_hit and location_hit and not missing
        status = "PASSED" if passed else "FAILED"
        repair_status = "VALIDATED" if passed else "REPAIR_REQUIRED"
    return {"run_id": "BR_" + uuid.uuid4().hex, "case_id": case["case_id"], "run_at": v2._now(), "reviewer": reviewer, "answer_status": result.get("answer_status"), "dense_runtime": result.get("dense_runtime"), "citation_source_hit": source_hit, "citation_location_hit": location_hit, "missing_required_terms": missing, "regression_status": status, "repair_status": repair_status, "answer_excerpt": answer[:1200], "citation_ids": [item.get("citation_id") for item in citations], "provider_http_requests": result.get("provider_http_requests"), "runtime_input_injection": 0, "formal_knowledge_publish": False}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()
